# Criado por José Eduardo Santana Martins
# Este arquivo serve para exportar o ETP ou TR em Word (sobre os modelos originais) e em PDF.
"""Exportação do documento.

- **Word:** abre o modelo (`recursos/contratacoes/etp.docx` ou `tr.docx`, o mesmo do SGI SPI antigo: Verdana 10, margens e página),
  limpa o corpo e reescreve: título, processo, link, seções (`1. TÍTULO`, em negrito) e itens com o marcador jurídico em negrito
  (`1.1.`, `I -`, `a)`). A tabela estruturada do TR sai logo depois do item 1.1.
- **PDF:** o mesmo conteúdo com `reportlab` (parágrafos, listas e tabelas).
"""

import copy
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

from app.models.contratacoes import DocumentoContratacao, ItemContratacao
from app.services.contratacoes import conteudo, html_blocos
from app.services.contratacoes.html_blocos import Trecho
from app.services.documentos.pdf import ESTILO_TEXTO, DocumentoPdf

RECURSOS = Path(__file__).resolve().parents[2] / "recursos" / "contratacoes"
ROTULOS_TIPO = {"etp": "ETP", "tr": "TR"}
ALINHAR = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "left": WD_ALIGN_PARAGRAPH.LEFT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def itens_em_ordem(secao) -> list[ItemContratacao]:
    """Itens da seção em pré-ordem (pai antes dos filhos), respeitando a ordem dos irmãos."""
    filhos: dict = {}
    for i in secao.itens:
        filhos.setdefault(i.pai_id, []).append(i)
    saida: list[ItemContratacao] = []

    def descer(pai_id) -> None:
        for i in sorted(filhos.get(pai_id, []), key=lambda x: (x.ordem, str(x.id))):
            saida.append(i)
            descer(i.id)

    descer(None)
    return saida


def _blocos_do_item(item: ItemContratacao) -> list[tuple]:
    return html_blocos.blocos(item.conteudo_html or conteudo.html_do_texto(item.conteudo))


# --- Word ---------------------------------------------------------------------------------------

def _escrever_trechos(paragrafo, partes: list[Trecho], negrito_total: bool = False) -> None:
    for t in partes:
        run = paragrafo.add_run(t.texto if t.texto != "\n" else "")
        if t.texto == "\n":
            run.add_break()
        run.bold = True if (t.negrito or negrito_total) else None
        run.italic = True if t.italico else None
        run.underline = True if (t.sublinhado or t.link) else None
        run.font.superscript = True if t.sobrescrito else None
        run.font.subscript = True if t.subscrito else None
        if t.realce:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _paragrafo(d, alinhamento: str | None = None):
    p = d.add_paragraph()
    p.alignment = ALINHAR.get(alinhamento or "justify", WD_ALIGN_PARAGRAPH.JUSTIFY)
    return p


def _tabela_html(d, linhas: list[list[list[Trecho]]]) -> None:
    colunas = max((len(l) for l in linhas), default=0)
    if not linhas or not colunas:
        return
    tabela = d.add_table(rows=len(linhas), cols=colunas)
    tabela.style = "Table Grid"
    for r, linha in enumerate(linhas):
        for c, celula in enumerate(linha):
            p = tabela.cell(r, c).paragraphs[0]
            _escrever_trechos(p, celula)
    d.add_paragraph()


def _tabela_tr(d, linhas) -> None:
    cabecalho = ["Item", "Descrição", "CATMAT/CATSER", "Siafísico", "UF", "Quantidade"]
    tabela = d.add_table(rows=1 + len(linhas), cols=len(cabecalho))
    tabela.style = "Table Grid"
    for c, titulo in enumerate(cabecalho):
        p = tabela.cell(0, c).paragraphs[0]
        run = p.add_run(titulo)
        run.bold = True
    for r, l in enumerate(linhas, start=1):
        valores = [str(l.ordem), l.descricao, l.catser_catmat, l.siafisico, l.unidade, f"{l.quantidade_objeto:f}".rstrip("0").rstrip(".") or "0"]
        for c, v in enumerate(valores):
            tabela.cell(r, c).paragraphs[0].add_run(v)
    d.add_paragraph()


def gerar_word(documento: DocumentoContratacao) -> bytes:
    modelo = RECURSOS / f"{documento.tipo}.docx"
    if not modelo.is_file():
        raise FileNotFoundError(f"Modelo de exportação não encontrado: {modelo.name}")
    d = docx.Document(str(modelo))
    corpo = d.element.body
    sect = corpo.sectPr
    sect_copia = copy.deepcopy(sect) if sect is not None else None
    for filho in list(corpo):
        corpo.remove(filho)
    if sect_copia is not None:
        corpo.append(sect_copia)

    titulo = _paragrafo(d)
    r = titulo.add_run(f"{ROTULOS_TIPO[documento.tipo]} - {documento.nome}")
    r.bold, r.font.size = True, Pt(12)
    if documento.processo.strip():
        p = _paragrafo(d)
        p.add_run("Processo: ").bold = True
        p.add_run(documento.processo)
    if documento.link_sei:
        p = _paragrafo(d)
        p.add_run("Link: ").bold = True
        p.add_run(documento.link_sei)

    for secao in sorted(documento.secoes, key=lambda s: s.ordem):
        _paragrafo(d).add_run(f"{secao.ordem}. {secao.titulo}").bold = True
        marcas = conteudo.marcadores(secao.ordem, secao.itens)
        for item in itens_em_ordem(secao):
            marca = marcas.get(item.id, "")
            blocos = _blocos_do_item(item) or [("p", [], None)]
            primeiro = True
            for bloco in blocos:
                tipo, dados, extra = bloco
                if tipo == "tabela":
                    _tabela_html(d, dados)
                    continue
                p = _paragrafo(d, extra if tipo == "p" else None)
                if tipo == "li":
                    p.paragraph_format.left_indent = Pt(24)
                if primeiro and marca:
                    p.add_run(f"{marca} ").bold = True
                if tipo == "li":
                    p.add_run(f"{extra} ")
                primeiro = False
                _escrever_trechos(p, dados, negrito_total=item.tipo == "subsecao")
            if item.linhas_tabela:
                _tabela_tr(d, sorted(item.linhas_tabela, key=lambda x: x.ordem))
    saida = BytesIO()
    d.save(saida)
    return saida.getvalue()


# --- PDF ----------------------------------------------------------------------------------------

ESTILO_PDF = ParagraphStyle("contratacao", parent=ESTILO_TEXTO, fontSize=9.5, leading=13, alignment=4, spaceAfter=3)
ESTILO_SECAO_PDF = ParagraphStyle("contratacao_secao", parent=ESTILO_PDF, fontName="Helvetica-Bold", fontSize=11, leading=15, spaceBefore=9, spaceAfter=4, alignment=0)


def _rml(partes: list[Trecho], negrito_total: bool = False) -> str:
    saida = []
    for t in partes:
        if t.texto == "\n":
            saida.append("<br/>")
            continue
        s = escape(t.texto)
        if t.negrito or negrito_total:
            s = f"<b>{s}</b>"
        if t.italico:
            s = f"<i>{s}</i>"
        if t.sublinhado or t.link:
            s = f"<u>{s}</u>"
        if t.realce:
            s = f'<font backColor="#fff2a8">{s}</font>'
        if t.sobrescrito:
            s = f"<super>{s}</super>"
        if t.subscrito:
            s = f"<sub>{s}</sub>"
        saida.append(s)
    return "".join(saida)


def gerar_pdf(documento: DocumentoContratacao, autor: str = "") -> bytes:
    pdf = DocumentoPdf(titulo=f"{ROTULOS_TIPO[documento.tipo]} - {documento.nome}", subtitulo=f"Processo: {documento.processo or '—'}", paisagem=False, autor=autor)
    for secao in sorted(documento.secoes, key=lambda s: s.ordem):
        pdf.bloco(Paragraph(escape(f"{secao.ordem}. {secao.titulo}"), ESTILO_SECAO_PDF))
        marcas = conteudo.marcadores(secao.ordem, secao.itens)
        for item in itens_em_ordem(secao):
            marca = marcas.get(item.id, "")
            primeiro = True
            for tipo, dados, extra in (_blocos_do_item(item) or [("p", [], None)]):
                if tipo == "tabela":
                    linhas = [[Paragraph(_rml(c) or " ", ESTILO_PDF) for c in l] for l in dados]
                    if linhas:
                        tabela = Table(linhas, hAlign="LEFT")
                        tabela.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                        pdf.bloco(tabela, Spacer(1, 2 * mm))
                    continue
                prefixo = f"<b>{escape(marca)}</b> " if primeiro and marca else ""
                if tipo == "li":
                    prefixo += f"{escape(str(extra))} "
                estilo = ParagraphStyle("li", parent=ESTILO_PDF, leftIndent=14) if tipo == "li" else ESTILO_PDF
                pdf.bloco(Paragraph(prefixo + _rml(dados, item.tipo == "subsecao"), estilo))
                primeiro = False
            if item.linhas_tabela:
                cab = ["Item", "Descrição", "CATMAT/CATSER", "Siafísico", "UF", "Qtd."]
                linhas = [[Paragraph(f"<b>{c}</b>", ESTILO_PDF) for c in cab]] + [
                    [Paragraph(escape(str(v)), ESTILO_PDF) for v in (l.ordem, l.descricao, l.catser_catmat, l.siafisico, l.unidade, f"{l.quantidade_objeto:f}".rstrip("0").rstrip("."))]
                    for l in sorted(item.linhas_tabela, key=lambda x: x.ordem)]
                tabela = Table(linhas, hAlign="LEFT", colWidths=[12 * mm, 70 * mm, 28 * mm, 24 * mm, 14 * mm, 20 * mm])
                tabela.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                pdf.bloco(tabela, Spacer(1, 2 * mm))
    return pdf.gerar()
