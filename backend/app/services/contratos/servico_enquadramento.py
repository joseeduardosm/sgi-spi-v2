# Criado por José Eduardo Santana Martins
# Este arquivo serve para conferir a prorrogação contra o checklist de 37 itens do Parecer Referencial NLC nº 1/2026 e exportá-lo em planilha.
"""Enquadramento da prorrogação no Parecer Referencial.

Cada item tem a **resposta da equipe** (`sim`, `nao`, `na`) e, quando o SGI consegue conferir sozinho, a **conferência automática** com a conta.
Vale a resposta da equipe; sem ela, vale a conferência; sem nenhuma das duas, o item está pendente. Qualquer item `nao` liga o aviso de que o caso
sai do Parecer Referencial e precisa da Consultoria Jurídica; o aviso **não bloqueia** o registro (decisão do usuário: só avisa).
"""

import io
import uuid
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.contratos import Contrato, ProcessoProrrogacao
from app.models.usuario import Usuario
from app.schemas.contratos.alteracoes import (
    ConferenciaAutomatica,
    GravacaoEnquadramento,
    GrupoEnquadramento,
    ItemEnquadramento,
    LeituraEnquadramento,
)
from app.services.contratos import calculos, servico_parametros
from app.services.contratos.catalogo_parecer_referencial import ITENS
from app.services.contratos.erros import ErroRegraContrato
from app.services.contratos.servico_contratos import exigir_edicao, obter_contrato
from app.services.servico_auditoria import auditar

MODELO_XLSX = Path(__file__).resolve().parents[2] / "recursos" / "prorrogacao" / "checklist-parecer-referencial.xlsx"
ROTULO = {"sim": "Sim", "nao": "Não", "na": "Não se aplica", "": ""}
DIREITA_14133 = date(2021, 4, 1)


def _conferencias(sessao: Session, contrato: Contrato, meses: int | None, hoje: date) -> dict[int, ConferenciaAutomatica]:
    """Conferências que o SGI faz sozinho (número do item → resultado e conta)."""
    c: dict[int, ConferenciaAutomatica] = {}

    def marcar(numero: int, resultado: str, explicacao: str) -> None:
        c[numero] = ConferenciaAutomatica(resultado=resultado, explicacao=explicacao)

    if contrato.data_inicio < DIREITA_14133:
        marcar(1, "nao", f"O contrato começou em {contrato.data_inicio:%d/%m/%Y}, antes da Lei nº 14.133/2021 (01/04/2021).")
    marcar(3, "sim" if hoje <= contrato.data_fim else "nao",
           f"A vigência atual termina em {contrato.data_fim:%d/%m/%Y} e hoje é {hoje:%d/%m/%Y}: "
           + ("o aditivo ainda pode ser assinado dentro da vigência." if hoje <= contrato.data_fim else "a vigência já terminou, o parecer não vale para vigência expirada."))
    if meses:
        marcar(5, "sim" if meses <= contrato.vigencia_inicial_meses else "nao",
               f"Prorrogação de {meses} mês(es) contra {contrato.vigencia_inicial_meses} mês(es) de vigência inicial.")
        total = contrato.vigencia_inicial_meses + sum(p.meses for p in contrato.prorrogacoes) + meses
        marcar(14, "sim" if total <= 120 else "nao",
               f"{contrato.vigencia_inicial_meses} (inicial) + {sum(p.meses for p in contrato.prorrogacoes)} (prorrogações anteriores) + {meses} (esta) = {total} mês(es), limite de 120 (art. 107).")
    limite = servico_parametros.valor("parecer_referencial_vigencia_ate")
    marcar(8, "sim" if hoje <= limite else "nao", f"O Parecer Referencial vale até {limite:%d/%m/%Y} (Parâmetros de Contratos); hoje é {hoje:%d/%m/%Y}.")
    if not contrato.exige_garantia:
        for numero in (27, 28):
            marcar(numero, "na", "O contrato está cadastrado sem exigência de garantia.")
    if contrato.modalidade_contratacao == "licitacao":
        for numero in (29, 30, 31):
            marcar(numero, "na", "O contrato veio de licitação: itens da dispensa por valor não se aplicam.")
    return c


def _respostas(processo: ProcessoProrrogacao | None) -> dict[int, dict]:
    bruto = ((processo.enquadramento if processo else None) or {}).get("respostas", {})
    return {int(k): v for k, v in bruto.items()}


def consultar(sessao: Session, contrato: Contrato, processo: ProcessoProrrogacao | None, hoje: date | None = None) -> LeituraEnquadramento:
    """Checklist com as respostas, as conferências automáticas e o resumo."""
    hoje = hoje or date.today()
    auto = _conferencias(sessao, contrato, processo.meses if processo else None, hoje)
    respostas = _respostas(processo)
    grupos: dict[str, list[ItemEnquadramento]] = {}
    contagem = {"sim": 0, "na": 0, "nao": 0, "": 0}
    motivos: list[str] = []
    for item in ITENS:
        r = respostas.get(item.numero, {})
        resposta = r.get("resposta", "")
        automatica = auto.get(item.numero)
        efetivo = resposta or (automatica.resultado if automatica else "")
        contagem[efetivo] += 1
        if efetivo == "nao":
            motivos.append(f"Item {item.numero}: {item.texto}")
        grupos.setdefault(item.grupo, []).append(ItemEnquadramento(
            numero=item.numero, texto=item.texto, referencia=item.referencia, automatica=automatica, resposta=resposta, doc=r.get("doc", ""), obs=r.get("obs", ""),
            efetivo=efetivo, contestado=bool(resposta and automatica and resposta != automatica.resultado)))
    return LeituraEnquadramento(
        grupos=[GrupoEnquadramento(nome=n, itens=i) for n, i in grupos.items()], total=len(ITENS), atendidos=contagem["sim"] + contagem["na"],
        nao_atendidos=contagem["nao"], pendentes=contagem[""], precisa_analise_juridica=contagem["nao"] > 0, motivos=motivos)


def salvar(sessao: Session, contrato_id: uuid.UUID, dados: GravacaoEnquadramento, autor: Usuario) -> LeituraEnquadramento:
    """Grava as respostas (substitui as anteriores; cria o rascunho se ainda não houver)."""
    from app.services.contratos.servico_prorrogacao import _rascunho_ou_novo

    contrato = obter_contrato(sessao, contrato_id)
    exigir_edicao(sessao, contrato, autor)
    validos = {i.numero for i in ITENS}
    if any(n not in validos for n in dados.respostas):
        raise ErroRegraContrato("Item desconhecido no checklist do Parecer Referencial (são 37, numerados de 1 a 37).")
    processo = _rascunho_ou_novo(sessao, contrato, autor)
    respostas = {str(n): r.model_dump() for n, r in dados.respostas.items() if r.resposta or r.doc.strip() or r.obs.strip()}
    processo.enquadramento = {"respostas": respostas}
    leitura = consultar(sessao, contrato, processo)
    auditar(sessao, autor.login, "contrato.prorrogacao.enquadramento", f"Contrato {contrato.numero}", autor_id=autor.id, alvo_tipo="contrato", alvo_id=contrato.id,
            dados={"atendidos": leitura.atendidos, "nao_atendidos": leitura.nao_atendidos, "pendentes": leitura.pendentes})
    sessao.commit()
    return leitura


def fotografia(sessao: Session, contrato: Contrato, processo: ProcessoProrrogacao, assinada_em: date | None = None) -> dict:
    """Resumo gravado na prorrogação registrada: respostas, o que valia em cada item e se precisava de análise jurídica (conferido na data da assinatura)."""
    leitura = consultar(sessao, contrato, processo, assinada_em)
    return {
        "respostas": (processo.enquadramento or {}).get("respostas", {}),
        "efetivos": {str(i.numero): i.efetivo for g in leitura.grupos for i in g.itens},
        "precisa_juridico": leitura.precisa_analise_juridica, "motivos": leitura.motivos,
    }


def planilha(sessao: Session, contrato: Contrato, processo: ProcessoProrrogacao | None) -> bytes:
    """Checklist no formato do Modelo 6 da SPI (planilha original com as respostas e os documentos preenchidos)."""
    leitura = consultar(sessao, contrato, processo)
    livro = load_workbook(MODELO_XLSX)
    folha = livro.active
    por_numero = {i.numero: i for g in leitura.grupos for i in g.itens}
    for item in ITENS:
        lido = por_numero[item.numero]
        folha.cell(row=item.linha, column=3, value=ROTULO[lido.efetivo] or None)
        documento = lido.doc.strip() or None
        if lido.obs.strip():
            documento = f"{documento}\n{lido.obs.strip()}" if documento else lido.obs.strip()
        folha.cell(row=item.linha, column=4, value=documento)
    folha.cell(row=1, column=2, value=f"Contrato {contrato.numero}" + (" — ATENÇÃO: caso fora do Parecer Referencial, precisa de análise da Consultoria Jurídica" if leitura.precisa_analise_juridica else ""))
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()
