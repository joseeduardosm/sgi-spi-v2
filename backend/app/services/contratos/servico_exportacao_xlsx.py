# Criado por José Eduardo Santana Martins
# Este arquivo serve para gerar a planilha XLSX de um contrato já preenchida, no formato da importação.
"""Exportação do contrato para XLSX ("Checklist de Alimentação do Sistema de Contratos").

Preenche uma cópia do mesmo modelo usado na importação (`app/recursos/modelo-importacao-contrato.xlsx`),
então a planilha baixada pode ser importada de novo. Detalhes:
- o preposto é o primeiro ativo da empresa (o contrato não guarda um preposto próprio);
- a equipe (gestor, fiscais) sai com os nomes vigentes; a importação ignora essas linhas;
- as linhas de legenda do modelo (abaixo dos títulos dos itens) são trocadas pelos itens do contrato.
"""

from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.models.contratos import Contrato
from app.services.contratos.servico_contratos import designacoes_vigentes

MODELO = Path(__file__).resolve().parents[2] / "recursos" / "modelo-importacao-contrato.xlsx"

PERIODICIDADES = {1: "Mensal", 2: "Bimestral", 3: "Trimestral", 6: "Semestral", 12: "Anual"}
MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
TIPOS = {"continuo": "Contínuo", "sob_demanda": "Sob demanda"}
# Linha do modelo → papel da equipe
LINHAS_EQUIPE = {22: "gestor", 23: "gestor_suplente", 24: "fiscal_administrativo", 25: "fiscal_administrativo_suplente",
                  26: "fiscal_tecnico", 27: "fiscal_tecnico_suplente"}
LINHA_TITULOS_ITENS = 29


def _mascara(documento: str, formato: str) -> str:
    """CNPJ (`00.000.000/0000-00`) ou CPF (`000.000.000-00`) a partir dos dígitos; outro tamanho fica como está."""
    return formato.format(*documento) if len(documento) == formato.count("{}") else documento


def _numero(valor) -> float | int:
    """Decimal para a célula, sem casas inúteis."""
    return int(valor) if valor == valor.to_integral_value() else float(valor)


def gerar(contrato: Contrato) -> bytes:
    """Planilha do contrato com cabeçalho, equipe vigente e itens financeiros."""
    livro = load_workbook(MODELO)
    folha = livro.worksheets[0]
    empresa = contrato.empresa
    preposto = next((p for p in empresa.prepostos if p.ativo), None)

    campos = {
        2: contrato.numero,
        3: _mascara(empresa.cnpj, "{}{}.{}{}{}.{}{}{}/{}{}{}{}-{}{}"), 4: empresa.razao_social, 5: empresa.nome_fantasia, 6: empresa.endereco,
        7: preposto.nome if preposto else "", 8: _mascara(preposto.cpf, "{}{}{}.{}{}{}.{}{}{}-{}{}") if preposto else "",
        9: preposto.email if preposto else "", 10: preposto.telefone if preposto else "",
        11: contrato.apelido, 12: contrato.data_inicio, 13: contrato.vigencia_inicial_meses, 14: contrato.vigencia_maxima_meses,
        15: PERIODICIDADES.get(contrato.periodicidade_meses, contrato.periodicidade_meses), 16: MESES[contrato.mes_reajuste - 1],
        17: contrato.objeto,
        18: contrato.sei_gestao_numero, 19: contrato.sei_gestao_link, 20: contrato.sei_execucao_numero, 21: contrato.sei_execucao_link,
    }
    for linha, valor in campos.items():
        celula = folha.cell(linha, 3, valor)
        if isinstance(valor, date):
            celula.number_format = "dd/mm/yyyy"
    # Equipe vigente: mais de uma pessoa no mesmo papel vira lista separada por vírgula
    vigentes = designacoes_vigentes(contrato)
    for linha, papel in LINHAS_EQUIPE.items():
        folha.cell(linha, 3, ", ".join(d.nome_usuario for d in vigentes if d.papel == papel))

    # Itens: as linhas de legenda do modelo (tipo e faturamento possíveis) dão lugar aos itens do contrato
    primeira = LINHA_TITULOS_ITENS + 1
    for linha in range(primeira, max(folha.max_row, primeira) + 1):
        for coluna in range(1, 12):
            folha.cell(linha, coluna).value = None
    for linha, item in enumerate(sorted(contrato.itens, key=lambda i: i.ordem), start=primeira):
        valores = [item.descricao, TIPOS.get(item.tipo, item.tipo), "Pró-rata" if item.calcula_pro_rata else "Sempre Integral",
                   item.codigo_classe, item.codigo_natureza_despesa, item.codigo_siafisico, item.codigo_catmat_catser,
                   _numero(item.quantidade_mensal), _numero(item.quantidade_total), _numero(item.valor_unitario), item.unidade_fornecimento]
        for coluna, valor in enumerate(valores, start=1):
            folha.cell(linha, coluna, valor)

    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()
