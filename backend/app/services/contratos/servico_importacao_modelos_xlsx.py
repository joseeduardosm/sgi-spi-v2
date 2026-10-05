# Criado por José Eduardo Santana Martins
# Este arquivo serve para gerar, ler e gravar planilhas XLSX de checklists e formulários de avaliação.
"""Importação de checklists e formulários de avaliação por planilha XLSX.

Planilhas (o modelo é gerado aqui mesmo, com exemplo e instruções):
- **Checklist** (aba "Checklist"): nome na célula ao lado de "Nome do checklist" e, abaixo, a tabela
  `Documento | Observação | Obrigatório | Com validade | Vale para outros contratos` (Sim/Não).
- **Formulário** (abas "Formulário", "Escala", "Faixas" e "Itens"): nome ao lado de "Nome do
  formulário"; escala `Nota | Legenda`; faixas `Mínimo | Máximo | Percentual liberado | Notas zero`;
  itens `Grupo | Item | Descrição | Peso`. Grupo em branco repete o de cima.

Colunas e abas são achadas pelo título (sem acento, em minúsculas), não pela posição. Cada problema
vira um erro com a linha da planilha. As regras de negócio (escala crescente, pesos somando 100%...)
são as dos próprios schemas de gravação, sem duplicação.

Fluxo: `ler_*` só lê e valida; `importar_*` repete a leitura e, sem erros, grava a versão inativa no
contrato ou o modelo global.
"""

import re
from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.core.erros import _mensagem_validacao
from app.models.usuario import Usuario
from app.schemas.contratos.execucao import (
    GravacaoChecklist,
    GravacaoFormulario,
    GravacaoModelo,
)
from app.schemas.contratos.importacao import (
    DocumentoPrevia,
    ErroImportacao,
    FaixaPrevia,
    GrupoPrevia,
    ItemFormularioPrevia,
    NotaPrevia,
    PreviaImportacaoModelo,
)
from app.services.contratos import servico_configuracao_execucao as configuracao
from app.services.contratos.erros import ErroRegraContrato
from app.services.contratos.servico_importacao_xlsx import TAMANHO_MAXIMO, ErroPlanilha, converter_decimal, normalizar, texto
from app.services.servico_auditoria import auditar

# Colunas de cada tabela: título normalizado → (chave, rótulo exibido nas mensagens)
COLUNAS_DOCUMENTOS = {
    "documento": ("nome", "Documento"),
    "observacao": ("observacao", "Observação"),
    "obrigatorio": ("obrigatorio", "Obrigatório"),
    "com validade": ("com_validade", "Com validade"),
    "vale para outros contratos": ("vale_outros_contratos", "Vale para outros contratos"),
}
COLUNAS_ESCALA = {"nota": ("valor", "Nota"), "legenda": ("legenda", "Legenda")}
COLUNAS_FAIXAS = {
    "minimo": ("minimo", "Mínimo"),
    "maximo": ("maximo", "Máximo"),
    "percentual liberado": ("percentual", "Percentual liberado"),
    "notas zero": ("notas_zero", "Notas zero"),
}
COLUNAS_ITENS = {
    "grupo": ("grupo", "Grupo"),
    "item": ("nome", "Item"),
    "descricao": ("descricao", "Descrição"),
    "peso": ("peso", "Peso"),
}
# Colunas que podem faltar ou ficar em branco
OPCIONAIS = {"observacao", "obrigatorio", "com_validade", "vale_outros_contratos", "maximo", "notas_zero", "descricao"}


# ---------------------------------------------------------------------------------------------
# Leitura
# ---------------------------------------------------------------------------------------------

@dataclass
class Leitura:
    """Resultado da leitura: linhas já convertidas, erros, avisos e (sem erros) o corpo de gravação."""

    nome: str = ""
    documentos: list[dict[str, Any]] = field(default_factory=list)
    escala: list[dict[str, Any]] = field(default_factory=list)
    faixas: list[dict[str, Any]] = field(default_factory=list)
    itens: list[dict[str, Any]] = field(default_factory=list)
    erros: list[ErroImportacao] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    checklist: GravacaoChecklist | None = None
    formulario: GravacaoFormulario | None = None

    def erro(self, campo: str, mensagem: str, linha: int | None = None) -> None:
        """Registra um erro (um por campo e linha basta)."""
        if not any(e.campo == campo and e.linha == linha for e in self.erros):
            self.erros.append(ErroImportacao(linha=linha, campo=campo, mensagem=mensagem))


def _abrir_livro(conteudo: bytes):
    """Abre o arquivo; vazio, grande demais ou que não é XLSX vira erro de regra (400)."""
    if not conteudo:
        raise ErroRegraContrato("O arquivo está vazio.")
    if len(conteudo) > TAMANHO_MAXIMO:
        raise ErroRegraContrato("O arquivo passa de 5 MB.")
    try:
        return load_workbook(BytesIO(conteudo), data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError) as erro:
        raise ErroRegraContrato("Não foi possível ler o arquivo. Envie uma planilha no formato .xlsx.") from erro


def _aba(livro, nome: str, leitura: Leitura):
    """Aba pelo nome normalizado; ausente vira erro do arquivo."""
    for folha in livro.worksheets:
        if normalizar(folha.title) == nome:
            return folha
    leitura.erro(f"Aba \"{nome.capitalize()}\"", "A aba não foi encontrada. Use o modelo baixado no sistema.")
    return None


def _nome(folha, rotulo: str, leitura: Leitura) -> str:
    """Valor da célula à direita do rótulo (ex.: "Nome do checklist"), procurado nas primeiras linhas."""
    for linha in folha.iter_rows(min_row=1, max_row=15):
        for celula in linha:
            if normalizar(celula.value) == rotulo:
                valor = texto(folha.cell(celula.row, celula.column + 1).value)
                if not valor:
                    leitura.erro(rotulo.capitalize(), "Informe o nome.", celula.row)
                return valor
    leitura.erro(rotulo.capitalize(), "O rótulo não foi encontrado na planilha. Use o modelo baixado no sistema.")
    return ""


def _tabela(folha, colunas: dict, titulo_aba: str, leitura: Leitura) -> list[tuple[int, dict[str, Any]]]:
    """Linhas de dados (número da linha, valores por chave) abaixo da linha de títulos.

    A linha de títulos é a primeira que traz o título da primeira coluna da tabela; as colunas são
    achadas pelo título. Linhas totalmente em branco são ignoradas.
    """
    primeira = next(iter(colunas))
    for linha in folha.iter_rows():
        posicoes = {}
        for celula in linha:
            chave = colunas.get(normalizar(celula.value))
            if chave:
                posicoes.setdefault(chave[0], celula.column)
        if next(iter(colunas.values()))[0] in posicoes:
            titulos = linha[0].row
            break
    else:
        leitura.erro(f"Aba \"{titulo_aba}\"", f"A coluna \"{colunas[primeira][1]}\" não foi encontrada.")
        return []
    rotulos = {c: r for c, r in colunas.values()}
    for chave in rotulos:
        if chave not in posicoes and chave not in OPCIONAIS:
            leitura.erro(f"Aba \"{titulo_aba}\"", f"A coluna \"{rotulos[chave]}\" não foi encontrada.", titulos)
    linhas = []
    for n in range(titulos + 1, folha.max_row + 1):
        valores = {chave: folha.cell(n, coluna).value for chave, coluna in posicoes.items()}
        if any(texto(v) for v in valores.values()):
            linhas.append((n, valores))
    return linhas


def _sim_nao(valor: Any, padrao: bool) -> bool:
    """Sim/Não da célula; em branco vale o padrão."""
    nome = normalizar(valor)
    if not nome:
        return padrao
    if nome in ("sim", "s", "x", "verdadeiro", "1"):
        return True
    if nome in ("nao", "n", "falso", "0"):
        return False
    raise ValueError("use Sim ou Não")


def _decimal(valor: Any) -> Decimal | None:
    """Decimal com 2 casas (formato brasileiro aceito); em branco vira nulo."""
    return None if not texto(valor) else converter_decimal(valor, 2)


def _inteiro(valor: Any) -> int | None:
    """Número inteiro; em branco vira nulo."""
    if not texto(valor):
        return None
    numero = converter_decimal(valor)
    if numero != numero.to_integral_value():
        raise ValueError("deve ser um número inteiro")
    return int(numero)


def _converter(leitura: Leitura, linha: int, valores: dict[str, Any], rotulos: dict[str, str],
               conversores: dict[str, Any]) -> dict[str, Any]:
    """Converte as colunas (as sem conversor são texto); erros apontam linha e rótulo da coluna."""
    saida = {}
    for chave, valor in valores.items():
        try:
            saida[chave] = conversores[chave](valor) if chave in conversores else texto(valor)
        except (ValueError, ArithmeticError) as erro:
            leitura.erro(rotulos[chave], str(erro) or "valor inválido", linha)
    saida["linha"] = linha
    return saida


def _erros_validacao(leitura: Leitura, erro: ValidationError, linhas: dict[tuple, int], rotular) -> None:
    """Converte os erros do Pydantic em erros com a linha da planilha.

    `linhas`: caminho do dado (prefixo do `loc`) → linha; `rotular(loc)` devolve o rótulo da coluna.
    """
    for detalhe in erro.errors():
        local = tuple(detalhe.get("loc", ()))
        mensagem = _mensagem_validacao(detalhe)
        # Linha mais específica: o caminho inteiro ou o mais longo prefixo conhecido
        linha = next((linhas[local[:n]] for n in range(len(local), 0, -1) if local[:n] in linhas), None)
        if linha is None and (grupo := re.search(r'grupo "(.*)"', mensagem)):
            # Regra do grupo (soma dos pesos): aponta a primeira linha dele
            linha = linhas.get(("grupo", grupo.group(1)))
        leitura.erro(rotular(local), mensagem, linha)


def _rotulo_checklist(local: tuple) -> str:
    """Rótulo da coluna de um erro do checklist."""
    if local == ("nome",):
        return "Nome do checklist"
    return {c: r for c, r in COLUNAS_DOCUMENTOS.values()}.get(str(local[-1]), "Documento")


def _rotulo_formulario(local: tuple) -> str:
    """Rótulo da coluna de um erro do formulário (a aba depende da posição do erro)."""
    campo = str(local[-1])
    if local == ("nome",):
        return "Nome do formulário"
    secao = local[1] if len(local) > 1 else None
    if secao == "escala":
        return {c: r for c, r in COLUNAS_ESCALA.values()}.get(campo, "Escala")
    if secao == "faixas":
        return {c: r for c, r in COLUNAS_FAIXAS.values()}.get(campo, "Faixas")
    if secao == "grupos" and len(local) >= 5:
        return {c: r for c, r in COLUNAS_ITENS.values()}.get(campo, "Item")
    return "Grupo" if secao == "grupos" else "Formulário"


def ler_checklist(conteudo: bytes) -> Leitura:
    """Lê e valida a planilha do checklist."""
    leitura = Leitura()
    livro = _abrir_livro(conteudo)
    folha = _aba(livro, "checklist", leitura)
    if folha is None:
        return leitura
    leitura.nome = _nome(folha, "nome do checklist", leitura)
    rotulos = {c: r for c, r in COLUNAS_DOCUMENTOS.values()}
    conversores = {"obrigatorio": lambda v: _sim_nao(v, True), "com_validade": lambda v: _sim_nao(v, False), "vale_outros_contratos": lambda v: _sim_nao(v, False)}
    linhas = {}
    for n, valores in _tabela(folha, COLUNAS_DOCUMENTOS, "Checklist", leitura):
        documento = _converter(leitura, n, valores, rotulos, conversores)
        if not documento.get("nome"):
            leitura.erro("Documento", "Informe o nome do documento.", n)
        leitura.documentos.append(documento)
        linhas[("itens", len(leitura.documentos) - 1)] = n
    if not leitura.documentos:
        leitura.erro("Documento", "Informe ao menos um documento.")
    if leitura.erros:
        return leitura
    try:
        leitura.checklist = GravacaoChecklist(
            nome=leitura.nome, itens=[{k: v for k, v in d.items() if k != "linha"} for d in leitura.documentos]
        )
    except ValidationError as erro:
        _erros_validacao(leitura, erro, linhas, _rotulo_checklist)
    return leitura


def ler_formulario(conteudo: bytes) -> Leitura:
    """Lê e valida as quatro abas do formulário de avaliação."""
    leitura = Leitura()
    livro = _abrir_livro(conteudo)
    abas = {n: _aba(livro, n, leitura) for n in ("formulario", "escala", "faixas", "itens")}
    if any(a is None for a in abas.values()):
        return leitura
    leitura.nome = _nome(abas["formulario"], "nome do formulario", leitura)

    # Escala: um erro de conversão aponta a linha; a ordem e as repetições são regra do schema
    rotulos = {c: r for c, r in COLUNAS_ESCALA.values()}
    linhas: dict[tuple, int] = {}
    for n, valores in _tabela(abas["escala"], COLUNAS_ESCALA, "Escala", leitura):
        leitura.escala.append(_converter(leitura, n, valores, rotulos, {"valor": _decimal}))
        linhas[("definicao", "escala", len(leitura.escala) - 1)] = n
    # Faixas
    rotulos_faixa = {c: r for c, r in COLUNAS_FAIXAS.values()}
    for n, valores in _tabela(abas["faixas"], COLUNAS_FAIXAS, "Faixas", leitura):
        leitura.faixas.append(_converter(leitura, n, valores, rotulos_faixa,
                                         {"minimo": _decimal, "maximo": _decimal, "percentual": _decimal, "notas_zero": _inteiro}))
        linhas[("definicao", "faixas", len(leitura.faixas) - 1)] = n
    # Itens: agrupados pelo nome do grupo, na ordem em que aparecem (grupo em branco repete o anterior)
    rotulos_item = {c: r for c, r in COLUNAS_ITENS.values()}
    grupos: dict[str, list[dict[str, Any]]] = {}
    atual = ""
    for n, valores in _tabela(abas["itens"], COLUNAS_ITENS, "Itens", leitura):
        item = _converter(leitura, n, valores, rotulos_item, {"peso": _decimal})
        atual = item.get("grupo") or atual
        if not atual:
            leitura.erro("Grupo", "Informe o grupo (a primeira linha não pode ficar sem grupo).", n)
            continue
        item["grupo"] = atual
        grupos.setdefault(atual, []).append(item)
        linhas.setdefault(("grupo", atual), n)
        leitura.itens.append(item)
    for g, (nome, itens) in enumerate(grupos.items()):
        for i, item in enumerate(itens):
            linhas[("definicao", "grupos", g, "itens", i)] = item["linha"]
        linhas[("definicao", "grupos", g)] = itens[0]["linha"]
    if not leitura.escala:
        leitura.erro("Escala", "Informe as notas da escala (ao menos duas).")
    if not leitura.faixas:
        leitura.erro("Faixas", "Informe ao menos uma faixa de liberação.")
    if not leitura.itens and not leitura.erros:
        leitura.erro("Itens", "Informe ao menos um item de avaliação.")
    if leitura.erros:
        return leitura

    def sem_linha(d: dict[str, Any]) -> dict[str, Any]:
        return {k: v for k, v in d.items() if k not in ("linha", "grupo")}

    try:
        leitura.formulario = GravacaoFormulario(nome=leitura.nome, definicao={
            "escala": [sem_linha(e) for e in leitura.escala],
            "faixas": [sem_linha(f) for f in leitura.faixas],
            "grupos": [{"nome": nome, "itens": [sem_linha(i) for i in itens]} for nome, itens in grupos.items()],
        })
    except ValidationError as erro:
        _erros_validacao(leitura, erro, linhas, _rotulo_formulario)
    return leitura


# ---------------------------------------------------------------------------------------------
# Prévia e gravação
# ---------------------------------------------------------------------------------------------

def previa(tipo: str, conteudo: bytes) -> PreviaImportacaoModelo:
    """Lê e valida a planilha sem gravar nada."""
    leitura = ler_checklist(conteudo) if tipo == "checklist" else ler_formulario(conteudo)
    grupos: dict[str, list[ItemFormularioPrevia]] = {}
    for i in leitura.itens:
        grupos.setdefault(i["grupo"], []).append(ItemFormularioPrevia(
            linha=i["linha"], nome=i.get("nome", ""), descricao=i.get("descricao", ""), peso=i.get("peso")))
    return PreviaImportacaoModelo(
        tipo=tipo,
        nome=leitura.nome,
        documentos=[DocumentoPrevia(linha=d["linha"], nome=d.get("nome", ""), observacao=d.get("observacao", ""),
                                    obrigatorio=d.get("obrigatorio", True), com_validade=d.get("com_validade", False),
                                    vale_outros_contratos=d.get("vale_outros_contratos", False))
                    for d in leitura.documentos],
        escala=[NotaPrevia(linha=e["linha"], valor=e.get("valor"), legenda=e.get("legenda", "")) for e in leitura.escala],
        faixas=[FaixaPrevia(linha=f["linha"], minimo=f.get("minimo"), maximo=f.get("maximo"), percentual=f.get("percentual"),
                            notas_zero=f.get("notas_zero")) for f in leitura.faixas],
        grupos=[GrupoPrevia(nome=nome, itens=itens) for nome, itens in grupos.items()],
        erros=leitura.erros,
        avisos=leitura.avisos,
        pode_importar=not leitura.erros,
    )


def importar_no_contrato(sessao: Session, tipo: str, contrato_id, conteudo: bytes, nome_arquivo: str, autor: Usuario) -> None:
    """Cria uma versão inativa do checklist ou formulário no contrato (as regras de edição são as do cadastro manual)."""
    leitura = ler_checklist(conteudo) if tipo == "checklist" else ler_formulario(conteudo)
    if leitura.erros:
        raise ErroPlanilha(leitura.erros)
    if tipo == "checklist":
        configuracao.criar_checklist(sessao, contrato_id, leitura.checklist, autor)
    else:
        configuracao.criar_formulario(sessao, contrato_id, leitura.formulario, autor)
    auditar(sessao, autor.login, f"contrato.{tipo}.importar_xlsx", f"Contrato {contrato_id}", autor_id=autor.id,
            alvo_tipo="contrato", alvo_id=contrato_id, dados={"arquivo": nome_arquivo, "nome": leitura.nome})
    sessao.commit()


def importar_como_modelo(sessao: Session, tipo: str, conteudo: bytes, nome_arquivo: str, autor: Usuario):
    """Cria um modelo global (checklist ou formulário) a partir da planilha."""
    leitura = ler_checklist(conteudo) if tipo == "checklist" else ler_formulario(conteudo)
    if leitura.erros:
        raise ErroPlanilha(leitura.erros)
    if tipo == "checklist":
        dados = GravacaoModelo(tipo=tipo, nome=leitura.nome, itens=leitura.checklist.itens)
    else:
        dados = GravacaoModelo(tipo=tipo, nome=leitura.nome, definicao=leitura.formulario.definicao)
    modelo = configuracao.salvar_modelo(sessao, dados, autor)
    auditar(sessao, autor.login, "contrato.modelo.importar_xlsx", f"Modelo {leitura.nome}", autor_id=autor.id,
            alvo_tipo="modelo", alvo_id=modelo.id, dados={"arquivo": nome_arquivo, "tipo": tipo})
    sessao.commit()
    return modelo


# ---------------------------------------------------------------------------------------------
# Geração do modelo da planilha
# ---------------------------------------------------------------------------------------------

CABECALHO = Font(bold=True, color="FFFFFF")
FUNDO = PatternFill("solid", fgColor="1F4E79")


def _titulos(folha, linha: int, titulos: list[str]) -> None:
    """Linha de títulos com a cor da identidade e larguras confortáveis."""
    for coluna, titulo in enumerate(titulos, start=1):
        celula = folha.cell(linha, coluna, titulo)
        celula.font, celula.fill = CABECALHO, FUNDO
        folha.column_dimensions[celula.column_letter].width = 34 if coluna <= 2 else 22


def _instrucoes(folha, linhas: list[str]) -> None:
    """Instruções no topo da aba (a leitura não depende do texto delas)."""
    folha["A1"] = linhas[0]
    folha["A1"].font = Font(bold=True, size=13)
    for n, linha in enumerate(linhas[1:], start=2):
        folha.cell(n, 1, linha).alignment = Alignment(wrap_text=False)


def _salvar(livro: Workbook) -> bytes:
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def gerar_modelo_checklist() -> bytes:
    """Planilha em branco do checklist, com um exemplo preenchido."""
    livro = Workbook()
    folha = livro.active
    folha.title = "Checklist"
    _instrucoes(folha, ["Modelo de importação de checklist", "Preencha o nome e os documentos mensais, na ordem. Obrigatório, Com validade e Vale para outros contratos: Sim ou Não (em branco: Sim, Não e Não). Vale para outros contratos (só com validade) marca o documento da empresa, reaproveitável entre contratos da mesma empresa."])
    folha["A4"], folha["B4"] = "Nome do checklist", "Checklist mensal padrão"
    _titulos(folha, 6, ["Documento", "Observação", "Obrigatório", "Com validade", "Vale para outros contratos"])
    exemplos = [("Certidão negativa de débitos", "Federal, estadual e municipal", "Sim", "Sim", "Sim"),
                ("Comprovante de pagamento de encargos", "", "Sim", "Não", "Não"),
                ("Relatório fotográfico", "Quando houver", "Não", "Não", "Não")]
    for n, linha in enumerate(exemplos, start=7):
        for coluna, valor in enumerate(linha, start=1):
            folha.cell(n, coluna, valor)
    return _salvar(livro)


def gerar_modelo_formulario() -> bytes:
    """Planilha em branco do formulário de avaliação, com um exemplo preenchido."""
    livro = Workbook()
    folha = livro.active
    folha.title = "Formulário"
    _instrucoes(folha, ["Modelo de importação de formulário de avaliação da qualidade",
                        "Preencha o nome aqui e as abas Escala, Faixas e Itens. Os pesos de cada grupo devem somar 100."])
    folha["A4"], folha["B4"] = "Nome do formulário", "Avaliação mensal da qualidade"
    folha.column_dimensions["A"].width, folha.column_dimensions["B"].width = 26, 40
    escala = livro.create_sheet("Escala")
    _instrucoes(escala, ["Escala de notas (ordem crescente, sem repetir)"])
    _titulos(escala, 3, ["Nota", "Legenda"])
    for n, linha in enumerate([(0, "Insatisfatório"), (5, "Regular"), (10, "Ótimo")], start=4):
        escala.cell(n, 1, linha[0]), escala.cell(n, 2, linha[1])
    faixas = livro.create_sheet("Faixas")
    _instrucoes(faixas, ["Faixas da nota final e % do pagamento liberado (Máximo em branco = sem teto; Notas zero é opcional)"])
    _titulos(faixas, 3, ["Mínimo", "Máximo", "Percentual liberado", "Notas zero"])
    for n, linha in enumerate([(0, 4.99, 70, None), (5, 7.99, 90, None), (8, None, 100, None)], start=4):
        for coluna, valor in enumerate(linha, start=1):
            faixas.cell(n, coluna, valor)
    itens = livro.create_sheet("Itens")
    _instrucoes(itens, ["Itens avaliados por grupo (grupo em branco repete o de cima; pesos do grupo somam 100)"])
    _titulos(itens, 3, ["Grupo", "Item", "Descrição", "Peso"])
    exemplos = [("Pessoal", "Pontualidade", "Cumprimento dos horários", 50), (None, "Uniforme", "Uso de uniforme e crachá", 50),
                ("Materiais", "Qualidade", "Materiais conforme especificação", 100)]
    for n, linha in enumerate(exemplos, start=4):
        for coluna, valor in enumerate(linha, start=1):
            itens.cell(n, coluna, valor)
    return _salvar(livro)
