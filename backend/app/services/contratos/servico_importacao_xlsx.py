# Criado por José Eduardo Santana Martins
# Este arquivo serve para ler a planilha XLSX de cadastro de contrato, montar a prévia e gravar a importação.
"""Importação de contrato por planilha XLSX ("Checklist de Alimentação do Sistema de Contratos").

Formato esperado (o modelo fica em `app/recursos/modelo-importacao-contrato.xlsx`):
- rótulos nas colunas A e B e valores na coluna C (mesclada até J), um campo por linha;
- nos blocos agrupados ("Empresa", "Empresa - Preposto", "Processo Gestão/Execução"), o rótulo do
  grupo fica em A (célula mesclada) e o do campo em B;
- as linhas da equipe (Gestor, Fiscais...) são ignoradas: a equipe é cadastrada depois, editando o contrato;
- a tabela de itens começa na linha cujo rótulo em A é "Descrição"; cada coluna é achada pelo
  título, e só entram as linhas com descrição preenchida.

Os campos são achados pelo rótulo (sem acento, em minúsculas), e não pela posição, para tolerar
linhas a mais. Cada problema vira um erro com a linha da planilha, para o usuário corrigir.

Fluxo: `previa()` só lê e valida; `importar()` repete a leitura e, sem erros, grava empresa (se nova),
preposto (se novo) e contrato numa transação só.
"""

import re
import unicodedata
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.erros import _mensagem_validacao
from app.models.contratos import Contrato, EmpresaContratada, PrepostoEmpresa
from app.models.usuario import Usuario
from app.schemas.contratos.contratos import GravacaoContrato, GravacaoItem
from app.schemas.contratos.empresas import GravacaoEmpresa, GravacaoPreposto
from app.schemas.contratos.importacao import (
    ContratoPrevia,
    EmpresaPrevia,
    ErroImportacao,
    ItemPrevia,
    PrepostoPrevia,
    PreviaImportacao,
)
from app.schemas.contratos.validadores import normalizar_cnpj, normalizar_cpf, somente_digitos
from app.services.contratos import calculos
from app.services.contratos.erros import ErroRegraContrato
from app.services.contratos.servico_contratos import criar_contrato
from app.services.contratos.servico_empresas import criar_empresa, salvar_preposto
from app.services.servico_auditoria import auditar

# Limite do arquivo enviado (a planilha modelo tem ~12 KB)
TAMANHO_MAXIMO = 5 * 1024 * 1024

# Rótulo normalizado na planilha → (chave interna, rótulo exibido nas mensagens).
# Nos blocos agrupados a chave do rótulo é "grupo|campo" (grupo na coluna A, campo na coluna B).
CAMPOS = {
    "nro do contrato": ("numero", "Nro do Contrato"),
    "numero do contrato": ("numero", "Número do contrato"),
    "empresa|cnpj": ("cnpj", "CNPJ"),
    "empresa|razao social": ("razao_social", "Razão Social"),
    "empresa|nome fantasia": ("nome_fantasia", "Nome Fantasia"),
    "empresa|endereco": ("endereco", "Endereço"),
    "empresa - preposto|nome": ("preposto_nome", "Preposto · Nome"),
    "empresa - preposto|cpf": ("preposto_cpf", "Preposto · CPF"),
    "empresa - preposto|e-mail": ("preposto_email", "Preposto · E-mail"),
    "empresa - preposto|email": ("preposto_email", "Preposto · E-mail"),
    "empresa - preposto|telefone": ("preposto_telefone", "Preposto · Telefone"),
    "apelido": ("apelido", "Apelido"),
    "data inicial": ("data_inicio", "Data Inicial"),
    "vigencia": ("vigencia_inicial_meses", "Vigência"),
    "vigencia inicial": ("vigencia_inicial_meses", "Vigência"),
    "vigencia maxima": ("vigencia_maxima_meses", "Vigência Máxima"),
    "periodicidade de execucao": ("periodicidade_meses", "Periodicidade de Execução"),
    "periodicidade": ("periodicidade_meses", "Periodicidade de Execução"),
    "mes de reajuste": ("mes_reajuste", "Mês de Reajuste"),
    "objeto": ("objeto", "Objeto"),
    "processo gestao|numero": ("sei_gestao_numero", "Processo Gestão · Número"),
    "processo gestao|link": ("sei_gestao_link", "Processo Gestão · Link"),
    "processo execucao|numero": ("sei_execucao_numero", "Processo Execução · Número"),
    "processo execucao|link": ("sei_execucao_link", "Processo Execução · Link"),
}
# Rótulos da equipe: lidos só para avisar que não são importados
ROTULOS_EQUIPE = ("gestor", "fiscal")

# Título normalizado da coluna na tabela de itens → (chave do item, rótulo exibido)
COLUNAS_ITENS = {
    "descricao": ("descricao", "Descrição"),
    "tipo": ("tipo", "Tipo"),
    "faturamento": ("calcula_pro_rata", "Faturamento"),
    "classe": ("codigo_classe", "Classe"),
    "nd": ("codigo_natureza_despesa", "ND"),
    "siafisico (bec)": ("codigo_siafisico", "SIAFISICO (BEC)"),
    "siafisico": ("codigo_siafisico", "SIAFISICO (BEC)"),
    "catmat/catser": ("codigo_catmat_catser", "CATMAT/CATSER"),
    "qtd mensal": ("quantidade_mensal", "QTD MENSAL"),
    "qtd na vigencia": ("quantidade_total", "QTD NA VIGÊNCIA"),
    "valor unitario": ("valor_unitario", "VALOR UNITÁRIO"),
}

PERIODICIDADES = {"mensal": 1, "bimestral": 2, "trimestral": 3, "semestral": 6, "anual": 12}
MESES = ["janeiro", "fevereiro", "marco", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]


# ---------------------------------------------------------------------------------------------
# Estrutura da leitura
# ---------------------------------------------------------------------------------------------

@dataclass
class Leitura:
    """Resultado da leitura da planilha, antes de consultar o banco."""

    # Valores convertidos por chave interna (ex.: "numero", "cnpj", "data_inicio")
    valores: dict[str, Any] = field(default_factory=dict)
    # Chave interna → (linha, rótulo), para apontar erros no lugar certo
    posicoes: dict[str, tuple[int, str]] = field(default_factory=dict)
    itens: list[dict[str, Any]] = field(default_factory=list)
    erros: list[ErroImportacao] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)

    def erro(self, chave_ou_rotulo: str, mensagem: str, linha: int | None = None) -> None:
        """Registra um erro; se a chave for de um campo lido, usa a linha e o rótulo dele."""
        if chave_ou_rotulo in self.posicoes:
            linha, rotulo = self.posicoes[chave_ou_rotulo]
        else:
            rotulo = chave_ou_rotulo
        # Um erro por campo basta (evita repetir o mesmo problema vindo da conversão e da validação)
        if not any(e.campo == rotulo and e.linha == linha for e in self.erros):
            self.erros.append(ErroImportacao(linha=linha, campo=rotulo, mensagem=mensagem))


# ---------------------------------------------------------------------------------------------
# Conversões de células
# ---------------------------------------------------------------------------------------------

def normalizar(texto: Any) -> str:
    """Rótulo comparável: sem acento, minúsculas, espaços simples e sem ":" no fim."""
    if texto is None:
        return ""
    sem_acento = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", sem_acento).strip().rstrip(":").strip().lower()


def texto(valor: Any) -> str:
    """Conteúdo da célula como texto; números inteiros perdem o ".0" que o Excel acrescenta."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip()


def converter_data(valor: Any) -> date:
    """Célula de data do Excel ou texto dd/mm/aaaa (também aceita aaaa-mm-dd)."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    bruto = texto(valor)
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(bruto, formato).date()
        except ValueError:
            continue
    raise ValueError("data inválida (use dd/mm/aaaa)")


def converter_meses(valor: Any) -> int:
    """Quantidade de meses: número ou texto com número (ex.: "12 meses")."""
    if isinstance(valor, (int, float)) and float(valor).is_integer():
        return int(valor)
    achado = re.search(r"\d+", texto(valor))
    if not achado:
        raise ValueError("informe a quantidade de meses (ex.: 12)")
    return int(achado.group())


def converter_periodicidade(valor: Any) -> int:
    """Mensal, Bimestral, Trimestral, Semestral, Anual ou o número de meses."""
    nome = normalizar(valor)
    if nome in PERIODICIDADES:
        return PERIODICIDADES[nome]
    numero = converter_meses(valor)
    if numero not in PERIODICIDADES.values():
        raise ValueError("use Mensal, Bimestral, Trimestral, Semestral ou Anual")
    return numero


def converter_mes(valor: Any) -> int:
    """Nome do mês (Janeiro...) ou número de 1 a 12."""
    nome = normalizar(valor)
    for indice, mes in enumerate(MESES, start=1):
        # Aceita o nome inteiro ou abreviado ("jan", "janeiro")
        if nome and (nome == mes or (len(nome) >= 3 and mes.startswith(nome))):
            return indice
    numero = converter_meses(valor)
    if not 1 <= numero <= 12:
        raise ValueError("informe um mês de 1 a 12 ou o nome do mês")
    return numero


def converter_tipo(valor: Any) -> str:
    """Contínuo ou Sob demanda → código da API."""
    nome = normalizar(valor)
    if nome.startswith("contin"):
        return "continuo"
    if nome.replace("-", " ").startswith("sob demanda"):
        return "sob_demanda"
    raise ValueError("use Contínuo ou Sob demanda")


def converter_faturamento(valor: Any) -> bool:
    """Pró-rata (proporcional em mês parcial) ou Sempre Integral → `calcula_pro_rata`."""
    nome = normalizar(valor).replace("-", " ").replace(" ", "")
    if nome in ("prorata", "comprorata"):
        return True
    if nome in ("sempreintegral", "integral"):
        return False
    raise ValueError("use Pró-rata ou Sempre Integral")


def converter_decimal(valor: Any) -> Decimal:
    """Número do Excel ou texto no formato brasileiro ("1.234,56", "R$ 10,50")."""
    if isinstance(valor, bool):
        raise ValueError("deve ser um número")
    if isinstance(valor, (int, float)):
        # repr evita o ruído binário (0.1 vira "0.1" e não 0.1000000000000000055...)
        return Decimal(repr(valor)) if isinstance(valor, float) else Decimal(valor)
    bruto = texto(valor).replace("R$", "").replace(" ", "").replace(" ", "")
    # Com vírgula, o formato é brasileiro: pontos são separadores de milhar
    if "," in bruto:
        bruto = bruto.replace(".", "").replace(",", ".")
    try:
        return Decimal(bruto)
    except InvalidOperation as erro:
        raise ValueError("deve ser um número (ex.: 1.234,56)") from erro


def converter_numero_contrato(valor: Any) -> str:
    """"12/2026" ou "012/2026" → "012/2026" (o formato validado pela API)."""
    bruto = texto(valor).replace(" ", "")
    achado = re.fullmatch(r"(\d{1,4})/(\d{4})", bruto)
    if not achado:
        raise ValueError("use o formato NNN/AAAA (ex.: 012/2026)")
    return f"{int(achado.group(1)):03d}/{achado.group(2)}"


def documento(valor: Any, digitos: int) -> str:
    """CNPJ/CPF como texto; se veio como número, recupera os zeros à esquerda perdidos pelo Excel."""
    if isinstance(valor, (int, float)):
        return str(int(valor)).zfill(digitos)
    return texto(valor)


# ---------------------------------------------------------------------------------------------
# Leitura da planilha
# ---------------------------------------------------------------------------------------------

def abrir_planilha(conteudo: bytes):
    """Abre o arquivo e devolve a primeira aba; arquivo que não é XLSX vira erro de regra (400)."""
    if not conteudo:
        raise ErroRegraContrato("O arquivo está vazio.")
    if len(conteudo) > TAMANHO_MAXIMO:
        raise ErroRegraContrato("O arquivo passa de 5 MB. Envie a planilha de cadastro do contrato.")
    try:
        # data_only: lê o valor calculado das fórmulas, e não o texto da fórmula
        livro = load_workbook(BytesIO(conteudo), data_only=True, read_only=False)
    except (BadZipFile, KeyError, OSError, ValueError) as erro:
        raise ErroRegraContrato("Não foi possível ler o arquivo. Envie uma planilha no formato .xlsx.") from erro
    return livro.worksheets[0]


# Conversor de cada campo do cabeçalho (os demais são texto simples)
CONVERSORES = {
    "numero": converter_numero_contrato,
    "data_inicio": converter_data,
    "vigencia_inicial_meses": converter_meses,
    "vigencia_maxima_meses": converter_meses,
    "periodicidade_meses": converter_periodicidade,
    "mes_reajuste": converter_mes,
    "cnpj": lambda v: documento(v, 14),
    "preposto_cpf": lambda v: documento(v, 11),
}


def ler_planilha(conteudo: bytes) -> Leitura:
    """Lê o cabeçalho (campo a campo) e a tabela de itens, convertendo os valores."""
    folha = abrir_planilha(conteudo)
    leitura = Leitura()
    grupo = ""
    linha_itens: int | None = None
    equipe_preenchida = False

    for linha in range(1, folha.max_row + 1):
        rotulo_a = normalizar(folha.cell(linha, 1).value)
        rotulo_b = normalizar(folha.cell(linha, 2).value)
        # A tabela de itens começa na linha com "Descrição" na coluna A
        if rotulo_a == "descricao":
            linha_itens = linha
            break
        # O rótulo em A abre um grupo (células mescladas deixam A vazio nas linhas seguintes do grupo)
        if rotulo_a:
            grupo = rotulo_a
        chave_rotulo = f"{grupo}|{rotulo_b}" if rotulo_b else rotulo_a
        valor = folha.cell(linha, 3).value
        if chave_rotulo in CAMPOS:
            chave, rotulo = CAMPOS[chave_rotulo]
            leitura.posicoes[chave] = (linha, rotulo)
            if valor is None or texto(valor) == "":
                leitura.valores[chave] = None
                continue
            try:
                conversor = CONVERSORES.get(chave, texto)
                leitura.valores[chave] = conversor(valor)
            except ValueError as erro:
                leitura.valores[chave] = None
                leitura.erro(chave, str(erro))
        elif rotulo_a.startswith(ROTULOS_EQUIPE) and texto(valor):
            equipe_preenchida = True

    if equipe_preenchida:
        leitura.avisos.append("A equipe de gestão e fiscalização não é importada: cadastre-a editando o contrato depois da importação.")
    if linha_itens is None:
        leitura.erro("Itens", "Não foi encontrada a tabela de itens (linha com \"Descrição\" na coluna A).")
    else:
        _ler_itens(folha, linha_itens, leitura)
    return leitura


def _ler_itens(folha, linha_titulos: int, leitura: Leitura) -> None:
    """Lê as linhas de itens abaixo dos títulos; linhas sem descrição (ex.: a legenda do modelo) são puladas."""
    # Coluna de cada campo, pelo título (a ordem das colunas pode variar)
    colunas: dict[str, tuple[int, str]] = {}
    for coluna in range(1, folha.max_column + 1):
        titulo = normalizar(folha.cell(linha_titulos, coluna).value)
        if titulo in COLUNAS_ITENS:
            chave, rotulo = COLUNAS_ITENS[titulo]
            colunas.setdefault(chave, (coluna, rotulo))
    faltando = [r for c, r in {v[0]: v[1] for v in COLUNAS_ITENS.values()}.items() if c not in colunas]
    if faltando:
        leitura.erro("Itens", "Colunas não encontradas na tabela de itens: " + ", ".join(faltando) + ".", linha_titulos)
        return

    conversores = {
        "tipo": converter_tipo,
        "calcula_pro_rata": converter_faturamento,
        "quantidade_mensal": converter_decimal,
        "quantidade_total": converter_decimal,
        "valor_unitario": converter_decimal,
    }
    for linha in range(linha_titulos + 1, folha.max_row + 1):
        descricao = texto(folha.cell(linha, colunas["descricao"][0]).value)
        if not descricao:
            continue
        numero_item = len(leitura.itens) + 1
        item: dict[str, Any] = {"linha": linha, "descricao": descricao}
        for chave, (coluna, rotulo) in colunas.items():
            if chave == "descricao":
                continue
            valor = folha.cell(linha, coluna).value
            if valor is None or texto(valor) == "":
                item[chave] = None
                continue
            try:
                item[chave] = conversores.get(chave, texto)(valor)
            except ValueError as erro:
                item[chave] = None
                leitura.erro(f"Item {numero_item} · {rotulo}", str(erro), linha)
        leitura.itens.append(item)
    if not leitura.itens:
        leitura.avisos.append("Nenhum item financeiro foi encontrado na planilha; o contrato será cadastrado sem itens.")


# ---------------------------------------------------------------------------------------------
# Validação e prévia
# ---------------------------------------------------------------------------------------------

@dataclass
class Analise:
    """Leitura validada: os objetos prontos para gravar e a prévia para a tela."""

    previa: PreviaImportacao
    contrato: GravacaoContrato | None
    empresa_existente: EmpresaContratada | None
    empresa_nova: GravacaoEmpresa | None
    preposto_novo: GravacaoPreposto | None


def _vazio(valor: Any) -> Any:
    """Valor da leitura para os schemas: None vira texto vazio (os schemas cuidam do obrigatório)."""
    return "" if valor is None else valor


def _item_para_gravacao(item: dict[str, Any]) -> dict[str, Any]:
    """Item lido no formato de `GravacaoItem` (quantidades vazias viram zero)."""
    return {
        "id": None,
        "descricao": item["descricao"],
        "tipo": item.get("tipo") or "",
        "calcula_pro_rata": item.get("calcula_pro_rata") if item.get("calcula_pro_rata") is not None else True,
        "codigo_classe": _vazio(item.get("codigo_classe")),
        "codigo_natureza_despesa": _vazio(item.get("codigo_natureza_despesa")),
        "codigo_siafisico": _vazio(item.get("codigo_siafisico")),
        "codigo_catmat_catser": _vazio(item.get("codigo_catmat_catser")),
        "quantidade_mensal": item.get("quantidade_mensal") if item.get("quantidade_mensal") is not None else Decimal(0),
        # Contínuo: a quantidade da vigência é calculada (mensal × meses); só o sob demanda usa este teto
        "quantidade_total": (item.get("quantidade_total") or Decimal(0)) if item.get("tipo") == "sob_demanda" else Decimal(0),
        "valor_unitario": item.get("valor_unitario") if item.get("valor_unitario") is not None else "",
    }


def _registrar_erros_validacao(leitura: Leitura, erro: ValidationError) -> None:
    """Converte os erros do Pydantic em erros com linha e rótulo da planilha."""
    for detalhe in erro.errors():
        local = detalhe.get("loc", ())
        mensagem = _mensagem_validacao(detalhe)
        if local and local[0] == "itens" and len(local) >= 2 and isinstance(local[1], int):
            item = leitura.itens[local[1]]
            # Campo do item (ou o item todo, quando a regra envolve mais de um campo)
            rotulo = next((r for c, r in COLUNAS_ITENS.values() if len(local) > 2 and c == local[2]), None)
            leitura.erro(f"Item {local[1] + 1}" + (f" · {rotulo}" if rotulo else ""), mensagem, item["linha"])
        elif local and local[0] in leitura.posicoes:
            leitura.erro(str(local[0]), mensagem)
        elif local and local[0] in ("vigencia_inicial_meses", "vigencia_maxima_meses") or not local:
            # Regra entre campos (ex.: vigência máxima menor que a inicial) aponta para a vigência máxima
            leitura.erro("vigencia_maxima_meses" if "vigencia_maxima_meses" in leitura.posicoes else "Contrato", mensagem)
        else:
            leitura.erro(str(local[0]), mensagem)


def analisar(sessao: Session, conteudo: bytes) -> Analise:
    """Lê a planilha, valida tudo (inclusive contra o banco) e monta a prévia."""
    leitura = ler_planilha(conteudo)
    v = leitura.valores

    # --- Empresa: reaproveitada pelo CNPJ ou cadastrada ---
    empresa_existente: EmpresaContratada | None = None
    empresa_nova: GravacaoEmpresa | None = None
    empresa_previa: EmpresaPrevia | None = None
    cnpj = ""
    try:
        cnpj = normalizar_cnpj(_vazio(v.get("cnpj")))
    except ValueError as erro:
        leitura.erro("cnpj", "campo obrigatório" if not v.get("cnpj") else str(erro))
    if cnpj:
        empresa_existente = sessao.scalar(select(EmpresaContratada).where(EmpresaContratada.cnpj == cnpj))
    if empresa_existente is not None:
        empresa_previa = EmpresaPrevia(
            existente=True, id=empresa_existente.id, cnpj=empresa_existente.cnpj, razao_social=empresa_existente.razao_social,
            nome_fantasia=empresa_existente.nome_fantasia, endereco=empresa_existente.endereco,
        )
        if not empresa_existente.ativa:
            leitura.erro("cnpj", "A empresa com este CNPJ está inativa. Reative-a no cadastro de empresas antes de importar.")
        # Dados diferentes na planilha não alteram o cadastro: só avisam
        for chave, atual in (("razao_social", empresa_existente.razao_social), ("nome_fantasia", empresa_existente.nome_fantasia),
                             ("endereco", empresa_existente.endereco)):
            informado = _vazio(v.get(chave))
            if informado and normalizar(informado) != normalizar(atual):
                rotulo = leitura.posicoes[chave][1]
                leitura.avisos.append(
                    f"A empresa já está cadastrada e será mantida como está: {rotulo} cadastrado \"{atual or '—'}\", na planilha \"{informado}\"."
                )
    elif cnpj or v.get("razao_social"):
        dados_empresa = {"cnpj": cnpj or _vazio(v.get("cnpj")), "razao_social": _vazio(v.get("razao_social")),
                         "nome_fantasia": _vazio(v.get("nome_fantasia")), "endereco": _vazio(v.get("endereco")), "ativa": True}
        try:
            empresa_nova = GravacaoEmpresa.model_validate(dados_empresa)
        except ValidationError as erro:
            _registrar_erros_validacao(leitura, erro)
        empresa_previa = EmpresaPrevia(existente=False, cnpj=cnpj, razao_social=dados_empresa["razao_social"],
                                       nome_fantasia=dados_empresa["nome_fantasia"], endereco=dados_empresa["endereco"])

    # --- Preposto (opcional): reaproveitado pelo CPF dentro da empresa ou cadastrado ---
    preposto_novo: GravacaoPreposto | None = None
    preposto_previa: PrepostoPrevia | None = None
    campos_preposto = ("preposto_nome", "preposto_cpf", "preposto_email", "preposto_telefone")
    if any(v.get(c) for c in campos_preposto):
        dados_preposto = {"cpf": _vazio(v.get("preposto_cpf")), "nome": _vazio(v.get("preposto_nome")),
                          "email": _vazio(v.get("preposto_email")), "telefone": _vazio(v.get("preposto_telefone")), "cargo": "", "ativo": True}
        try:
            validado = GravacaoPreposto.model_validate(dados_preposto)
        except ValidationError as erro:
            validado = None
            for detalhe in erro.errors():
                # Os campos do schema do preposto têm nomes próprios; aponta para a linha do bloco do preposto
                campo = detalhe.get("loc", ("",))[0]
                leitura.erro(f"preposto_{'email' if campo == 'email' else campo}", _mensagem_validacao(detalhe))
        existente = None
        if validado and empresa_existente is not None:
            existente = sessao.scalar(select(PrepostoEmpresa).where(
                PrepostoEmpresa.empresa_id == empresa_existente.id, PrepostoEmpresa.cpf == validado.cpf))
        if existente is not None:
            preposto_previa = PrepostoPrevia(existente=True, cpf=existente.cpf, nome=existente.nome, email=existente.email, telefone=existente.telefone)
            if normalizar(existente.nome) != normalizar(validado.nome):
                leitura.avisos.append(f"O preposto de CPF informado já está cadastrado como \"{existente.nome}\" e será mantido como está.")
        else:
            preposto_novo = validado
            preposto_previa = PrepostoPrevia(existente=False, cpf=somente_digitos(dados_preposto["cpf"]), nome=dados_preposto["nome"],
                                             email=dados_preposto["email"], telefone=dados_preposto["telefone"])

    # --- Contrato e itens (validados pelo mesmo schema do cadastro pela tela) ---
    dados_contrato = {
        "numero": v.get("numero"),
        # Id provisório: a empresa real só existe na gravação
        "empresa_id": empresa_existente.id if empresa_existente else uuid.uuid4(),
        "apelido": _vazio(v.get("apelido")),
        "objeto": _vazio(v.get("objeto")),
        "data_inicio": v.get("data_inicio"),
        "vigencia_inicial_meses": v.get("vigencia_inicial_meses"),
        "vigencia_maxima_meses": v.get("vigencia_maxima_meses"),
        "periodicidade_meses": v.get("periodicidade_meses"),
        "mes_reajuste": v.get("mes_reajuste"),
        "sei_gestao_numero": _vazio(v.get("sei_gestao_numero")),
        "sei_gestao_link": _vazio(v.get("sei_gestao_link")),
        "sei_execucao_numero": _vazio(v.get("sei_execucao_numero")),
        "sei_execucao_link": _vazio(v.get("sei_execucao_link")),
        "situacao_forcada": None,
        "equipe": {},
        "itens": [_item_para_gravacao(i) for i in leitura.itens],
        "versao": None,
    }
    # Campos vazios saem do dicionário: o schema responde "campo obrigatório" em vez de erro de tipo
    dados_contrato = {chave: valor for chave, valor in dados_contrato.items() if valor is not None or chave in ("situacao_forcada", "versao")}
    contrato: GravacaoContrato | None = None
    try:
        contrato = GravacaoContrato.model_validate(dados_contrato)
    except ValidationError as erro:
        _registrar_erros_validacao(leitura, erro)

    # Regra entre campos conferida aqui também: o validador do schema só roda quando os demais campos são válidos
    inicial, maxima = v.get("vigencia_inicial_meses"), v.get("vigencia_maxima_meses")
    if inicial and maxima and maxima < inicial:
        leitura.erro("vigencia_maxima_meses", "a vigência máxima deve ser maior ou igual à vigência inicial")

    # Número já usado por outro contrato
    if v.get("numero"):
        sequencial, ano = (int(p) for p in v["numero"].split("/"))
        if sessao.scalar(select(Contrato.id).where(Contrato.sequencial == sequencial, Contrato.ano == ano)) is not None:
            leitura.erro("numero", f"Já existe um contrato com o número {v['numero']}.")

    previa = PreviaImportacao(
        contrato=_contrato_previa(v),
        empresa=empresa_previa,
        preposto=preposto_previa,
        itens=[ItemPrevia(**{k: val for k, val in i.items() if k in ItemPrevia.model_fields and val is not None} | {
            "quantidade_total": i.get("quantidade_total") if i.get("tipo") == "sob_demanda" else None}) for i in leitura.itens],
        valor_global_estimado=_valor_estimado(leitura),
        # Na ordem da planilha, para o usuário corrigir de cima para baixo
        erros=sorted(leitura.erros, key=lambda e: e.linha or 0),
        avisos=leitura.avisos,
        pode_importar=not leitura.erros,
    )
    return Analise(previa, contrato if not leitura.erros else None, empresa_existente, empresa_nova, preposto_novo)


def _contrato_previa(v: dict[str, Any]) -> ContratoPrevia:
    """Dados do contrato para a prévia, com a data final calculada quando possível."""
    fim = None
    if v.get("data_inicio") and v.get("vigencia_inicial_meses"):
        fim = calculos.calcular_data_fim(v["data_inicio"], v["vigencia_inicial_meses"])
    campos = {c: v.get(c) for c in ContratoPrevia.model_fields if c != "data_fim" and v.get(c) is not None}
    return ContratoPrevia(**campos, data_fim=fim)


def _valor_estimado(leitura: Leitura) -> Decimal | None:
    """Valor global estimado da vigência inicial (mesma conta simples da tela de cadastro)."""
    meses = leitura.valores.get("vigencia_inicial_meses")
    if not meses or not leitura.itens:
        return None
    total = Decimal(0)
    for item in leitura.itens:
        preco = item.get("valor_unitario") or Decimal(0)
        if item.get("tipo") == "sob_demanda":
            total += (item.get("quantidade_total") or Decimal(0)) * preco
        else:
            total += (item.get("quantidade_mensal") or Decimal(0)) * preco * meses
    return calculos.arredondar(total)


# ---------------------------------------------------------------------------------------------
# Operações usadas pelas rotas
# ---------------------------------------------------------------------------------------------

def previa(sessao: Session, conteudo: bytes) -> PreviaImportacao:
    """Lê e valida a planilha sem gravar nada."""
    return analisar(sessao, conteudo).previa


class ErroPlanilha(Exception):
    """A planilha tem erros: a importação não grava nada. `erros` vai no corpo da resposta 400."""

    def __init__(self, erros: list[ErroImportacao]) -> None:
        super().__init__(f"A planilha tem {len(erros)} erro(s). Corrija-os e envie de novo.")
        self.erros = erros


def importar(sessao: Session, conteudo: bytes, nome_arquivo: str, autor: Usuario) -> Contrato:
    """Grava empresa (se nova), preposto (se novo) e contrato numa transação só."""
    analise = analisar(sessao, conteudo)
    if not analise.previa.pode_importar or analise.contrato is None:
        raise ErroPlanilha(analise.previa.erros)
    try:
        empresa = analise.empresa_existente or criar_empresa(sessao, analise.empresa_nova, autor, commit=False)
        if analise.preposto_novo is not None:
            salvar_preposto(sessao, empresa.id, analise.preposto_novo, autor, commit=False)
        dados = analise.contrato.model_copy(update={"empresa_id": empresa.id})
        contrato = criar_contrato(sessao, dados, autor, commit=False)
        auditar(sessao, autor.login, "contrato.importar_xlsx", f"Contrato {contrato.numero}", autor_id=autor.id,
                alvo_tipo="contrato", alvo_id=contrato.id,
                dados={"arquivo": nome_arquivo, "empresa_nova": analise.empresa_existente is None,
                       "preposto_novo": analise.preposto_novo is not None, "itens": len(dados.itens)})
        sessao.commit()
    except Exception:
        # Qualquer falha desfaz tudo: nem empresa, nem preposto, nem contrato ficam gravados
        sessao.rollback()
        raise
    return contrato
