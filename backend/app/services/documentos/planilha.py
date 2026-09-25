# Criado por José Eduardo Santana Martins
# Este arquivo serve para gerar planilhas XLSX com formatação brasileira.
"""Geração de planilhas XLSX com cabeçalho institucional e formatos brasileiros.

Quem usa descreve as abas com `Aba` e `Coluna` (título, formato, largura) e passa as linhas;
este módulo cuida do estilo, do cabeçalho congelado e da conversão dos valores.
"""

from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Formatos numéricos do Excel: moeda em reais, quantidade com 4 casas, percentual e data
FORMATO_MOEDA = '"R$" #,##0.00'
FORMATO_QUANTIDADE = "#,##0.0000"
FORMATO_PERCENTUAL = "0.00%"
FORMATO_DATA = "dd/mm/yyyy"

# Cores da identidade visual (sem o "#", como o openpyxl espera)
_VERMELHO = "C82331"
_VERMELHO_SUAVE = "FCECEE"


@dataclass
class Coluna:
    """Uma coluna da aba: título, formato numérico e largura (em caracteres)."""
    titulo: str
    # Formato numérico do Excel (use as constantes FORMATO_*); vazio = texto/geral
    formato: str = ""
    largura: float = 16


@dataclass
class Aba:
    """Uma aba da planilha: colunas, linhas de dados, título, rodapé de totais e observações."""
    nome: str
    colunas: Sequence[Coluna]
    linhas: Iterable[Sequence[Any]]
    titulo: str = ""
    rodape: Sequence[Any] | None = None
    observacoes: list[str] = field(default_factory=list)


def _valor(valor: Any) -> Any:
    """Converte o valor para um tipo que o openpyxl grava corretamente."""
    # O openpyxl grava Decimal como número; datas continuam datas para permitir filtros no Excel
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, date):
        return valor
    return valor


def gerar_planilha(abas: Sequence[Aba]) -> bytes:
    """Planilha com uma aba por item de `abas`: título, cabeçalho congelado, formatos e rodapé."""
    livro = Workbook()
    # O Workbook nasce com uma aba vazia; ela é removida para usar só as abas pedidas
    livro.remove(livro.active)
    borda = Border(bottom=Side(style="thin", color=_VERMELHO))
    for aba in abas:
        # O Excel limita o nome da aba a 31 caracteres
        folha = livro.create_sheet(aba.nome[:31])
        linha = 1
        # Título na linha 1; nesse caso, o cabeçalho das colunas começa na linha 3
        if aba.titulo:
            folha.cell(row=1, column=1, value=aba.titulo).font = Font(bold=True, size=13, color=_VERMELHO)
            linha = 3
        # Cabeçalho das colunas: negrito, fundo vermelho claro, filete vermelho embaixo
        for indice, coluna in enumerate(aba.colunas, start=1):
            celula = folha.cell(row=linha, column=indice, value=coluna.titulo)
            celula.font = Font(bold=True)
            celula.fill = PatternFill("solid", fgColor=_VERMELHO_SUAVE)
            celula.border = borda
            celula.alignment = Alignment(vertical="center", wrap_text=True)
            folha.column_dimensions[get_column_letter(indice)].width = coluna.largura
        # Congela o cabeçalho: ele continua visível ao rolar a planilha
        folha.freeze_panes = folha.cell(row=linha + 1, column=1)
        # Linhas de dados, aplicando o formato de cada coluna
        for valores in aba.linhas:
            linha += 1
            for indice, (coluna, valor) in enumerate(zip(aba.colunas, valores, strict=False), start=1):
                celula = folha.cell(row=linha, column=indice, value=_valor(valor))
                if coluna.formato:
                    celula.number_format = coluna.formato
        # Rodapé (ex.: linha de totais) em negrito
        if aba.rodape:
            linha += 1
            for indice, (coluna, valor) in enumerate(zip(aba.colunas, aba.rodape, strict=False), start=1):
                celula = folha.cell(row=linha, column=indice, value=_valor(valor))
                celula.font = Font(bold=True)
                if coluna.formato:
                    celula.number_format = coluna.formato
        # Observações em itálico abaixo da tabela, com uma linha em branco entre elas
        for observacao in aba.observacoes:
            linha += 2
            folha.cell(row=linha, column=1, value=observacao).font = Font(italic=True, color="586372")
    # Salva em memória e devolve os bytes (nada é gravado em disco aqui)
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()
