# Criado por José Eduardo Santana Martins
# Este arquivo serve para a carga em lote dos dados funcionais do RH por planilha XLSX (CGP).
"""Carga em lote dos dados funcionais (CGP): modelo, prévia e importação.

- **Modelo:** planilha já preenchida com todos os servidores ativos (login, nome, setor) e os valores atuais. A CGP
  completa as colunas e devolve o arquivo.
- **Prévia:** lê e confere tudo sem gravar; devolve, por linha, o que vai mudar e os erros.
- **Importação:** confere de novo e, sem nenhum erro, grava tudo numa transação só (auditoria
  `rh.cadastro.funcionais` por servidor e `rh.cadastro.funcionais_lote` do lote).

Regras:
- a linha é identificada pelo **login**; "Nome" e "Setor" são só informativos;
- **célula vazia mantém o valor atual** (nada é apagado pela planilha);
- "Dias disponíveis no período vigente" ajusta os dias creditados do período para que sobrem exatamente esses dias,
  descontando o que já está agendado no sistema. Sem isso, o período nasce com o crédito cheio do parâmetro (30 dias);
- as validações são as mesmas da tela (`GravacaoDadosFuncionais` e `servico_cadastro.salvar_funcionais`).
"""

import re
from dataclasses import dataclass, field
from datetime import time
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.configuracao import obter_configuracao
from app.models.rh import DadosFuncionais
from app.models.usuario import Usuario
from app.schemas.rh import GravacaoDadosFuncionais
from app.services.contratos.servico_importacao_xlsx import normalizar, texto
from app.services.rh import servico_periodos
from app.services.rh.papeis import exigir_cgp, setor_do_usuario
from app.services.rh.servico_cadastro import ErroCadastro, ajustar_periodo_vigente, salvar_funcionais
from app.services.servico_auditoria import auditar

TAMANHO_MAXIMO = 5 * 1024 * 1024

# (chave, título da coluna, largura, dica)
COLUNAS = [
    ("login", "Login", 20, "Obrigatório. Identifica o servidor (não altere)."),
    ("nome", "Nome", 34, "Só informativo."),
    ("setor", "Setor", 36, "Só informativo."),
    ("autorizador", "Autorizador (login)", 20, "Login de quem aprova as férias e a LP."),
    ("substituto", "Substituto (login)", 20, "Login de quem aprova no lugar deste servidor, quando ele (como autorizador) estiver afastado."),
    ("topo", "Topo da hierarquia", 12, "Sim ou Não."),
    ("inicio", "Início do período aquisitivo (dd/mm)", 16, "Ex.: 15/03."),
    ("disponiveis", "Dias disponíveis no período vigente", 14, "Dias que ainda restam no período atual (sem contar o que já está agendado no sistema)."),
    ("exercicio", "Exercício da LP", 10, "Ano do saldo de licença-prêmio. Ex.: 2026."),
    ("lp", "Dias de LP", 10, "Saldo de licença-prêmio no exercício."),
    ("jornada", "Jornada (horas/semana)", 12, "Ex.: 40."),
    ("plantao", "Regime de plantão", 11, "Sim ou Não."),
    ("horario", "Horário de trabalho", 18, "Ex.: 9:00 às 18:00."),
    ("estudante", "Horário de estudante", 11, "Sim ou Não."),
    ("intervalo", "Intervalo de almoço e descanso", 18, "Ex.: 12:00 às 13:00."),
    ("rg_cin", "RG/CIN nº", 16, "Ex.: 12.345.678-9."),
    ("rs_pv", "RS/PV nº", 16, "Ex.: 1.234.567/8."),
]
TITULO_PARA_CHAVE = {normalizar(titulo): chave for chave, titulo, _, _ in COLUNAS}
FAIXA = re.compile(r"^\s*(\d{1,2})(?::|h)?(\d{2})?\s*(?:às|as|a|-|–)\s*(\d{1,2})(?::|h)?(\d{2})?\s*$", re.IGNORECASE)


class ErroPlanilhaRh(ErroCadastro):
    """Arquivo ilegível ou com erros (400)."""


@dataclass
class LinhaImportacao:
    linha: int
    login: str
    nome: str = ""
    usuario_id: int | None = None
    mudancas: list[str] = field(default_factory=list)
    erros: list[str] = field(default_factory=list)
    dados: dict | None = None
    disponiveis: int | None = None


def _hora_texto(valor: time | None) -> str:
    return f"{valor.hour}:{valor.minute:02d}" if valor else ""


def _faixa_texto(inicio: time | None, fim: time | None) -> str:
    return f"{_hora_texto(inicio)} às {_hora_texto(fim)}" if inicio and fim else ""


def _sim_nao(valor: bool | None) -> str:
    return "Sim" if valor else "Não"


# ---------------------------------------------------------------------------------------------
# Modelo
# ---------------------------------------------------------------------------------------------

def modelo(sessao: Session, autor: Usuario) -> bytes:
    """Planilha com todos os servidores ativos e os valores atuais (menos a conta root)."""
    exigir_cgp(sessao, autor)
    login_admin = obter_configuracao().login_admin.strip().lower()
    usuarios = sorted((u for u in sessao.scalars(select(Usuario).where(Usuario.ativo.is_(True))) if u.login.strip().lower() != login_admin),
                      key=lambda u: (u.nome_completo or u.login).lower())
    logins = {u.id: u.login for u in sessao.scalars(select(Usuario))}
    livro = Workbook()
    aba = livro.active
    aba.title = "Dados funcionais"
    cabecalho = PatternFill("solid", fgColor="B0222E")
    informativo = PatternFill("solid", fgColor="EEF0F2")
    for coluna, (_, titulo, largura, _) in enumerate(COLUNAS, start=1):
        celula = aba.cell(row=1, column=coluna, value=titulo)
        celula.font, celula.fill = Font(bold=True, color="FFFFFF"), cabecalho
        celula.alignment = Alignment(wrap_text=True, vertical="center")
        aba.column_dimensions[celula.column_letter].width = largura
    aba.row_dimensions[1].height = 42
    aba.freeze_panes = "B2"
    for linha, u in enumerate(usuarios, start=2):
        f = sessao.get(DadosFuncionais, u.id)
        periodo = servico_periodos.vigente(sessao, u.id, _hoje(), criar=False) if f else None
        disponivel = None
        if periodo is not None:
            disponivel = servico_periodos.situacao(sessao, periodo, _hoje()).disponivel
        valores = [
            u.login, u.nome_completo or u.login, setor_do_usuario(u),
            logins.get(f.autorizador_id, "") if f else "", logins.get(f.substituto_id, "") if f else "",
            _sim_nao(f.sem_superior) if f else "", servico_periodos.texto_inicio(f) or "" if f else "", disponivel,
            f.exercicio if f else None, f.saldo_lp_dias if f else None, f.jornada_semanal_horas if f else None,
            _sim_nao(f.regime_plantao) if f else "", _faixa_texto(f.horario_trabalho_inicio, f.horario_trabalho_fim) if f else "",
            _sim_nao(f.horario_estudante) if f else "", _faixa_texto(f.intervalo_inicio, f.intervalo_fim) if f else "",
            (f.rg_cin or "") if f else "", (f.rs_pv or "") if f else "",
        ]
        for coluna, valor in enumerate(valores, start=1):
            celula = aba.cell(row=linha, column=coluna, value=valor)
            # Início (dd/mm) e documentos como texto, para o Excel não converter em data ou número
            if COLUNAS[coluna - 1][0] in ("inicio", "rg_cin", "rs_pv", "horario", "intervalo"):
                celula.number_format = "@"
            if coluna <= 3:
                celula.fill = informativo
    sim_nao = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
    aba.add_data_validation(sim_nao)
    for chave in ("topo", "plantao", "estudante"):
        letra = aba.cell(row=1, column=[c[0] for c in COLUNAS].index(chave) + 1).column_letter
        sim_nao.add(f"{letra}2:{letra}{max(len(usuarios) + 1, 2)}")
    instrucoes = livro.create_sheet("Instruções")
    instrucoes.column_dimensions["A"].width = 38
    instrucoes.column_dimensions["B"].width = 100
    instrucoes.append(["Carga dos dados funcionais do RH (SGI SPI)"])
    instrucoes["A1"].font = Font(bold=True, size=13)
    for linha in (
        ["Como usar", "Preencha a aba \"Dados funcionais\" e envie em RH › Validações › Importar planilha. Confira a prévia antes de gravar."],
        ["Célula vazia", "Mantém o valor atual: a planilha nunca apaga dados."],
        ["Linhas", "Uma por servidor, identificado pelo login. Pode apagar as linhas de quem não quer alterar."],
        [],
    ):
        instrucoes.append(linha)
    instrucoes.append(["Coluna", "Como preencher"])
    instrucoes["A6"].font = instrucoes["B6"].font = Font(bold=True)
    for _, titulo, _, dica in COLUNAS:
        instrucoes.append([titulo, dica])
    saida = BytesIO()
    livro.save(saida)
    sessao.rollback()
    return saida.getvalue()


def _hoje():
    from app.services.rh.servico_afastamentos import hoje

    return hoje()


# ---------------------------------------------------------------------------------------------
# Leitura e conferência
# ---------------------------------------------------------------------------------------------

def _abrir(conteudo: bytes):
    if not conteudo:
        raise ErroPlanilhaRh("O arquivo está vazio.")
    if len(conteudo) > TAMANHO_MAXIMO:
        raise ErroPlanilhaRh("O arquivo passa de 5 MB.")
    try:
        return load_workbook(BytesIO(conteudo), data_only=True).worksheets[0]
    except (BadZipFile, KeyError, OSError, ValueError) as erro:
        raise ErroPlanilhaRh("Não foi possível ler o arquivo. Envie a planilha no formato .xlsx (baixe o modelo).") from erro


def _sim_nao_celula(valor: Any, rotulo: str, erros: list[str]) -> bool | None:
    t = normalizar(texto(valor))
    if not t:
        return None
    if t in ("sim", "s", "x", "1", "true", "verdadeiro"):
        return True
    if t in ("nao", "n", "0", "false", "falso"):
        return False
    erros.append(f"{rotulo}: use Sim ou Não.")
    return None


def _inteiro(valor: Any, rotulo: str, erros: list[str]) -> int | None:
    t = texto(valor).strip()
    if not t:
        return None
    try:
        return int(float(t.replace(",", ".")))
    except ValueError:
        erros.append(f"{rotulo}: número inválido ({t}).")
        return None


def _faixa(valor: Any, rotulo: str, erros: list[str]) -> tuple[time, time] | None:
    t = texto(valor).strip()
    if not t:
        return None
    m = FAIXA.match(t)
    if not m:
        erros.append(f"{rotulo}: use o formato 9:00 às 18:00 ({t}).")
        return None
    try:
        return time(int(m[1]), int(m[2] or 0)), time(int(m[3]), int(m[4] or 0))
    except ValueError:
        erros.append(f"{rotulo}: horário inválido ({t}).")
        return None


def _inicio(valor: Any) -> str:
    """dd/mm como texto (se o Excel converteu em data, volta a dd/mm)."""
    if hasattr(valor, "day") and hasattr(valor, "month"):
        return f"{valor.day:02d}/{valor.month:02d}"
    return texto(valor).strip()


def conferir(sessao: Session, conteudo: bytes) -> list[LinhaImportacao]:
    """Lê a planilha e monta, por linha, os dados finais (valores atuais + células preenchidas) e os erros."""
    aba = _abrir(conteudo)
    linhas = aba.iter_rows(values_only=True)
    titulos = next(linhas, None) or ()
    indices = {TITULO_PARA_CHAVE[normalizar(t)]: i for i, t in enumerate(titulos) if normalizar(t) in TITULO_PARA_CHAVE}
    if "login" not in indices:
        raise ErroPlanilhaRh("A planilha não tem a coluna \"Login\". Baixe o modelo e use-o como base.")
    usuarios = {u.login.strip().lower(): u for u in sessao.scalars(select(Usuario))}
    resultado: list[LinhaImportacao] = []
    vistos: set[str] = set()
    for numero, valores in enumerate(linhas, start=2):
        celula = lambda chave: valores[indices[chave]] if chave in indices and indices[chave] < len(valores) else None  # noqa: E731
        login = texto(celula("login")).strip()
        if not login and not any(v not in (None, "") for v in valores):
            continue
        item = LinhaImportacao(linha=numero, login=login)
        resultado.append(item)
        usuario = usuarios.get(login.lower())
        if usuario is None:
            item.erros.append(f"Login \"{login}\" não encontrado." if login else "Login em branco.")
            continue
        if login.lower() in vistos:
            item.erros.append("Login repetido na planilha.")
            continue
        vistos.add(login.lower())
        item.usuario_id, item.nome = usuario.id, usuario.nome_completo or usuario.login
        atual = sessao.get(DadosFuncionais, usuario.id)
        base = {
            "autorizador_id": atual.autorizador_id if atual else None, "substituto_id": atual.substituto_id if atual else None,
            "sem_superior": atual.sem_superior if atual else False, "inicio_periodo_aquisitivo": servico_periodos.texto_inicio(atual) if atual else None,
            "exercicio": atual.exercicio if atual else None, "saldo_lp_dias": atual.saldo_lp_dias if atual else 0,
            "jornada_semanal_horas": atual.jornada_semanal_horas if atual else None, "regime_plantao": atual.regime_plantao if atual else False,
            "horario_trabalho_inicio": atual.horario_trabalho_inicio if atual else None, "horario_trabalho_fim": atual.horario_trabalho_fim if atual else None,
            "horario_estudante": atual.horario_estudante if atual else False, "intervalo_inicio": atual.intervalo_inicio if atual else None,
            "intervalo_fim": atual.intervalo_fim if atual else None, "rg_cin": atual.rg_cin if atual else None, "rs_pv": atual.rs_pv if atual else None,
        }
        novo = dict(base)
        erros = item.erros
        for chave, campo in (("autorizador", "autorizador_id"), ("substituto", "substituto_id")):
            valor = texto(celula(chave)).strip()
            if valor:
                outro = usuarios.get(valor.lower())
                if outro is None or not outro.ativo:
                    situacao = "não encontrado" if outro is None else "inativo"
                    erros.append(f"{'Autorizador' if chave == 'autorizador' else 'Substituto'}: login \"{valor}\" {situacao}.")
                else:
                    novo[campo] = outro.id
        for chave, campo, rotulo in (("topo", "sem_superior", "Topo da hierarquia"), ("plantao", "regime_plantao", "Regime de plantão"),
                                     ("estudante", "horario_estudante", "Horário de estudante")):
            valor = _sim_nao_celula(celula(chave), rotulo, erros)
            if valor is not None:
                novo[campo] = valor
        if _inicio(celula("inicio")):
            novo["inicio_periodo_aquisitivo"] = _inicio(celula("inicio"))
        for chave, campo, rotulo in (("exercicio", "exercicio", "Exercício da LP"), ("lp", "saldo_lp_dias", "Dias de LP"),
                                     ("jornada", "jornada_semanal_horas", "Jornada")):
            valor = _inteiro(celula(chave), rotulo, erros)
            if valor is not None:
                novo[campo] = valor
        for chave, (campo_i, campo_f), rotulo in (("horario", ("horario_trabalho_inicio", "horario_trabalho_fim"), "Horário de trabalho"),
                                                 ("intervalo", ("intervalo_inicio", "intervalo_fim"), "Intervalo")):
            faixa = _faixa(celula(chave), rotulo, erros)
            if faixa:
                novo[campo_i], novo[campo_f] = faixa
        for chave in ("rg_cin", "rs_pv"):
            valor = texto(celula(chave)).strip()
            if valor:
                novo[chave] = valor
        item.disponiveis = _inteiro(celula("disponiveis"), "Dias disponíveis", erros)
        if item.disponiveis is not None and not 0 <= item.disponiveis <= 365:
            erros.append("Dias disponíveis: de 0 a 365.")
        if item.disponiveis is not None and not novo["inicio_periodo_aquisitivo"]:
            erros.append("Dias disponíveis exigem o início do período aquisitivo.")
        # Mesmas regras da tela
        try:
            dados = GravacaoDadosFuncionais(**novo).model_dump()
        except ValidationError as erro:
            erros.extend(f"{'/'.join(str(p) for p in e['loc']) or 'Dados'}: {e['msg'].removeprefix('Value error, ')}" for e in erro.errors())
            continue
        if dados["inicio_periodo_aquisitivo"]:
            try:
                servico_periodos.ler_inicio(dados["inicio_periodo_aquisitivo"])
            except ValueError as erro:
                erros.append(f"Início do período aquisitivo: {erro}")
        for campo in ("autorizador_id", "substituto_id"):
            if dados[campo] == usuario.id:
                erros.append("O servidor não pode ser o próprio autorizador nem o próprio substituto.")
        if erros:
            continue
        item.dados = dados
        comparavel = {**base, "inicio_periodo_aquisitivo": base["inicio_periodo_aquisitivo"]}
        item.mudancas = [c for c in dados if dados[c] != comparavel.get(c)]
        # Dias disponíveis: só é mudança se diferirem do saldo atual (ou se o início do período mudou)
        if item.disponiveis is not None:
            vigente = servico_periodos.vigente(sessao, usuario.id, _hoje(), criar=False) if atual else None
            disponivel_atual = servico_periodos.situacao(sessao, vigente, _hoje()).disponivel if vigente else None
            if "inicio_periodo_aquisitivo" in item.mudancas or item.disponiveis != disponivel_atual:
                item.mudancas.append("dias disponíveis do período vigente")
            else:
                item.disponiveis = None
    return resultado


def importar(sessao: Session, conteudo: bytes, autor: Usuario) -> list[LinhaImportacao]:
    """Confere de novo e, sem nenhum erro, grava tudo numa transação."""
    exigir_cgp(sessao, autor)
    linhas = conferir(sessao, conteudo)
    com_erro = [l for l in linhas if l.erros]
    if com_erro:
        raise ErroPlanilhaRh(f"A planilha tem {len(com_erro)} linha(s) com erro. Corrija e envie de novo (nada foi gravado).")
    try:
        for item in linhas:
            if not item.mudancas:
                continue
            salvar_funcionais(sessao, item.usuario_id, item.dados, autor, commit=False)
            if item.disponiveis is not None:
                periodo = servico_periodos.vigente(sessao, item.usuario_id, _hoje())
                agendado = servico_periodos.usado(sessao, item.usuario_id, periodo.inicio, periodo.fim)
                ajustar_periodo_vigente(sessao, item.usuario_id, item.disponiveis + agendado, autor, commit=False)
        auditar(sessao, autor.login, "rh.cadastro.funcionais_lote", "Planilha de dados funcionais",
                f"{sum(1 for l in linhas if l.mudancas)} servidor(es) alterado(s)", autor_id=autor.id)
        sessao.commit()
    except ErroCadastro:
        sessao.rollback()
        raise
    return linhas
