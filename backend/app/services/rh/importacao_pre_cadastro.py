# Criado por José Eduardo Santana Martins
# Este arquivo serve para o formulário XLSX da ficha de pré-cadastro (uma pessoa por arquivo): gerar o modelo, conferir a prévia e importar a ficha preenchida.
"""Ficha de pré-cadastro em planilha (RH, controle total).

- **Modelo:** um formulário de **uma pessoa**, em uma aba só ("Ficha"): instruções no topo, seções (identificação, documentos, contato, endereço, vínculo, formação e experiências),
  rótulos bloqueados e só as células de resposta liberadas. Onde há opções fixas a célula tem lista; se o texto não segue o padrão (CPF, data, e-mail, UF...) o Excel avisa e explica
  como preencher. A planilha vai protegida (sem senha) e o RH a envia à pessoa, que preenche e devolve.
- **Leitura:** cada célula de resposta tem um **nome definido** (`f_cpf`, `fo1_curso`, `ex2_empresa`...): o arquivo é lido por eles, não pela posição, então ajustes de largura ou
  de altura não quebram a importação.
- **Prévia:** confere tudo sem gravar (as mesmas regras da tela: CPF, PIS, CEP, UF, CPF já cadastrado, datas).
- **Importação:** cria a ficha como **rascunho**; o RH completa o setor e o cargo pretendidos, o login proposto e o início previsto, que não são da pessoa.
"""

from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.models.usuario import Usuario
from app.schemas.rh_pre_cadastro import Experiencia, Formacao, GravacaoFicha
from app.services.contratos.servico_importacao_xlsx import normalizar, texto
from app.services.rh import servico_pre_cadastro
from app.services.rh.servico_pre_cadastro import ErroFicha, so_digitos
from app.services.servico_auditoria import auditar

TAMANHO_MAXIMO = 5 * 1024 * 1024
ABA = "Ficha"
LINHAS_FORMACAO = 5
LINHAS_EXPERIENCIA = 5
VERMELHO, CINZA, CINZA_BORDA, AMARELO = "B0222E", "F1F2F4", "BFC5CC", "FFF8E1"

UFS = "AC,AL,AP,AM,BA,CE,DF,ES,GO,MA,MT,MS,MG,PA,PB,PR,PE,PI,RJ,RN,RS,RO,RR,SC,SP,SE,TO"
LISTAS = {
    "uf": UFS,
    "sexo": "Feminino,Masculino,Outro,Prefere não informar",
    "estado_civil": "Solteiro(a),Casado(a),União estável,Separado(a),Divorciado(a),Viúvo(a)",
    "vinculo": "Servidor efetivo cedido,Comissionado sem vínculo,Empregado público,Outro",
    "nivel": "Ensino médio,Técnico,Graduação,Pós-graduação,Mestrado,Doutorado",
    "sim_nao": "Sim,Não",
}


@dataclass(frozen=True)
class Campo:
    chave: str
    rotulo: str
    tipo: str = "texto"  # texto, data, data_nasc, data_passada, lista, cpf, pis, cep, email, telefone, ano
    maximo: int = 100
    lista: str = ""
    obrigatorio: bool = False
    largo: bool = False
    dica: str = ""


# Seções do formulário: o que a pessoa preenche (setor/cargo pretendidos, login proposto e início previsto são do RH e ficam fora)
SECOES: list[tuple[str, list[Campo]]] = [
    ("1. Identificação", [
        Campo("nome_completo", "Nome completo", maximo=200, obrigatorio=True, largo=True, dica="Nome como consta no documento de identidade."),
        Campo("nome_social", "Nome social", maximo=200),
        Campo("data_nascimento", "Data de nascimento", "data_nasc", dica="Formato dd/mm/aaaa. Ex.: 15/03/1990."),
        Campo("sexo", "Sexo", "lista", lista="sexo"),
        Campo("estado_civil", "Estado civil", "lista", lista="estado_civil"),
        Campo("nacionalidade", "Nacionalidade", maximo=60, dica="Ex.: Brasileira."),
        Campo("naturalidade_cidade", "Naturalidade (cidade)", maximo=100),
        Campo("naturalidade_uf", "Naturalidade (UF)", "lista", lista="uf"),
        Campo("nome_mae", "Nome da mãe", maximo=200, largo=True),
        Campo("nome_pai", "Nome do pai", maximo=200, largo=True),
    ]),
    ("2. Documentos", [
        Campo("cpf", "CPF", "cpf", obrigatorio=True, dica="11 números, com ou sem pontos e traço. Ex.: 529.982.247-25."),
        Campo("rg_numero", "RG (número)", maximo=30),
        Campo("rg_orgao", "RG (órgão emissor)", maximo=30, dica="Ex.: SSP."),
        Campo("rg_uf", "RG (UF)", "lista", lista="uf"),
        Campo("rg_emissao", "RG (data de emissão)", "data_passada", dica="Formato dd/mm/aaaa."),
        Campo("cin", "CIN (nova identidade)", maximo=30, dica="Carteira de Identidade Nacional, se já tiver."),
        Campo("pis_pasep", "PIS/PASEP", "pis", dica="11 números, com ou sem pontuação."),
        Campo("titulo_numero", "Título de eleitor", maximo=20),
        Campo("titulo_zona", "Zona", maximo=5),
        Campo("titulo_secao", "Seção", maximo=5),
        Campo("ctps_numero", "CTPS (número)", maximo=20),
        Campo("ctps_serie", "CTPS (série)", maximo=10),
        Campo("ctps_uf", "CTPS (UF)", "lista", lista="uf"),
        Campo("cnh", "CNH", maximo=20),
    ]),
    ("3. Contato", [
        Campo("email_pessoal", "E-mail pessoal", "email", maximo=200, largo=True, dica="O e-mail em que você recebe mensagens. Ex.: nome@exemplo.com."),
        Campo("telefone", "Telefone", "telefone", dica="Com DDD. Ex.: (11) 3333-4444."),
        Campo("celular", "Celular", "telefone", dica="Com DDD. Ex.: (11) 99999-8888."),
        Campo("emergencia_nome", "Contato de emergência (nome)", maximo=200, largo=True),
        Campo("emergencia_parentesco", "Parentesco", maximo=60, dica="Ex.: Mãe, cônjuge, irmão."),
        Campo("emergencia_telefone", "Telefone de emergência", "telefone"),
    ]),
    ("4. Endereço", [
        Campo("cep", "CEP", "cep", dica="8 números, com ou sem traço. Ex.: 01310-100."),
        Campo("logradouro", "Logradouro", maximo=200, largo=True),
        Campo("numero", "Número", maximo=20),
        Campo("complemento", "Complemento", maximo=100),
        Campo("bairro", "Bairro", maximo=100),
        Campo("cidade", "Cidade", maximo=100),
        Campo("uf", "UF", "lista", lista="uf"),
    ]),
    ("5. Vínculo com o serviço público", [
        Campo("vinculo_tipo", "Tipo de vínculo", "lista", lista="vinculo"),
        Campo("orgao_origem", "Órgão de origem", maximo=200),
        Campo("cargo_origem", "Cargo de origem", maximo=200),
        Campo("rs_pv", "RS/PV (matrícula)", maximo=30, dica="Se for servidor, o número da matrícula."),
    ]),
]
DATAS = {"data_nascimento", "rg_emissao"}
FORMACAO = [
    Campo("nivel", "Nível", "lista", lista="nivel"), Campo("curso", "Curso", maximo=200), Campo("instituicao", "Instituição", maximo=200),
    Campo("ano_conclusao", "Ano de conclusão", "ano", dica="Ano com 4 números. Ex.: 2015. Deixe vazio se estiver em andamento."), Campo("em_andamento", "Em andamento", "lista", lista="sim_nao"),
]
EXPERIENCIA = [
    Campo("empresa", "Órgão ou empresa", maximo=200), Campo("cargo", "Cargo", maximo=200), Campo("inicio", "Início", "data_passada", dica="Formato dd/mm/aaaa."),
    Campo("fim", "Fim", "data_passada", dica="Formato dd/mm/aaaa. Deixe vazio se ainda trabalha lá."), Campo("descricao", "Descrição das atividades", maximo=4000),
]

INSTRUCOES = (
    "COMO PREENCHER\n"
    "• Preencha apenas as células BRANCAS. As cinzas são o formulário e estão bloqueadas.\n"
    "• Onde a célula tiver uma seta ▼ ao ser selecionada, escolha uma das opções da lista.\n"
    "• Datas: dd/mm/aaaa (ex.: 15/03/1990). CPF, PIS e CEP: só números, ou com a pontuação de costume.\n"
    "• Se você digitar algo fora do padrão, o Excel avisa e explica como corrigir. Ao selecionar uma célula, uma dica mostra o que escrever.\n"
    "• Os campos com * são obrigatórios. Em Formação e Experiências, preencha quantas linhas precisar; deixe em branco as que sobrarem.\n"
    "• Ao terminar, salve o arquivo (.xlsx) e devolva ao RH. Não renomeie a aba nem desproteja a planilha."
)


class ErroPlanilhaFichas(ErroFicha):
    """Arquivo ilegível ou com erros (400)."""


@dataclass
class LinhaFicha:
    linha: int
    nome: str = ""
    cpf: str = ""
    login: str = ""
    cargo_rotulo: str = ""
    erros: list[str] = field(default_factory=list)
    dados: GravacaoFicha | None = None
    ficha_id: int | None = None


# --- Modelo ------------------------------------------------------------------------------------------------

def _borda(cor: str = CINZA_BORDA) -> Border:
    lado = Side(style="thin", color=cor)
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def _validacao(campo: Campo, celula: str) -> DataValidation | None:
    """Regra da célula de resposta, com a mensagem de erro que ensina o preenchimento e a dica que aparece ao selecionar."""
    tipo = campo.tipo
    sem = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({celula},".",""),"-","")," ",""),"/","")'
    if tipo == "lista":
        v = DataValidation(type="list", formula1=f'"{LISTAS[campo.lista]}"', allow_blank=True)
        v.error, v.errorTitle = "Escolha uma das opções da lista (clique na seta ▼ ao lado da célula).", "Opção inválida"
    elif tipo in ("cpf", "pis", "cep"):
        tamanho = {"cpf": 11, "pis": 11, "cep": 8}[tipo]
        nome = {"cpf": "CPF", "pis": "PIS/PASEP", "cep": "CEP"}[tipo]
        v = DataValidation(type="custom", formula1=f"=OR(LEN({celula})=0,AND(ISNUMBER(--{sem}),LEN({sem})={tamanho}))", allow_blank=True)
        v.error, v.errorTitle = f"{nome} deve ter {tamanho} números (pode usar a pontuação de costume). Ex.: " + {"cpf": "529.982.247-25", "pis": "123.45678.91-9", "cep": "01310-100"}[tipo] + ".", f"{nome} inválido"
    elif tipo == "email":
        v = DataValidation(type="custom", formula1=f'=OR(LEN({celula})=0,AND(ISNUMBER(FIND("@",{celula})),ISNUMBER(FIND(".",{celula},FIND("@",{celula}))),ISERROR(FIND(" ",{celula}))))', allow_blank=True)
        v.error, v.errorTitle = "Digite um e-mail completo, sem espaços. Ex.: nome@exemplo.com.", "E-mail inválido"
    elif tipo == "telefone":
        v = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="20", allow_blank=True)
        v.error, v.errorTitle = "Digite o telefone com DDD, em até 20 caracteres. Ex.: (11) 99999-8888.", "Telefone inválido"
    elif tipo in ("data", "data_nasc", "data_passada"):
        v = DataValidation(type="date", operator="between", formula1="DATE(1900,1,1)", formula2="TODAY()", allow_blank=True)
        v.error, v.errorTitle = "Digite uma data válida no formato dd/mm/aaaa (ex.: 15/03/1990), que não seja futura.", "Data inválida"
    elif tipo == "ano":
        v = DataValidation(type="whole", operator="between", formula1="1950", formula2="2100", allow_blank=True)
        v.error, v.errorTitle = "Digite o ano com 4 números, entre 1950 e 2100. Ex.: 2015.", "Ano inválido"
    else:
        v = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=str(campo.maximo), allow_blank=True)
        v.error, v.errorTitle = f"Texto longo demais: use no máximo {campo.maximo} caracteres.", "Texto longo demais"
    v.errorStyle, v.showErrorMessage, v.showInputMessage = "stop", True, True
    v.promptTitle = campo.rotulo[:32]
    v.prompt = (campo.dica or {"lista": "Escolha uma opção da lista.", "texto": f"Digite até {campo.maximo} caracteres."}.get(tipo, "Preencha esta célula."))[:255]
    return v


class _Formulario:
    """Montagem do formulário: células de rótulo (bloqueadas), de resposta (liberadas, com regra e nome definido) e títulos."""

    def __init__(self) -> None:
        self.livro = Workbook()
        self.aba = self.livro.active
        self.aba.title = ABA
        self.linha = 1
        self.fonte = Font(name="Calibri", size=10)

    def _estilo(self, celula, *, fundo: str | None = None, negrito: bool = False, cor: str = "000000", tamanho: int = 10, borda: bool = False, quebra: bool = True, horizontal: str = "left") -> None:
        celula.font = Font(name="Calibri", size=tamanho, bold=negrito, color=cor)
        celula.alignment = Alignment(wrap_text=quebra, vertical="center", horizontal=horizontal)
        if fundo:
            celula.fill = PatternFill("solid", fgColor=fundo)
        if borda:
            celula.border = _borda()

    def faixa(self, valor: str, altura: float, **estilo) -> None:
        """Linha inteira (B:G) mesclada: título, subtítulo, instruções."""
        self.aba.merge_cells(start_row=self.linha, start_column=2, end_row=self.linha, end_column=7)
        for coluna in range(2, 8):
            self._estilo(self.aba.cell(row=self.linha, column=coluna), **estilo)
        self.aba.cell(row=self.linha, column=2, value=valor)
        self.aba.row_dimensions[self.linha].height = altura
        self.linha += 1

    def resposta(self, campo: Campo, nome_celula: str, coluna: int, ate: int | None = None) -> None:
        """Célula de resposta (liberada) na coluna indicada (mesclada até `ate`), com regra, dica e nome definido."""
        ate = ate or coluna
        if ate > coluna:
            self.aba.merge_cells(start_row=self.linha, start_column=coluna, end_row=self.linha, end_column=ate)
        for c in range(coluna, ate + 1):
            celula = self.aba.cell(row=self.linha, column=c)
            self._estilo(celula, fundo="FFFFFF", borda=True)
            celula.protection = Protection(locked=False)
            if campo.tipo in ("data", "data_nasc", "data_passada"):
                celula.number_format = "dd/mm/yyyy"
            elif campo.tipo != "ano":
                celula.number_format = "@"
        principal = self.aba.cell(row=self.linha, column=coluna)
        validacao = _validacao(campo, principal.coordinate)
        if validacao:
            self.aba.add_data_validation(validacao)
            validacao.add(principal.coordinate)
        self.livro.defined_names[nome_celula] = DefinedName(nome_celula, attr_text=f"'{ABA}'!{principal.coordinate.replace(principal.column_letter, '$' + principal.column_letter).replace(str(principal.row), '$' + str(principal.row))}")

    def rotulo(self, texto_: str, coluna: int, obrigatorio: bool = False) -> None:
        celula = self.aba.cell(row=self.linha, column=coluna, value=texto_ + (" *" if obrigatorio else ""))
        self._estilo(celula, fundo=CINZA, negrito=True, cor="B0222E" if obrigatorio else "3D454E", borda=True)

    def secao(self, titulo: str, campos: list[Campo]) -> None:
        self.faixa(titulo, 22, fundo=VERMELHO, negrito=True, cor="FFFFFF", tamanho=11)
        fila: list[Campo] = []
        for campo in campos:
            if campo.largo:
                self._linha_campos(fila)
                fila = []
                self.rotulo(campo.rotulo, 2, campo.obrigatorio)
                self.resposta(campo, f"f_{campo.chave}", 3, 7)
                self.aba.row_dimensions[self.linha].height = 24
                self.linha += 1
            else:
                fila.append(campo)
                if len(fila) == 3:
                    self._linha_campos(fila)
                    fila = []
        self._linha_campos(fila)
        self.linha += 1

    def _linha_campos(self, campos: list[Campo]) -> None:
        if not campos:
            return
        for n, campo in enumerate(campos):
            self.rotulo(campo.rotulo, 2 + 2 * n, campo.obrigatorio)
            self.resposta(campo, f"f_{campo.chave}", 3 + 2 * n)
        self.aba.row_dimensions[self.linha].height = 24
        self.linha += 1

    def tabela(self, titulo: str, subtitulo: str, campos: list[Campo], prefixo: str, quantidade: int, colunas: list[tuple[int, int]], altura: float = 24) -> None:
        """Tabela de linhas repetidas (formações, experiências): cabeçalho bloqueado e `quantidade` linhas de resposta."""
        self.faixa(titulo, 22, fundo=VERMELHO, negrito=True, cor="FFFFFF", tamanho=11)
        self.faixa(subtitulo, 20, fundo=CINZA, cor="586372", tamanho=9)
        for campo, (de, ate) in zip(campos, colunas):
            if ate > de:
                self.aba.merge_cells(start_row=self.linha, start_column=de, end_row=self.linha, end_column=ate)
            for c in range(de, ate + 1):
                self._estilo(self.aba.cell(row=self.linha, column=c), fundo=CINZA, negrito=True, cor="3D454E", borda=True)
            self.aba.cell(row=self.linha, column=de, value=campo.rotulo)
        self.aba.row_dimensions[self.linha].height = 22
        self.linha += 1
        for n in range(1, quantidade + 1):
            for campo, (de, ate) in zip(campos, colunas):
                self.resposta(campo, f"{prefixo}{n}_{campo.chave}", de, ate)
            self.aba.row_dimensions[self.linha].height = altura
            self.linha += 1
        self.linha += 1


def modelo(sessao: Session | None = None) -> bytes:
    """Formulário em branco (uma aba, protegida, só as células de resposta liberadas)."""
    f = _Formulario()
    aba = f.aba
    aba.sheet_view.showGridLines = False
    for coluna, largura in zip("ABCDEFGH", (2, 22, 26, 22, 26, 22, 26, 2)):
        aba.column_dimensions[coluna].width = largura
    f.faixa("FICHA DE PRÉ-CADASTRO", 34, fundo=VERMELHO, negrito=True, cor="FFFFFF", tamanho=16)
    f.faixa("Secretaria de Parcerias em Investimentos · Coordenadoria de Gestão de Pessoas", 20, fundo=CINZA, cor="586372", tamanho=10)
    f.faixa(INSTRUCOES, 124, fundo=AMARELO, tamanho=10)
    f.linha += 1
    for titulo, campos in SECOES:
        f.secao(titulo, campos)
    f.tabela("6. Formação", "Inclua quantas formações precisar (até 5); deixe as linhas que sobrarem em branco.", FORMACAO, "fo", LINHAS_FORMACAO, [(2, 2), (3, 3), (4, 5), (6, 6), (7, 7)])
    f.tabela("7. Experiências anteriores", "Órgãos e empresas por onde você passou (até 5); deixe as linhas que sobrarem em branco.", EXPERIENCIA, "ex", LINHAS_EXPERIENCIA, [(2, 2), (3, 3), (4, 4), (5, 5), (6, 7)], altura=48)
    # Protegida sem senha: a pessoa só edita as células liberadas; o RH pode desproteger se precisar
    aba.protection.sheet = True
    aba.protection.selectLockedCells = False
    aba.protection.formatColumns = True
    aba.protection.formatRows = True
    aba.page_setup.orientation = "portrait"
    aba.page_setup.paperSize = aba.PAPERSIZE_A4
    aba.page_setup.fitToWidth, aba.page_setup.fitToHeight = 1, 0
    aba.sheet_properties.pageSetUpPr.fitToPage = True
    aba.print_area = f"A1:H{f.linha}"
    saida = BytesIO()
    f.livro.save(saida)
    return saida.getvalue()


# --- Leitura e conferência -------------------------------------------------------------------------------------

def _abrir(conteudo: bytes):
    if not conteudo:
        raise ErroPlanilhaFichas("O arquivo está vazio.")
    if len(conteudo) > TAMANHO_MAXIMO:
        raise ErroPlanilhaFichas("O arquivo passa de 5 MB.")
    try:
        livro = load_workbook(BytesIO(conteudo), data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError) as erro:
        raise ErroPlanilhaFichas("Não foi possível ler o arquivo. Envie a ficha no formato .xlsx (baixe o modelo).") from erro
    if "f_nome_completo" not in livro.defined_names:
        raise ErroPlanilhaFichas("Esta planilha não é o formulário da ficha de pré-cadastro. Baixe o modelo e use-o como base.")
    return livro


def _celula(livro, nome: str) -> Any:
    """Valor da célula de resposta de um nome definido (None se o nome não existir)."""
    if nome not in livro.defined_names:
        return None
    for aba, coordenada in livro.defined_names[nome].destinations:
        return livro[aba][coordenada.replace("$", "")].value
    return None


def _data(valor: Any, rotulo: str, erros: list[str]) -> date | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(valor).strip(), formato).date()
        except ValueError:
            continue
    erros.append(f"{rotulo}: use a data no formato dd/mm/aaaa.")
    return None


def _sim_nao(valor: Any, rotulo: str, erros: list[str]) -> bool:
    bruto = normalizar(valor)
    if bruto in ("", "nao", "n"):
        return False
    if bruto in ("sim", "s"):
        return True
    erros.append(f"{rotulo}: use Sim ou Não.")
    return False


def _ano(valor: Any, rotulo: str, erros: list[str]) -> int | None:
    if valor in (None, ""):
        return None
    try:
        return int(float(str(valor).replace(",", ".")))
    except ValueError:
        erros.append(f"{rotulo}: use um ano com 4 números.")
        return None


def conferir(sessao: Session, conteudo: bytes, autor: Usuario) -> list[LinhaFicha]:
    """Lê o formulário (uma pessoa) e devolve uma linha com a ficha final e os erros, sem gravar nada."""
    livro = _abrir(conteudo)
    item = LinhaFicha(linha=1)
    erros = item.erros
    valores: dict[str, Any] = {}
    for _, campos in SECOES:
        for campo in campos:
            bruto = _celula(livro, f"f_{campo.chave}")
            valores[campo.chave] = _data(bruto, campo.rotulo, erros) if campo.chave in DATAS else texto(bruto)
    item.nome, item.cpf = valores["nome_completo"], so_digitos(valores["cpf"])
    if not item.nome:
        erros.append("Nome completo em branco (campo obrigatório).")
    if not item.cpf:
        erros.append("CPF em branco (campo obrigatório).")
    formacoes, experiencias = [], []
    for n in range(1, LINHAS_FORMACAO + 1):
        lido = {c.chave: _celula(livro, f"fo{n}_{c.chave}") for c in FORMACAO}
        if any(v not in (None, "") for v in lido.values()):
            formacoes.append(Formacao(nivel=texto(lido["nivel"]), curso=texto(lido["curso"]), instituicao=texto(lido["instituicao"]),
                                      ano_conclusao=_ano(lido["ano_conclusao"], f"Formação {n}, ano de conclusão", erros), em_andamento=_sim_nao(lido["em_andamento"], f"Formação {n}, em andamento", erros)))
    for n in range(1, LINHAS_EXPERIENCIA + 1):
        lido = {c.chave: _celula(livro, f"ex{n}_{c.chave}") for c in EXPERIENCIA}
        if any(v not in (None, "") for v in lido.values()):
            inicio, fim = _data(lido["inicio"], f"Experiência {n}, início", erros), _data(lido["fim"], f"Experiência {n}, fim", erros)
            if inicio and fim and fim < inicio:
                erros.append(f"Experiência {n}: o fim não pode ser anterior ao início.")
            experiencias.append(Experiencia(empresa=texto(lido["empresa"]), cargo=texto(lido["cargo"]), inicio=inicio, fim=fim, descricao=texto(lido["descricao"])))
    if erros:
        return [item]
    try:
        item.dados = GravacaoFicha(**valores, status="rascunho", formacoes=formacoes, experiencias=experiencias)
    except ValidationError as erro:
        erros.extend(f"{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in erro.errors())
        return [item]
    # Regras da tela (CPF, PIS, CEP, UF, CPF já cadastrado) sem gravar: grava dentro de um savepoint e desfaz
    try:
        with sessao.begin_nested():
            servico_pre_cadastro.criar(sessao, item.dados, autor, commit=False)
            raise _Desfazer()
    except _Desfazer:
        pass
    except ErroFicha as erro:
        erros.append(str(erro))
    return [item]


class _Desfazer(Exception):
    """Interrompe o savepoint da conferência para não gravar nada."""


def importar(sessao: Session, conteudo: bytes, autor: Usuario) -> list[LinhaFicha]:
    """Confere de novo e, sem erro, cria a ficha (rascunho)."""
    linhas = conferir(sessao, conteudo, autor)
    if any(l.erros for l in linhas):
        raise ErroPlanilhaFichas("A ficha tem erro(s). Corrija a planilha e envie de novo (nada foi gravado).")
    try:
        for item in linhas:
            item.ficha_id = servico_pre_cadastro.criar(sessao, item.dados, autor, commit=False).id
        auditar(sessao, autor.login, "rh.pre_cadastro.importar", "Ficha de pré-cadastro por planilha", f"{linhas[0].nome}", autor_id=autor.id, dados={"ficha_id": linhas[0].ficha_id})
        sessao.commit()
    except ErroFicha:
        sessao.rollback()
        raise
    return linhas
