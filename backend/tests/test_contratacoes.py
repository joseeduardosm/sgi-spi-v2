# Criado por José Eduardo Santana Martins
# Este arquivo serve para testar Contratações: acesso por papel, árvore, sanitização, lote, revisões, versões, histórico, conferência e Word/PDF.
"""Contratações (`/api/contratacoes`): ETP e TR com versionamento."""

import io
import uuid

import pytest
from docx import Document

from app.core.banco import FabricaSessao
from app.models.contratacoes import DocumentoContratacao
from app.services.contratacoes import docx_importacao
from tests.apoio_contratos import criar_contrato
from tests.conftest import cabecalho, criar_usuario

URL = "/api/contratacoes"


@pytest.fixture
def equipe(cliente, admin):
    """Ana cria (MODIFICACAO), Bruno edita e Carla revisa quando compartilhados; Dora (LEITURA) e Eva (CONTROLE_TOTAL); Fora não está na ACL."""
    ids = {login: criar_usuario(login, nome_completo=f"{login.title()} Souza", email=f"{login}@sp.gov.br") for login in ("ana", "bruno", "carla", "dora", "eva", "fora")}
    recurso = cliente.post("/api/acl/recursos", json={"nome": "Contratações", "slug": "contratacoes"}, headers=admin).json()
    for nivel, logins in (("MODIFICACAO", ("ana",)), ("LEITURA", ("bruno", "carla", "dora")), ("CONTROLE_TOTAL", ("eva",))):
        r = cliente.post("/api/acl/regras", json={"recurso_id": recurso["id"], "nivel": nivel, "usuarios_ids": [ids[x] for x in logins], "setores_ids": []}, headers=admin)
        assert r.status_code == 201, r.text
    return ids, {login: cabecalho(cliente, login) for login in ids}


def _criar(cliente, h, tipo="etp", nome="ETP Teste", processo="SPI-1/2026"):
    r = cliente.post(f"{URL}/documentos", json={"tipo": tipo, "nome": nome, "processo": processo}, headers=h)
    assert r.status_code == 201, r.text
    return r.json()


def _secao(cliente, h, doc_id, titulo="OBJETO"):
    r = cliente.post(f"{URL}/documentos/{doc_id}/secoes", json={"titulo": titulo}, headers=h)
    assert r.status_code == 201, r.text
    return r.json()["secoes"][-1]["id"]


def _item(cliente, h, doc_id, secao_id, texto, tipo="item", pai_id=None, html=None):
    r = cliente.post(f"{URL}/documentos/{doc_id}/itens", json={"secao_id": secao_id, "tipo": tipo, "pai_id": pai_id, "conteudo": texto, "conteudo_html": html}, headers=h)
    assert r.status_code == 201, r.text
    doc = r.json()
    itens = [i for s in doc["secoes"] for i in s["itens"] if i["secao_id"] == secao_id]
    return itens, doc


def _abrir(cliente, h, doc_id):
    return cliente.get(f"{URL}/documentos/{doc_id}", headers=h).json()


def test_acesso_e_compartilhamento(cliente, equipe):
    ids, h = equipe
    assert cliente.get(f"{URL}/documentos", headers=h["fora"]).status_code == 403
    assert cliente.post(f"{URL}/documentos", json={"tipo": "etp", "nome": "x"}, headers=h["dora"]).status_code == 403  # LEITURA não cria
    assert any(u["login"] == "bruno" for u in cliente.get(f"{URL}/opcoes-usuarios", params={"busca": "bru"}, headers=h["ana"]).json())
    assert cliente.get(f"{URL}/opcoes-usuarios", headers=h["dora"]).status_code == 403
    doc = _criar(cliente, h["ana"])
    assert doc["meu_papel"] == "criador"
    # Sem ser membro: não vê, e o 404 não revela que existe
    assert cliente.get(f"{URL}/documentos/{doc['id']}", headers=h["bruno"]).status_code == 404
    assert cliente.get(f"{URL}/documentos", headers=h["bruno"]).json()["itens"] == []
    # Editor e revisor
    assert cliente.put(f"{URL}/documentos/{doc['id']}/membros/{ids['bruno']}", json={"papel": "editor"}, headers=h["ana"]).status_code == 200
    assert cliente.put(f"{URL}/documentos/{doc['id']}/membros/{ids['carla']}", json={"papel": "revisor"}, headers=h["ana"]).status_code == 200
    assert cliente.put(f"{URL}/documentos/{doc['id']}/membros/{ids['dora']}", json={"papel": "editor"}, headers=h["bruno"]).status_code == 403  # editor não compartilha
    sec = _secao(cliente, h["bruno"], doc["id"])  # editor edita
    assert cliente.post(f"{URL}/documentos/{doc['id']}/secoes", json={"titulo": "X"}, headers=h["carla"]).status_code == 403  # revisor não edita
    assert [d["id"] for d in cliente.get(f"{URL}/documentos", headers=h["carla"]).json()["itens"]] == [doc["id"]]
    assert _abrir(cliente, h["carla"], doc["id"])["pode_editar"] is False
    # CONTROLE_TOTAL vê tudo; excluir só o criador/administração
    assert cliente.get(f"{URL}/documentos/{doc['id']}", headers=h["eva"]).status_code == 200
    assert cliente.delete(f"{URL}/documentos/{doc['id']}", headers=h["bruno"]).status_code == 403
    assert cliente.delete(f"{URL}/documentos/{doc['id']}/membros/{ids['bruno']}", headers=h["ana"]).status_code == 200
    assert cliente.get(f"{URL}/documentos/{doc['id']}", headers=h["bruno"]).status_code == 404
    assert sec


def test_arvore_marcadores_mover_e_duplicar(cliente, equipe):
    ids, h = equipe
    doc = _criar(cliente, h["ana"])
    sec = _secao(cliente, h["ana"], doc["id"])
    itens, _ = _item(cliente, h["ana"], doc["id"], sec, "Primeiro")
    a = itens[0]["id"]
    _item(cliente, h["ana"], doc["id"], sec, "Segundo")
    _item(cliente, h["ana"], doc["id"], sec, "Filho", tipo="subitem", pai_id=a)
    _item(cliente, h["ana"], doc["id"], sec, "Inciso um", tipo="inciso", pai_id=a)
    itens, d = _item(cliente, h["ana"], doc["id"], sec, "Alínea", tipo="alinea", pai_id=a)
    assert [(i["marcador"], i["conteudo"]) for i in itens] == [("1.1.", "Primeiro"), ("1.1.1.", "Filho"), ("I -", "Inciso um"), ("a)", "Alínea"), ("1.2.", "Segundo")]
    # Mover o primeiro (com a subárvore) para depois do segundo
    r = cliente.post(f"{URL}/documentos/{doc['id']}/itens/{a}/mover", json={"secao_id": sec, "pai_id": None, "ordem": 2}, headers=h["ana"])
    assert r.status_code == 200, r.text
    marc = [(i["marcador"], i["conteudo"]) for i in r.json()["secoes"][0]["itens"]]
    assert marc[0] == ("1.1.", "Segundo") and marc[1] == ("1.2.", "Primeiro") and ("1.2.1.", "Filho") in marc
    # Um item não pode virar filho de si mesmo
    assert cliente.post(f"{URL}/documentos/{doc['id']}/itens/{a}/mover", json={"secao_id": sec, "pai_id": a, "ordem": 1}, headers=h["ana"]).status_code == 400
    # Duplicar o item copia a subárvore; limpar-filhos apaga só os filhos
    r = cliente.post(f"{URL}/documentos/{doc['id']}/itens/{a}/duplicar", json=None, headers=h["ana"])
    assert r.status_code == 201 and len(r.json()["secoes"][0]["itens"]) == 9
    r = cliente.post(f"{URL}/documentos/{doc['id']}/itens/{a}/limpar-filhos", headers=h["ana"])
    assert len(r.json()["secoes"][0]["itens"]) == 6
    # Cópia do documento
    copia = cliente.post(f"{URL}/documentos/{doc['id']}/duplicar", headers=h["ana"])
    assert copia.status_code == 201 and copia.json()["nome"].startswith("Cópia de") and len(copia.json()["secoes"][0]["itens"]) == 6


def test_sanitizacao_do_html(cliente, equipe):
    ids, h = equipe
    doc = _criar(cliente, h["ana"])
    sec = _secao(cliente, h["ana"], doc["id"])
    sujo = '<p onclick="x()">Olá <strong>mundo</strong><script>alert(1)</script><a href="javascript:alert(1)">l</a></p><table><tr><td colspan="2" style="text-align:center">c</td></tr></table>'
    itens, _ = _item(cliente, h["ana"], doc["id"], sec, None, html=sujo)
    html = itens[0]["conteudo_html"]
    assert "script" not in html and "onclick" not in html and "javascript:" not in html
    assert "<strong>mundo</strong>" in html and "colspan" in html and "text-align" in html
    assert itens[0]["conteudo"].startswith("Olá mundo")


def test_entrada_em_lote_com_previa(cliente, equipe):
    ids, h = equipe
    doc = _criar(cliente, h["ana"])
    sec = _secao(cliente, h["ana"], doc["id"])
    texto = "# Objeto\n## Detalhe\n** inciso do detalhe\n$$ alinea\n# Prazo"
    previa = cliente.post(f"{URL}/documentos/{doc['id']}/lote/previa", json={"secao_id": sec, "texto": texto}, headers=h["ana"])
    assert previa.status_code == 200
    assert [(i["tipo"], i["profundidade"]) for i in previa.json()["itens"]][:3] == [("item", 0), ("subitem", 1), ("inciso", 2)]
    assert _abrir(cliente, h["ana"], doc["id"])["secoes"][0]["itens"] == []  # a prévia não grava
    r = cliente.post(f"{URL}/documentos/{doc['id']}/lote", json={"secao_id": sec, "texto": texto}, headers=h["ana"])
    assert r.status_code == 201 and r.json()["criados"] == 5
    assert all(i["precisa_revisao"] for i in _abrir(cliente, h["ana"], doc["id"])["secoes"][0]["itens"])
    assert cliente.post(f"{URL}/documentos/{doc['id']}/lote", json={"secao_id": sec, "texto": "sem marcador"}, headers=h["ana"]).status_code == 400


def test_revisoes_proposta_e_aplicacao(cliente, equipe):
    ids, h = equipe
    doc = _criar(cliente, h["ana"])
    sec = _secao(cliente, h["ana"], doc["id"])
    itens, _ = _item(cliente, h["ana"], doc["id"], sec, "Texto original")
    item = itens[0]["id"]
    for login, papel in (("bruno", "editor"), ("carla", "revisor")):
        cliente.put(f"{URL}/documentos/{doc['id']}/membros/{ids[login]}", json={"papel": papel}, headers=h["ana"])
    # Revisor propõe; exige comentário
    assert cliente.post(f"{URL}/documentos/{doc['id']}/itens/{item}/revisoes", json={"comentario": " "}, headers=h["carla"]).status_code == 400
    r = cliente.post(f"{URL}/documentos/{doc['id']}/itens/{item}/revisoes", json={"comentario": "Troque", "conteudo_proposto": "Texto novo"}, headers=h["carla"])
    assert r.status_code == 201, r.text
    rev = r.json()["secoes"][0]["itens"][0]["revisoes"][0]
    assert rev["autor_nome"] and rev["aplicada_em"] is None
    # Revisor não aplica; conferência bloqueia concluir enquanto a proposta está aberta
    rota = f"{URL}/documentos/{doc['id']}/itens/{item}/revisoes/{rev['id']}"
    assert cliente.post(f"{rota}/aplicar", headers=h["carla"]).status_code == 403
    assert cliente.put(f"{URL}/documentos/{doc['id']}/situacao", json={"situacao": "concluido", "confirmar": True}, headers=h["ana"]).status_code == 409
    # Editor (não só o criador) aplica
    r = cliente.post(f"{rota}/aplicar", headers=h["bruno"])
    assert r.status_code == 200, r.text
    novo = r.json()["secoes"][0]["itens"][0]
    assert novo["conteudo"] == "Texto novo" and novo["revisoes"][0]["aplicada_em"] and novo["revisoes"][0]["aplicada_por_nome"]
    assert cliente.post(f"{rota}/aplicar", headers=h["bruno"]).status_code == 409  # já aplicada
    # Comentário simples: resolver
    r = cliente.post(f"{URL}/documentos/{doc['id']}/itens/{item}/revisoes", json={"comentario": "Ok?"}, headers=h["carla"])
    rid = r.json()["secoes"][0]["itens"][0]["revisoes"][1]["id"]
    r = cliente.post(f"{URL}/documentos/{doc['id']}/itens/{item}/revisoes/{rid}/resolver", json={"resolvida": True}, headers=h["bruno"])
    assert r.json()["secoes"][0]["itens"][0]["revisoes"][1]["resolvida_em"]
    # Quem aplicou a proposta gerou versão do tipo "proposta"
    tipos = [v["tipo"] for v in cliente.get(f"{URL}/documentos/{doc['id']}/versoes", headers=h["ana"]).json()]
    assert "proposta" in tipos


def test_versoes_alteracoes_e_restauracao(cliente, equipe):
    ids, h = equipe
    doc = _criar(cliente, h["ana"])
    sec = _secao(cliente, h["ana"], doc["id"])
    itens, _ = _item(cliente, h["ana"], doc["id"], sec, "A contratada fornecerá serviços de limpeza")
    a = itens[0]["id"]
    _item(cliente, h["ana"], doc["id"], sec, "Item que será retirado")
    v = cliente.post(f"{URL}/documentos/{doc['id']}/versoes", json={"resumo": "Primeira redação"}, headers=h["ana"])
    assert v.status_code == 201
    n1 = v.json()["numero"]
    # Edita um item e retira outro
    cliente.put(f"{URL}/documentos/{doc['id']}/itens/{a}", json={"conteudo": "A contratada fornecerá serviços de vigilância"}, headers=h["ana"])
    d = _abrir(cliente, h["ana"], doc["id"])
    retirado = d["secoes"][0]["itens"][1]["id"]
    cliente.delete(f"{URL}/documentos/{doc['id']}/itens/{retirado}", headers=h["ana"])
    n2 = cliente.post(f"{URL}/documentos/{doc['id']}/versoes", json={"resumo": "Mudou o objeto"}, headers=h["ana"]).json()["numero"]
    assert n2 > n1
    # O que foi retirado e incrementado nesta versão
    alt = cliente.get(f"{URL}/documentos/{doc['id']}/versoes/{n2}/alteracoes", headers=h["ana"]).json()
    assert [x["html"] for x in alt["itens"] if x["mudanca"] == "removido"] == ["<p>Item que será retirado</p>"]
    alterado = next(x for x in alt["itens"] if x["mudanca"] == "alterado")
    assert "<del>limpeza</del>" in alterado["diferenca"] and "<ins>vigilância</ins>" in alterado["diferenca"]
    # Visualizar a versão antiga, e histórico do item com antes e depois
    foto = cliente.get(f"{URL}/documentos/{doc['id']}/versoes/{n1}", headers=h["ana"]).json()["foto"]
    assert len(foto["itens"]) == 2
    hist = cliente.get(f"{URL}/documentos/{doc['id']}/itens/{a}/historico", headers=h["ana"]).json()
    assert hist[0]["mudanca"] == "editou" and "limpeza" in hist[0]["antes_html"] and "vigilância" in hist[0]["depois_html"]
    # Restaurar só aquele item
    r = cliente.post(f"{URL}/documentos/{doc['id']}/historico/{hist[0]['id']}/restaurar", json={"estado": "antes"}, headers=h["ana"])
    assert r.status_code == 200 and "limpeza" in r.json()["secoes"][0]["itens"][0]["conteudo"]
    # Restaurar a versão 1 devolve o item retirado e cria uma nova versão
    r = cliente.post(f"{URL}/documentos/{doc['id']}/versoes/{n1}/restaurar", headers=h["ana"])
    assert r.status_code == 200, r.text
    assert [i["conteudo"] for i in r.json()["secoes"][0]["itens"]] == ["A contratada fornecerá serviços de limpeza", "Item que será retirado"]
    versoes = cliente.get(f"{URL}/documentos/{doc['id']}/versoes", headers=h["ana"]).json()
    assert versoes[0]["tipo"] == "restauracao" and versoes[0]["resumo"] == f"Restaurou a versão {n1}"
    assert cliente.get(f"{URL}/documentos/{doc['id']}/versoes/99", headers=h["ana"]).status_code == 404
    # Revisor também vê versões, mas não restaura
    cliente.put(f"{URL}/documentos/{doc['id']}/membros/{ids['carla']}", json={"papel": "revisor"}, headers=h["ana"])
    assert cliente.get(f"{URL}/documentos/{doc['id']}/versoes", headers=h["carla"]).status_code == 200
    assert cliente.post(f"{URL}/documentos/{doc['id']}/versoes/{n1}/restaurar", headers=h["carla"]).status_code == 403


def test_conferencia_e_conclusao(cliente, equipe):
    ids, h = equipe
    doc = _criar(cliente, h["ana"], processo="")
    sec = _secao(cliente, h["ana"], doc["id"])
    _secao(cliente, h["ana"], doc["id"], "VAZIA")
    _item(cliente, h["ana"], doc["id"], sec, "Texto")
    c = cliente.get(f"{URL}/documentos/{doc['id']}/conferencia", headers=h["ana"]).json()
    assert c["pode_concluir"] is True and any("sem itens" in a for a in c["alertas"]) and any("SEI" in a for a in c["alertas"])
    # Com alertas, precisa confirmar
    r = cliente.put(f"{URL}/documentos/{doc['id']}/situacao", json={"situacao": "concluido"}, headers=h["ana"])
    assert r.status_code == 409 and r.json()["codigo"] == "conferencia"
    r = cliente.put(f"{URL}/documentos/{doc['id']}/situacao", json={"situacao": "concluido", "confirmar": True}, headers=h["ana"])
    assert r.status_code == 200 and r.json()["situacao"] == "concluido"
    assert cliente.get(f"{URL}/painel", headers=h["ana"]).json()["por_situacao"]["concluido"] == 1


def test_vinculo_com_contrato(cliente, equipe, admin):
    ids, h = equipe
    doc = _criar(cliente, h["ana"])
    contrato = criar_contrato(cliente, admin)
    r = cliente.put(f"{URL}/documentos/{doc['id']}/contrato", json={"contrato_id": contrato["id"]}, headers=h["ana"])
    assert r.status_code == 200 and r.json()["contrato_numero"] == contrato["numero"]
    assert [d["id"] for d in cliente.get(f"{URL}/por-contrato/{contrato['id']}", headers=h["ana"]).json()] == [doc["id"]]
    assert cliente.put(f"{URL}/documentos/{doc['id']}/contrato", json={"contrato_id": "00000000-0000-0000-0000-000000000001"}, headers=h["ana"]).status_code == 404
    assert cliente.put(f"{URL}/documentos/{doc['id']}/contrato", json={"contrato_id": None}, headers=h["ana"]).json()["contrato_id"] is None


def test_tabela_do_tr(cliente, equipe):
    ids, h = equipe
    doc = _criar(cliente, h["ana"], tipo="tr", nome="TR Teste")
    sec = _secao(cliente, h["ana"], doc["id"], "DO OBJETO")
    itens, _ = _item(cliente, h["ana"], doc["id"], sec, "Objeto")
    primeiro = itens[0]
    assert primeiro["tem_tabela_tr"] is True
    r = cliente.post(f"{URL}/documentos/{doc['id']}/itens/{primeiro['id']}/linhas-tr", json={"descricao": "Notebook", "unidade": "UN", "quantidade_objeto": "10.5"}, headers=h["ana"])
    assert r.status_code == 201 and r.json()["secoes"][0]["itens"][0]["linhas_tabela"][0]["descricao"] == "Notebook"
    outro, _ = _item(cliente, h["ana"], doc["id"], sec, "Outro")
    assert cliente.post(f"{URL}/documentos/{doc['id']}/itens/{outro[1]['id']}/linhas-tr", json={"descricao": "x"}, headers=h["ana"]).status_code == 400


def _montar(cliente, h, tipo="tr"):
    doc = _criar(cliente, h, tipo=tipo, nome="Documento de ida e volta")
    sec = _secao(cliente, h, doc["id"], "OBJETO DA CONTRATAÇÃO")
    itens, _ = _item(cliente, h, doc["id"], sec, None, html="<p>Contratação de <strong>serviços</strong> de TI</p>")
    a = itens[0]["id"]
    _item(cliente, h, doc["id"], sec, "Subitem", tipo="subitem", pai_id=a)
    _item(cliente, h, doc["id"], sec, "Inciso", tipo="inciso", pai_id=a)
    _item(cliente, h, doc["id"], sec, "Alínea", tipo="alinea", pai_id=a)
    _item(cliente, h, doc["id"], sec, None, html="<table><tbody><tr><td>c1</td><td>c2</td></tr></tbody></table>")
    sec2 = _secao(cliente, h, doc["id"], "JUSTIFICATIVA")
    _item(cliente, h, doc["id"], sec2, "Porque sim")
    return doc


def test_exportar_word_pdf_e_ida_e_volta(cliente, equipe):
    ids, h = equipe
    doc = _montar(cliente, h["ana"])
    r = cliente.get(f"{URL}/documentos/{doc['id']}/exportar/word", headers=h["ana"])
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    texto = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "1. OBJETO DA CONTRATAÇÃO" in texto and "1.1. Contratação de serviços de TI" in texto and "1.1.1. Subitem" in texto and "I - Inciso" in texto and "a) Alínea" in texto
    pdf = cliente.get(f"{URL}/documentos/{doc['id']}/exportar/pdf", headers=h["ana"])
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    # Ida e volta: o Word exportado importa de volta com a mesma árvore
    previa = cliente.post(f"{URL}/importar-word", files={"arquivo": ("exportado.docx", r.content, "application/octet-stream")}, data={"tipo": "tr", "nome": "Reimportado"}, headers=h["ana"])
    assert previa.status_code == 200, previa.text
    dados = previa.json()
    assert [s["titulo"] for s in dados["secoes"]] == ["OBJETO DA CONTRATAÇÃO", "JUSTIFICATIVA"] and dados["documento_id"] is None
    assert [(i["tipo"], i["conteudo"]) for i in dados["secoes"][0]["itens"] if i["conteudo"] != "c1 c2"] == [("item", "Contratação de serviços de TI"), ("subitem", "Subitem"), ("inciso", "Inciso"), ("alinea", "Alínea")]
    assert dados["totais"]["tabelas"] == 1
    # Confirmar cria o documento do importador
    conf = cliente.post(f"{URL}/importar-word", files={"arquivo": ("exportado.docx", r.content, "application/octet-stream")}, data={"tipo": "tr", "nome": "Reimportado", "confirmar": "true"}, headers=h["ana"])
    novo = _abrir(cliente, h["ana"], conf.json()["documento_id"])
    assert novo["nome"] == "Reimportado" and [(i["marcador"], i["conteudo"]) for i in novo["secoes"][0]["itens"]][:4] == [("1.1.", "Contratação de serviços de TI"), ("1.1.1.", "Subitem"), ("I -", "Inciso"), ("a)", "Alínea")]
    assert "<strong>serviços</strong>" in novo["secoes"][0]["itens"][0]["conteudo_html"]
    assert cliente.get(f"{URL}/documentos/{novo['id']}/versoes", headers=h["ana"]).json()[0]["tipo"] == "importacao"
    # Quem só tem LEITURA não importa
    assert cliente.post(f"{URL}/importar-word", files={"arquivo": ("e.docx", r.content, "application/octet-stream")}, data={"tipo": "tr"}, headers=h["dora"]).status_code == 403


def test_importacao_rejeita_arquivo_invalido(cliente, equipe):
    ids, h = equipe
    r = cliente.post(f"{URL}/importar-word", files={"arquivo": ("x.docx", b"nao sou zip", "application/octet-stream")}, data={"tipo": "etp"}, headers=h["ana"])
    assert r.status_code == 400 and r.json()["codigo"] == "docx_invalido"
    vazio = io.BytesIO()
    Document().save(vazio)
    r = cliente.post(f"{URL}/importar-word", files={"arquivo": ("v.docx", vazio.getvalue(), "application/octet-stream")}, data={"tipo": "etp"}, headers=h["ana"])
    assert r.status_code == 400 and r.json()["codigo"] == "docx_vazio"


def test_importacao_estilos_e_comentarios():
    d = Document()
    d.add_paragraph("1. OBJETO", style=None)
    p = d.add_paragraph()
    p.add_run("texto solto sem numeração")
    d.add_paragraph("OU alternativa")
    buf = io.BytesIO()
    d.save(buf)
    previa = docx_importacao.ler(buf.getvalue(), "a.docx")
    itens = previa["secoes"][0]["itens"] if previa["secoes"][0]["titulo"] != "Informações iniciais" else previa["secoes"][0]["itens"]
    assert any(i["aviso"] for i in itens) and any(i["precisa_revisao"] for i in itens)


def test_sessao_gera_versao_antes_da_primeira_edicao(cliente, equipe):
    ids, h = equipe
    doc = _criar(cliente, h["ana"])
    sec = _secao(cliente, h["ana"], doc["id"])
    _item(cliente, h["ana"], doc["id"], sec, "Primeiro")
    versoes = cliente.get(f"{URL}/documentos/{doc['id']}/versoes", headers=h["ana"]).json()
    # A criação grava a versão 1; as edições da mesma sessão não geram uma por clique
    assert versoes[-1]["tipo"] == "criacao" and len(versoes) <= 2
    with FabricaSessao() as sessao:
        assert sessao.get(DocumentoContratacao, uuid.UUID(doc["id"])).ultima_edicao_em is not None
