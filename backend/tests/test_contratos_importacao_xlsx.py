# Criado por José Eduardo Santana Martins
# Este arquivo serve para testar a importação de contrato por planilha XLSX (prévia, gravação, erros e ACL).
"""Importação de contrato por XLSX: as planilhas são geradas a partir do modelo oficial, preenchido pelos testes."""

from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.core.banco import FabricaSessao
from app.models.contratos import Contrato, EmpresaContratada
from tests.apoio_contratos import criar_empresa, gerar_cnpj, gerar_cpf
from tests.conftest import cabecalho, criar_usuario

URL = "/api/contratos/importacao-xlsx"
MODELO = Path(__file__).resolve().parents[1] / "app" / "recursos" / "modelo-importacao-contrato.xlsx"
CNPJ = gerar_cnpj("445556660001")
CPF = gerar_cpf("987654321")

# Valores do cabeçalho por linha do modelo (coluna C)
CABECALHO = {
    2: "7/2026", 3: CNPJ, 4: "Limpa Tudo Serviços Ltda", 5: "Limpa Tudo", 6: "Rua B, 20",
    7: "João Preposto", 8: CPF, 9: "joao@limpatudo.com", 10: "(11) 99999-0000",
    11: "Limpeza anexo", 12: datetime(2026, 3, 1), 13: "12 meses", 14: 60, 15: "Mensal", 16: "Março",
    17: "Limpeza predial do anexo", 18: "SPI-PRC-2026/00010", 19: "https://sei.sp.gov.br/g", 20: "SPI-PRC-2026/00011",
    21: "https://sei.sp.gov.br/e",
}
# Itens a partir da linha 32 (as linhas 30–31 do modelo são a legenda das opções, sem descrição)
ITENS = [
    ["Limpeza diária", "Contínuo", "Pró-rata", "01", 339039, "123", "456", 2, None, "1.000,00"],
    ["Material de limpeza", "Sob demanda", "Sempre Integral", "02", "339030", "124", "457", None, 100, 10.5],
]


def planilha(cabecalho: dict | None = None, itens: list | None = None) -> bytes:
    """Preenche o modelo com os valores informados (os padrões acima, sobrescritos pelo argumento)."""
    livro = load_workbook(MODELO)
    folha = livro.active
    for linha, valor in {**CABECALHO, **(cabecalho or {})}.items():
        folha.cell(linha, 3, valor)
    for deslocamento, valores in enumerate(ITENS if itens is None else itens):
        for coluna, valor in enumerate(valores, start=1):
            folha.cell(32 + deslocamento, coluna, valor)
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def enviar(cliente, h, conteudo: bytes, caminho: str = "/previa", nome: str = "contrato.xlsx"):
    """Envia a planilha como multipart (campo `arquivo`)."""
    tipo = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return cliente.post(URL + caminho, files={"arquivo": (nome, conteudo, tipo)}, headers=h)


def test_previa_converte_os_valores_e_nao_grava(cliente, admin):
    """A prévia converte datas, meses, nomes e números no formato brasileiro, sem gravar nada."""
    r = enviar(cliente, admin, planilha())
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["pode_importar"] and p["erros"] == []
    c = p["contrato"]
    assert (c["numero"], c["data_inicio"], c["data_fim"]) == ("007/2026", "2026-03-01", "2027-02-28")
    assert (c["vigencia_inicial_meses"], c["periodicidade_meses"], c["mes_reajuste"]) == (12, 1, 3)
    assert p["empresa"]["existente"] is False and p["empresa"]["cnpj"] == CNPJ
    assert p["preposto"] == {"existente": False, "cpf": CPF, "nome": "João Preposto", "email": "joao@limpatudo.com", "telefone": "(11) 99999-0000"}
    assert [(i["tipo"], i["calcula_pro_rata"], i["codigo_natureza_despesa"]) for i in p["itens"]] == [
        ("continuo", True, "339039"), ("sob_demanda", False, "339030")]
    assert p["itens"][0]["valor_unitario"] == "1000.00" and p["itens"][1]["quantidade_total"] == "100.0000"
    # 2 × 1.000 × 12 + 100 × 10,50
    assert p["valor_global_estimado"] == "25050.00"
    with FabricaSessao() as sessao:
        assert sessao.query(Contrato).count() == 0 and sessao.query(EmpresaContratada).count() == 0


def test_importacao_grava_empresa_preposto_contrato_e_itens(cliente, admin):
    """Sem erros, a importação cria tudo e devolve o detalhe do contrato (sem equipe)."""
    r = enviar(cliente, admin, planilha(), caminho="")
    assert r.status_code == 201, r.text
    contrato = r.json()
    assert contrato["numero"] == "007/2026" and contrato["equipe"] == [] and len(contrato["itens"]) == 2
    assert contrato["empresa"]["cnpj"] == CNPJ
    empresa = cliente.get(f"/api/contratos/empresas/{contrato['empresa']['id']}", headers=admin).json()
    assert [p["cpf"] for p in empresa["prepostos"]] == [CPF]
    # Enviar de novo: o número já existe
    r = enviar(cliente, admin, planilha(), caminho="")
    assert r.status_code == 400 and r.json()["erros"] == [{"linha": 2, "campo": "Nro do Contrato", "mensagem": "Já existe um contrato com o número 007/2026."}]


def test_empresa_existente_e_reaproveitada_com_aviso(cliente, admin):
    """CNPJ já cadastrado: a empresa é mantida como está e a divergência vira aviso."""
    existente = criar_empresa(cliente, admin, base="44555666", razao="Outro Nome Ltda")
    assert existente["cnpj"] == CNPJ
    p = enviar(cliente, admin, planilha()).json()
    assert p["pode_importar"] and p["empresa"]["existente"] and p["empresa"]["razao_social"] == "Outro Nome Ltda"
    assert any("Razão Social" in a for a in p["avisos"])
    contrato = enviar(cliente, admin, planilha(), caminho="").json()
    assert contrato["empresa"]["id"] == existente["id"]
    with FabricaSessao() as sessao:
        assert sessao.query(EmpresaContratada).count() == 1


def test_erros_apontam_linha_e_campo_e_nada_e_gravado(cliente, admin):
    """CNPJ inválido, tipo desconhecido, link sem http e vigência máxima menor bloqueiam a importação."""
    itens = [["Limpeza", "Mensalista", "Pró-rata", "01", "339039", "123", "456", 2, None, "10,00"]]
    conteudo = planilha({3: "11.111.111/1111-11", 14: 6, 19: "sei.sp.gov.br"}, itens)
    p = enviar(cliente, admin, conteudo).json()
    erros = {(e["linha"], e["campo"]) for e in p["erros"]}
    assert (3, "CNPJ") in erros and (32, "Item 1 · Tipo") in erros and (19, "Processo Gestão · Link") in erros
    assert (14, "Vigência Máxima") in erros and not p["pode_importar"]
    r = enviar(cliente, admin, conteudo, caminho="")
    assert r.status_code == 400 and r.json()["codigo"] == "invalido" and r.json()["erros"]
    with FabricaSessao() as sessao:
        assert sessao.query(Contrato).count() == 0 and sessao.query(EmpresaContratada).count() == 0


def test_equipe_preenchida_so_gera_aviso(cliente, admin):
    """A equipe da planilha não é importada: fica registrado um aviso."""
    p = enviar(cliente, admin, planilha({22: "Fulano Gestor"})).json()
    assert p["pode_importar"] and any("equipe" in a for a in p["avisos"])


def test_arquivo_que_nao_e_xlsx_e_recusado(cliente, admin):
    """Extensão diferente ou conteúdo que não é planilha: 400."""
    assert enviar(cliente, admin, b"a;b;c", nome="contrato.csv").status_code == 400
    r = enviar(cliente, admin, b"isto nao e uma planilha")
    assert r.status_code == 400 and "xlsx" in r.json()["detalhe"]


def test_acl_da_importacao_e_modelo(cliente, admin):
    """Sem a liberação em `importacao-contratos`, o usuário recebe 403 mesmo podendo cadastrar contratos."""
    fulano, liberado = criar_usuario("fulano"), criar_usuario("liberado")
    recurso = cliente.post("/api/acl/recursos", json={"nome": "Importação", "slug": "importacao-contratos"}, headers=admin).json()
    corpo = {"recurso_id": recurso["id"], "nivel": "MODIFICACAO", "usuarios_ids": [liberado], "setores_ids": []}
    assert cliente.post("/api/acl/regras", json=corpo, headers=admin).status_code == 201
    r = enviar(cliente, cabecalho(cliente, "fulano"), planilha())
    assert r.status_code == 403 and r.json()["recurso"] == "importacao-contratos"
    assert fulano
    assert enviar(cliente, cabecalho(cliente, "liberado"), planilha()).status_code == 200
    modelo = cliente.get(f"{URL}/modelo", headers=cabecalho(cliente, "liberado"))
    assert modelo.status_code == 200 and modelo.content[:2] == b"PK"
