# Criado por José Eduardo Santana Martins
# Este arquivo serve para testar o pós-assinatura da prorrogação: prazo do PNCP em dias úteis, situação, datas, planilhas de garantia e avisos.
"""Prorrogação: PNCP (20/10 dias úteis), DOE, garantia (planilhas para a DOF) e lembretes."""

import io
import itertools
import uuid
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.core.banco import FabricaSessao
from app.models.contratos import Contrato
from app.services import calendario_util
from app.services.contratos import servico_pos_assinatura
from app.tarefas import notificacoes as tarefa_notificacoes
from tests.apoio_contratos import PDF, criar_contrato, criar_empresa, dados_contrato, definir_parametros, preparar_registro_prorrogacao


_SEQUENCIA = itertools.count(1)


def _registrar(cliente, admin, assinada="2026-10-05", **extras) -> tuple[dict, str, dict]:
    n = next(_SEQUENCIA)
    empresa = criar_empresa(cliente, admin, base=f"{30000000 + n}")
    contrato = criar_contrato(cliente, admin, empresa_id=empresa["id"], numero=f"{n:03d}/2026", vigencia_maxima_meses=60, **extras)
    base = f"/api/contratos/{contrato['id']}"
    cliente.put(f"{base}/prorrogacao", json={"meses": 12}, headers=admin)
    preparar_registro_prorrogacao(cliente, base, admin)
    r = cliente.post(f"{base}/prorrogacao/registrar", data={"assinada_em": assinada}, files={"termo": ("ta.pdf", PDF)}, headers=admin)
    assert r.status_code == 200, r.text
    return contrato, base, r.json()[0]


def test_prazo_do_pncp_conta_dias_uteis_pela_modalidade(cliente, admin):
    # 05/10/2026 é segunda-feira: 20 dias úteis → 02/11/2026 (sem feriados cadastrados); direta: 10 → 19/10/2026
    _, _, licitacao = _registrar(cliente, admin, modalidade_contratacao="licitacao")
    assert licitacao["pncp_prazo_ate"] == "2026-11-02"
    _, _, direta = _registrar(cliente, admin, modalidade_contratacao="dispensa")
    assert direta["pncp_prazo_ate"] == "2026-10-19"
    _, _, sem = _registrar(cliente, admin)
    assert sem["pncp_prazo_ate"] is None and sem["situacao_pncp"] == "sem_prazo"
    # O número de dias vem dos Parâmetros de Contratos
    definir_parametros(pncp_dias_uteis_direta=5)
    _, _, curto = _registrar(cliente, admin, modalidade_contratacao="inexigibilidade")
    assert curto["pncp_prazo_ate"] == "2026-10-12"


def test_situacao_do_pncp_no_prazo_vencido_e_publicado(cliente, admin):
    _, _, p = _registrar(cliente, admin, modalidade_contratacao="dispensa")
    with FabricaSessao() as sessao:
        contrato = sessao.query(Contrato).one() if sessao.query(Contrato).count() == 1 else sessao.query(Contrato).first()
        prorrogacao = contrato.prorrogacoes[0]
        assert servico_pos_assinatura.situacao_pncp(sessao, prorrogacao, date(2026, 10, 14)) == ("no_prazo", 3)
        assert servico_pos_assinatura.situacao_pncp(sessao, prorrogacao, date(2026, 10, 20)) == ("vencido", None)
        prorrogacao.pncp_publicado_em = date(2026, 10, 20)
        assert servico_pos_assinatura.situacao_pncp(sessao, prorrogacao, date(2026, 10, 30)) == ("publicado", None)


def test_pos_assinatura_grava_datas_e_recusa_antes_da_assinatura(cliente, admin):
    contrato, base, p = _registrar(cliente, admin, modalidade_contratacao="licitacao", exige_garantia=True, garantia_modalidade="seguro_garantia")
    assert p["garantia_pendente"] is True and p["situacao_pncp"] in ("no_prazo", "vencido")
    url = f"{base}/prorrogacoes/{p['id']}/pos-assinatura"
    corpo = {"pncp_publicado_em": "2026-10-20", "pncp_numero": "PNCP-123", "doe_publicado_em": "2026-10-21", "garantia_comprovada_em": "2026-10-22",
             "garantia_enviada_dof_em": "2026-10-23", "contratos_gov_atualizado_em": "2026-10-24", "dados_garantia": {"responsavel_nome": "Ana"}}
    r = cliente.put(url, json=corpo, headers=admin)
    assert r.status_code == 200, r.text
    novo = r.json()[0]
    assert novo["situacao_pncp"] == "publicado" and novo["pncp_numero"] == "PNCP-123" and novo["garantia_pendente"] is False and novo["dados_garantia"] == {"responsavel_nome": "Ana"}
    assert cliente.put(url, json={**corpo, "pncp_publicado_em": "2026-10-01"}, headers=admin).status_code == 400
    assert cliente.put(url, json={**corpo, "dados_garantia": {"inventado": "x"}}, headers=admin).status_code == 400


def test_planilha_de_garantia_em_dinheiro_e_de_seguro(cliente, admin):
    _, base, p = _registrar(cliente, admin, modalidade_contratacao="licitacao", exige_garantia=True, garantia_modalidade="caucao_dinheiro", garantia_percentual="5", garantia_valor="1000.00")
    cliente.put(f"{base}/prorrogacoes/{p['id']}/pos-assinatura", json={"dados_garantia": {"dare": "9988", "responsavel_nome": "Ana Souza", "ultima_nl": "2026NL00123"}}, headers=admin)
    r = cliente.get(f"{base}/prorrogacoes/{p['id']}/garantia/planilha", headers=admin)
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    f = load_workbook(io.BytesIO(r.content)).active
    assert f["D8"].value == "I - Caução em dinheiro" and f["C14"].value == "ACME Serviços Ltda" and f["C18"].value == "R$ 1.000,00" and f["D17"].value == "5,00%"
    assert "9988" in f["F29"].value and f["A35"].value.endswith("Ana Souza") and f["F24"].value == "Nº 2026NL00123" and "01/01/2027 a 31/12/2027" in f["F23"].value

    empresa = criar_empresa(cliente, admin, base="22333444", razao="Beta Ltda")
    contrato = criar_contrato(cliente, admin, empresa_id=empresa["id"], numero="902/2026", vigencia_maxima_meses=60, exige_garantia=True, garantia_modalidade="seguro_garantia",
                              garantia_emissor="Seguradora X", garantia_cnpj_emissor="11.222.333/0001-81", garantia_apolice="AP-77", garantia_valor="2500.00")
    b2 = f"/api/contratos/{contrato['id']}"
    cliente.put(f"{b2}/prorrogacao", json={"meses": 12}, headers=admin)
    preparar_registro_prorrogacao(cliente, b2, admin)
    p2 = cliente.post(f"{b2}/prorrogacao/registrar", data={"assinada_em": "2026-10-05"}, files={"termo": ("ta.pdf", PDF)}, headers=admin).json()[0]
    f2 = load_workbook(io.BytesIO(cliente.get(f"{b2}/prorrogacoes/{p2['id']}/garantia/planilha", headers=admin).content)).active
    assert f2["D9"].value == "Seguradora X" and "11.222.333/0001-81" in f2["H9"].value and f2["A27"].value.endswith("AP-77") and f2["C13"].value == "Beta Ltda"
    # Sem garantia cadastrada, não há planilha
    _, b3, p3 = _registrar(cliente, admin)
    assert cliente.get(f"{b3}/prorrogacoes/{p3['id']}/garantia/planilha", headers=admin).status_code == 400


def test_lembretes_do_pncp_e_da_garantia_saem_uma_vez(cliente, admin):
    from tests.conftest import cabecalho, criar_usuario
    from tests.apoio_contratos import restringir_contratos

    gestora = criar_usuario("gestora")
    restringir_contratos(cliente, admin, {gestora: "MODIFICACAO"})
    contrato = criar_contrato(cliente, admin, numero="903/2026", equipe={"gestor": gestora}, vigencia_maxima_meses=60, modalidade_contratacao="dispensa", exige_garantia=True, garantia_modalidade="seguro_garantia")
    base = f"/api/contratos/{contrato['id']}"
    cliente.put(f"{base}/prorrogacao", json={"meses": 12}, headers=admin)
    preparar_registro_prorrogacao(cliente, base, admin)
    cliente.post(f"{base}/prorrogacao/registrar", data={"assinada_em": "2026-10-05"}, files={"termo": ("ta.pdf", PDF)}, headers=admin)
    with FabricaSessao() as sessao:
        contratos = [sessao.get(Contrato, uuid.UUID(contrato["id"]))]
        # Prazo 19/10; aviso 5 dias úteis antes (12/10); dia 13/10: faltam 4 dias úteis → avisa PNCP; garantia: já passaram 5 dias úteis da assinatura em 12/10
        assert tarefa_notificacoes._avisar_pos_assinatura(sessao, contratos, date(2026, 10, 6)) == 0  # noqa: SLF001
        primeiro = tarefa_notificacoes._avisar_pos_assinatura(sessao, contratos, date(2026, 10, 13))  # noqa: SLF001
        assert primeiro == 2
        assert tarefa_notificacoes._avisar_pos_assinatura(sessao, contratos, date(2026, 10, 14)) == 0  # noqa: SLF001
        # Vencido: aviso novo, uma vez só
        assert tarefa_notificacoes._avisar_pos_assinatura(sessao, contratos, date(2026, 10, 21)) == 1  # noqa: SLF001
        assert tarefa_notificacoes._avisar_pos_assinatura(sessao, contratos, date(2026, 10, 22)) == 0  # noqa: SLF001
