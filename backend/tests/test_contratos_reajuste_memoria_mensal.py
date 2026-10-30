# Criado por José Eduardo Santana Martins
# Este arquivo serve para testar a memória de cálculo mês a mês do reajuste (meses pagos, em haver, resumo executivo e nova versão pela OB).
"""Reajuste: memória por mês medido (com a situação da OB), resumo executivo e planilha com uma aba por mês."""

import io
from datetime import date
from decimal import Decimal

import openpyxl

from tests.apoio_contratos import PDF, definir_parametros, juntar_nf
from tests.test_contratos_retencao import _ambiente, cenario  # noqa: F401 - fixtures (competência 01/2026 com medição concluída, data fixa de 15/03/2026)
from tests.test_contratos_execucao import _url


def _reajustar(cliente, contrato, h):
    painel = cliente.post(_url(contrato, "/reajustes"), json={"sequencia_vigencia": 1, "mes_referencia": "2026-01-01"}, headers=h).json()
    reajuste = painel["em_andamento"]
    url = _url(contrato, f"/reajustes/{reajuste['id']}")
    itens = [{"item_id": i["item_id"], "indice_percentual": "5"} for i in reajuste["itens"]]
    return url, cliente.put(f"{url}/memoria", json={"itens": itens}, headers=h).json()["em_andamento"]


def test_memoria_mes_a_mes_com_mes_pago_e_resumo_executivo(cliente, cenario):  # noqa: F811
    contrato, base, gestora = cenario
    definir_parametros(corte_alertas_a_partir_de=date(2026, 9, 1))  # competência retroativa: medição → NF → OB
    assert juntar_nf(cliente, base, gestora, "2105.00", "123").status_code == 200
    # Medido, sem OB: já entra na memória, marcado como aguardando a OB
    url, r = _reajustar(cliente, contrato, gestora)
    assert [(m["competencia"], m["paga"]) for m in r["resumo_meses"]["meses"]] == [("01/2026", False)]
    diferenca = Decimal(r["diferenca_retroativa"])
    assert Decimal(r["resumo_meses"]["total_em_haver"]) == diferenca
    cliente.post(f"{url}/evidencia", files={"arquivo": ("ev.pdf", PDF)}, headers=gestora)
    v1 = cliente.post(f"{url}/memoria/arquivos", headers=gestora).json()["em_andamento"]["memorias"]
    assert len(v1) == 1
    # A OB sai: o mês passa a ser pago, o hash muda e nasce a versão 2
    assert cliente.post(f"{base}/ordem-bancaria", files={"arquivo": ("ob.pdf", PDF)}, headers=gestora).status_code == 200
    r = cliente.get(_url(contrato, "/reajustes"), headers=gestora).json()["em_andamento"]
    resumo = r["resumo_meses"]
    assert [(m["competencia"], m["paga"]) for m in resumo["meses"]] == [("01/2026", True)]
    assert Decimal(resumo["total_em_haver"]) == diferenca and Decimal(resumo["meses"][0]["em_haver"]) == diferenca
    assert Decimal(resumo["meses"][0]["valor_reajustado"]) - Decimal(resumo["meses"][0]["valor_pago"]) == diferenca
    memorias = cliente.post(f"{url}/memoria/arquivos", headers=gestora).json()["em_andamento"]["memorias"]
    assert len(memorias) == 2
    # PDF: seção do mês pago e o resumo executivo
    from pypdf import PdfReader

    pdf = cliente.get(f"{url}/arquivos/{memorias[-1]['pdf']['anexo_id']}", headers=gestora).content
    texto = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "Memória de 01/2026 (OB paga)" in texto and "Memória de cálculo executiva" in texto and "Total em haver" in texto
    # Planilha: aba do reajuste, uma aba por mês pago e o resumo executivo
    xlsx = cliente.get(f"{url}/arquivos/{memorias[-1]['xlsx']['anexo_id']}", headers=gestora).content
    pasta = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert pasta.sheetnames == ["Reajuste", "01-2026", "Resumo"]
    mes = pasta["01-2026"]
    textos = [str(c.value) for linha in mes.iter_rows() for c in linha if c.value is not None]
    assert any("SIAFÍSICO" in t for t in textos) and any("Subtotal 01/2026" in t for t in textos)
    resumo_texto = " ".join(str(c.value) for linha in pasta["Resumo"].iter_rows() for c in linha if c.value is not None)
    assert "01/2026" in resumo_texto and "Total em haver" in resumo_texto and "Valor reajustado" in resumo_texto


def test_mes_medido_sem_ob_entra_na_memoria(cliente, cenario):  # noqa: F811
    contrato, base, gestora = cenario
    # Sem OB concluída o mês medido entra na memória e no total, marcado como aguardando a OB
    url, r = _reajustar(cliente, contrato, gestora)
    assert [(m["competencia"], m["paga"]) for m in r["resumo_meses"]["meses"]] == [("01/2026", False)]
    assert Decimal(r["resumo_meses"]["total_em_haver"]) == Decimal(r["diferenca_retroativa"])


def test_reajuste_comum_esta_disponivel_em_todas_as_vigencias(cliente, cenario):  # noqa: F811
    """Também a vigência retroativa (e as já pagas) aparece para o reajuste comum; uma vez reajustada, some da lista."""
    contrato, base, gestora = cenario
    painel = cliente.get(_url(contrato, "/reajustes"), headers=gestora).json()
    assert [v["sequencia"] for v in painel["vigencias_disponiveis"]] == [1]


def test_pdf_da_memoria_tem_uma_pagina_para_itens_cada_mes_e_o_resumo(cliente, cenario):  # noqa: F811
    from pypdf import PdfReader

    contrato, base, gestora = cenario
    definir_parametros(corte_alertas_a_partir_de=date(2026, 9, 1))
    assert juntar_nf(cliente, base, gestora, "2105.00", "123").status_code == 200
    assert cliente.post(f"{base}/ordem-bancaria", files={"arquivo": ("ob.pdf", PDF)}, headers=gestora).status_code == 200
    url, _ = _reajustar(cliente, contrato, gestora)
    cliente.post(f"{url}/evidencia", files={"arquivo": ("ev.pdf", PDF)}, headers=gestora)
    memoria = cliente.post(f"{url}/memoria/arquivos", headers=gestora).json()["em_andamento"]["memorias"][-1]
    paginas = [p.extract_text() for p in PdfReader(io.BytesIO(cliente.get(f"{url}/arquivos/{memoria['pdf']['anexo_id']}", headers=gestora).content)).pages]
    # Identificação, itens, mês pago e resumo executivo: cada um começa na sua página
    assert "Identificação" in paginas[0] and "Itens" not in paginas[0].split("Identificação")[0]
    assert any(p.lstrip().startswith("Itens") or "Itens (" in p[:200] for p in paginas[1:2])
    assert sum("Memória de 01/2026" in p for p in paginas) == 1 and "Memória de 01/2026" not in paginas[1]
    assert "Memória de cálculo executiva" in paginas[-1] and "Memória de 01/2026" not in paginas[-1]
    # O resumo executivo é uma tabela (mês, valor pago, valor reajustado, em haver) com o total no rodapé
    assert "Valor reajustado" in paginas[-1] and "Em haver" in paginas[-1] and "Total em haver" in paginas[-1] and "01/2026" in paginas[-1]


def test_calculadora_ao_vivo_nao_grava_nada(cliente, cenario):  # noqa: F811
    """A prévia devolve os totais com os índices digitados, sem salvar: o banco continua com índice 0 e os totais antigos."""
    contrato, base, gestora = cenario
    painel = cliente.post(_url(contrato, "/reajustes"), json={"sequencia_vigencia": 1, "mes_referencia": "2026-01-01"}, headers=gestora).json()
    reajuste = painel["em_andamento"]
    url = _url(contrato, f"/reajustes/{reajuste['id']}")
    itens = [{"item_id": i["item_id"], "indice_percentual": "10"} for i in reajuste["itens"]]
    previa = cliente.post(f"{url}/memoria/previa", json={"itens": itens}, headers=gestora)
    assert previa.status_code == 200, previa.text
    dados = previa.json()
    assert Decimal(dados["base_reajustada"]) == (Decimal(dados["base_atual"]) * Decimal("1.1")).quantize(Decimal("0.01")) and dados["itens"][0]["valor_unitario_reajustado"] == "1100.0000"
    assert Decimal(dados["valor_global_reajustado"]) > Decimal(dados["valor_global_atual"])
    # Nada foi gravado: o reajuste salvo segue com índice 0 e totais iguais aos atuais
    salvo = cliente.get(_url(contrato, "/reajustes"), headers=gestora).json()["em_andamento"]
    assert Decimal(salvo["itens"][0]["indice_percentual"]) == 0 and salvo["base_reajustada"] == salvo["base_atual"]
    assert cliente.post(f"{url}/memoria/previa", json={"itens": [{"item_id": "00000000-0000-0000-0000-000000000000", "indice_percentual": "1"}]}, headers=gestora).status_code == 400
