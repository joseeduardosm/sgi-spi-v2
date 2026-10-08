# Criado por José Eduardo Santana Martins
# Este arquivo serve para testar o Módulo Melhorias: envio com prints, triagem pela ACL, avisos, conversão em tarefa e relatórios.
"""Módulo Melhorias (`/api/melhorias`)."""

import json
from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from sqlalchemy import select

from app.core.banco import FabricaSessao
from app.models.mensagem import EntregaMensagem, Mensagem
from tests.apoio_rh import simular_smtp
from tests.conftest import cabecalho, criar_usuario

URL = "/api/melhorias"


@pytest.fixture(autouse=True)
def _smtp(monkeypatch):
    simular_smtp(monkeypatch)


@pytest.fixture
def pessoas(cliente, admin):
    """Ana envia sugestões, Tito faz a triagem (CONTROLE_TOTAL), Leo só tem leitura no recurso."""
    ids = {login: criar_usuario(login, nome_completo=nome, email=f"{login}@sp.gov.br") for login, nome in (
        ("ana", "Ana Autora"), ("tito", "Tito Triagem"), ("leo", "Leo Leitor"))}
    r = cliente.post("/api/acl/recursos", json={"nome": "Melhorias", "slug": "melhorias"}, headers=admin)
    assert r.status_code == 201, r.text
    for login, nivel in (("tito", "CONTROLE_TOTAL"), ("leo", "LEITURA")):
        corpo = {"recurso_id": r.json()["id"], "nivel": nivel, "usuarios_ids": [ids[login]], "setores_ids": []}
        assert cliente.post("/api/acl/regras", json=corpo, headers=admin).status_code == 201
    return ids, {k: cabecalho(cliente, k) for k in ids}


def _png(largura=800, altura=500) -> bytes:
    saida = BytesIO()
    Image.new("RGB", (largura, altura), (40, 90, 200)).save(saida, "PNG")
    return saida.getvalue()


def _enviar(cliente, h, texto="O filtro de contratos poderia lembrar a última busca.", tela="/contratos/painel", arquivos=()):
    return cliente.post(f"{URL}/sugestoes", data={"dados": json.dumps({"texto": texto, "tela": tela})},
                        files=[("arquivos", a) for a in arquivos], headers=h)


def _avisos(prefixo: str) -> list[tuple[int, str]]:
    with FabricaSessao() as s:
        return sorted(s.execute(select(EntregaMensagem.destinatario_id, Mensagem.assunto).join(Mensagem).where(Mensagem.chave.startswith(prefixo))))


def test_envio_com_prints_e_aviso_a_triagem(cliente, admin, pessoas):
    ids, h = pessoas
    # Formato recusado pelo conteúdo, texto vazio e prints demais
    r = _enviar(cliente, h["ana"], arquivos=[("tela.png", b"<html>nada</html>", "image/png")])
    assert r.status_code == 400 and "tela.png" in r.json()["detalhe"]
    assert _enviar(cliente, h["ana"], texto="   ").status_code == 422
    assert _enviar(cliente, h["ana"], arquivos=[(f"p{i}.png", _png(), "image/png") for i in range(4)]).status_code == 400

    r = _enviar(cliente, h["ana"], arquivos=[("print.png", _png(), "image/png")])
    assert r.status_code == 201, r.text
    s = r.json()
    assert s["numero"] == 1 and s["situacao"] == "nova" and s["modulo"] == "contratos" and len(s["prints"]) == 1
    assert "observacao_interna" not in s
    # Aviso a quem faz a triagem (Tito e o SuperRoot), não à autora nem ao Leo
    destinatarios = {d for d, _ in _avisos("melhoria-nova:")}
    assert ids["tito"] in destinatarios and ids["ana"] not in destinatarios and ids["leo"] not in destinatarios
    # O print: a autora e a triagem baixam; o Leo não
    url = s["prints"][0]["url"]
    assert cliente.get(url, headers=h["ana"]).status_code == 200
    assert cliente.get(url, headers=h["tito"]).status_code == 200
    assert cliente.get(url, headers=h["leo"]).status_code == 404
    # Minhas sugestões: só as próprias
    assert cliente.get(f"{URL}/minhas", headers=h["ana"]).json()["total"] == 1
    assert cliente.get(f"{URL}/minhas", headers=h["leo"]).json()["total"] == 0
    assert cliente.get(f"{URL}/minhas/1", headers=h["leo"]).status_code == 404


def test_triagem_resposta_ao_autor_e_observacao_interna(cliente, admin, pessoas):
    ids, h = pessoas
    _enviar(cliente, h["ana"])
    _enviar(cliente, h["leo"], texto="Tela de férias sem atalho para o calendário", tela="/rh/ferias")
    # Sem CONTROLE_TOTAL não há triagem
    assert cliente.get(f"{URL}/sugestoes", headers=h["leo"]).status_code == 403
    assert cliente.get(f"{URL}/acesso", headers=h["leo"]).json() == {"triagem": False}
    assert cliente.get(f"{URL}/acesso", headers=h["tito"]).json() == {"triagem": True}

    lista = cliente.get(f"{URL}/sugestoes", params={"modulo": "rh"}, headers=h["tito"]).json()
    assert lista["total"] == 1 and lista["totais"]["nova"] == 1 and set(lista["modulos"]) == {"contratos", "rh"}
    assert cliente.get(f"{URL}/sugestoes", params={"busca": "#1"}, headers=h["tito"]).json()["itens"][0]["numero"] == 1

    corpo = {"situacao": "em_analise", "resposta_publica": "Vamos avaliar no próximo ciclo.", "observacao_interna": "Depende do cache do filtro."}
    r = cliente.put(f"{URL}/sugestoes/1", json=corpo, headers=h["tito"])
    assert r.status_code == 200, r.text
    assert r.json()["situacao"] == "em_analise" and r.json()["eventos"][0]["situacao_nova"] == "em_analise"
    # SLA da triagem: a primeira mudança de situação é a resposta (dentro do prazo de 3 dias úteis) e ainda falta resolver
    sla = r.json()["sla"]
    assert (sla["meta_resposta_dias"], sla["meta_resolucao_dias"], sla["situacao_resposta"]) == (3, 15, "cumprido") and sla["resolvido_em"] is None
    assert "sla" not in cliente.get(f"{URL}/minhas/1", headers=h["ana"]).json()  # o autor não vê o SLA interno
    # A autora recebe o aviso e vê a resposta, mas não a observação interna
    assert _avisos("melhoria-tratada:") == [(ids["ana"], "Sugestão de melhoria #1: Em análise")]
    minha = cliente.get(f"{URL}/minhas/1", headers=h["ana"]).json()
    assert minha["resposta_publica"] == "Vamos avaliar no próximo ciclo." and "observacao_interna" not in minha
    # Só a observação interna muda: grava, sem novo aviso
    corpo["observacao_interna"] = "Conversar com o time de contratos."
    assert cliente.put(f"{URL}/sugestoes/1", json=corpo, headers=h["tito"]).status_code == 200
    assert len(_avisos("melhoria-tratada:")) == 1


def test_converter_em_tarefa_e_relatorios(cliente, admin, pessoas):
    ids, h = pessoas
    _enviar(cliente, h["ana"], arquivos=[("print.png", _png(), "image/png")])
    prazo = (datetime.now(timezone.utc) + timedelta(days=10)).isoformat()
    r = cliente.post(f"{URL}/sugestoes/1/tarefa", json={"titulo": "Lembrar a última busca no painel", "prazo": prazo}, headers=h["tito"])
    assert r.status_code == 200, r.text
    assert r.json()["situacao"] == "aceita" and r.json()["tarefa_numero"]
    tarefa = cliente.get(f"/api/tarefas/{r.json()['tarefa_numero']}", headers=h["tito"]).json()
    assert "Sugestão de melhoria #1" in tarefa["descricao"] and "Ana Autora" in tarefa["descricao"]
    assert cliente.post(f"{URL}/sugestoes/1/tarefa", json={"titulo": "De novo", "prazo": prazo}, headers=h["tito"]).status_code == 409
    assert (ids["ana"], "Sugestão de melhoria #1: Aceita") in _avisos("melhoria-tratada:")

    x = cliente.get(f"{URL}/sugestoes/exportar", headers=h["tito"])
    assert x.status_code == 200
    folha = load_workbook(BytesIO(x.content)).active
    valores = [str(c.value) for linha in folha.iter_rows() for c in linha]
    assert any("filtro de contratos" in v for v in valores) and "Aceita" in valores
    p = cliente.get(f"{URL}/sugestoes/relatorio", params={"situacao": "aceita"}, headers=h["tito"])
    assert p.status_code == 200 and p.headers["content-type"] == "application/pdf"
    texto = "".join(pagina.extract_text() for pagina in PdfReader(BytesIO(p.content)).pages)
    assert "Relatório de sugestões de melhoria" in texto and "#1" in texto and "Ana Autora" in texto
    assert cliente.get(f"{URL}/sugestoes/relatorio", headers=h["leo"]).status_code == 403
    assert cliente.get(f"{URL}/sugestoes/relatorio", params={"inicio": "2026-02-01", "fim": "2026-01-01"}, headers=h["tito"]).status_code == 400


def test_sem_regras_so_superroot_faz_triagem(cliente, admin):
    criar_usuario("bia", nome_completo="Bia", email="bia@sp.gov.br")
    hb = cabecalho(cliente, "bia")
    assert _enviar(cliente, hb).status_code == 201
    assert cliente.get(f"{URL}/sugestoes", headers=hb).status_code == 403
    assert cliente.get(f"{URL}/sugestoes", headers=admin).json()["total"] == 1
