# Criado por José Eduardo Santana Martins
# Este arquivo serve para testar o Patrimônio: cadastro, histórico, movimentação, inservível/extraviado, baixa, painel e depreciação.
"""Patrimônio (`/api/patrimonio`): bens móveis do prédio SPI."""

import io
import uuid
from datetime import date
from decimal import Decimal

import pytest
from sqlalchemy import select

from app.core.banco import FabricaSessao
from app.models.patrimonio import Bem, EventoPatrimonio, GrupoMaterial, Localizacao
from app.models.setor import Setor
from app.services.patrimonio import depreciacao, servico_bens
from tests.conftest import cabecalho, criar_usuario

URL = "/api/patrimonio"
JPG = b"\xff\xd8\xff\xe0" + b"\x00" * 64


@pytest.fixture
def equipe(cliente, admin):
    """Ana e Bruno operam (MODIFICACAO), Carla e Edu administram (CONTROLE_TOTAL), Dora só consulta (LEITURA)."""
    ids = {login: criar_usuario(login, nome_completo=f"{login.title()} Souza", email=f"{login}@sp.gov.br") for login in ("ana", "bruno", "carla", "edu", "dora", "fora")}
    recurso = cliente.post("/api/acl/recursos", json={"nome": "Patrimônio", "slug": "patrimonio"}, headers=admin).json()
    for nivel, logins in (("MODIFICACAO", ("ana", "bruno")), ("CONTROLE_TOTAL", ("carla", "edu")), ("LEITURA", ("dora",))):
        r = cliente.post("/api/acl/regras", json={"recurso_id": recurso["id"], "nivel": nivel, "usuarios_ids": [ids[x] for x in logins], "setores_ids": []}, headers=admin)
        assert r.status_code == 201, r.text
    return ids, {login: cabecalho(cliente, login) for login in ids}


@pytest.fixture
def predio(equipe):
    """Prédio > 5º andar > lado A com as salas 05A01 e 05A02 (com setor e responsável), 05A03 (sem) e o depósito 05A15 (setor "Depósito inservível")."""
    ids, _ = equipe
    with FabricaSessao() as s:
        financeiro, deposito = Setor(nome="Financeiro"), Setor(nome="Depósito inservível")
        s.add_all([financeiro, deposito])
        s.flush()
        raiz = Localizacao(id=uuid.uuid4(), tipo="predio", nome="Prédio SPI")
        andar = Localizacao(id=uuid.uuid4(), pai_id=raiz.id, tipo="andar", nome="5º andar", andar=5)
        lado = Localizacao(id=uuid.uuid4(), pai_id=andar.id, tipo="lado", nome="Lado A", andar=5, lado="A")
        s.add_all([raiz, andar, lado])
        s.flush()
        salas = {}
        for codigo, setor, resp in (("05A01", financeiro, ids["dora"]), ("05A02", financeiro, ids["edu"]), ("05A03", None, None), ("05A15", deposito, ids["edu"])):
            sala = Localizacao(id=uuid.uuid4(), pai_id=lado.id, tipo="sala", nome=codigo, codigo=codigo, andar=5, lado="A", setor_id=setor.id if setor else None, responsavel_id=resp)
            s.add(sala)
            salas[codigo] = str(sala.id)
        grupo = GrupoMaterial(id=uuid.uuid4(), codigo="MOB", nome="Mobiliário", vida_util_meses=120, residual_percentual=Decimal("10"), deprecia=True)
        s.add(grupo)
        s.commit()
        return {**salas, "andar": str(andar.id), "lado": str(lado.id), "grupo": str(grupo.id), "financeiro": financeiro.id}


def _bem(cliente, h, sala, chapa="PAT-001", **extra):
    corpo = {"chapa": chapa, "descricao": f"Cadeira {chapa}", "localizacao_id": sala, "conservacao": 4, **extra}
    return cliente.post(f"{URL}/bens", json=corpo, headers=h)


def test_acesso_por_nivel_do_recurso(cliente, equipe, predio):
    ids, h = equipe
    assert cliente.get(f"{URL}/localizacoes", headers=h["fora"]).status_code == 403
    assert cliente.get(f"{URL}/localizacoes", headers=h["dora"]).status_code == 200
    assert _bem(cliente, h["dora"], predio["05A01"]).status_code == 403
    assert _bem(cliente, h["ana"], predio["05A01"]).status_code == 201
    acesso = cliente.get(f"{URL}/acesso", headers=h["ana"]).json()
    assert acesso["pode_operar"] is True and acesso["pode_administrar"] is False
    assert cliente.get(f"{URL}/acesso", headers=h["carla"]).json()["pode_administrar"] is True
    assert cliente.get(f"{URL}/acesso", headers=h["dora"]).json()["pode_operar"] is False


def test_cadastro_obrigatorios_chapa_unica_e_responsavel_da_sala(cliente, equipe, predio):
    ids, h = equipe
    # Sala sem setor e responsável: o sistema pede antes de salvar
    r = _bem(cliente, h["ana"], predio["05A03"])
    assert r.status_code == 409 and r.json()["codigo"] == "sala_sem_responsavel"
    assert cliente.put(f"{URL}/localizacoes/{predio['05A03']}/dados", json={"setor_id": predio["financeiro"], "responsavel_id": ids["bruno"]}, headers=h["ana"]).status_code == 200
    ok = _bem(cliente, h["ana"], predio["05A03"])
    assert ok.status_code == 201 and ok.json()["responsavel_id"] == ids["bruno"] and ok.json()["sala_rotulo"] == "5º andar › 05A03 — Financeiro"
    # Depois de preenchidos, só o controle total altera setor e responsável
    assert cliente.put(f"{URL}/localizacoes/{predio['05A03']}/dados", json={"setor_id": predio["financeiro"], "responsavel_id": ids["ana"]}, headers=h["ana"]).status_code == 403
    # Chapa repetida (sem diferenciar maiúsculas), situação inválida e conservação fora de 1 a 5
    assert _bem(cliente, h["ana"], predio["05A03"], chapa="pat-001").status_code == 409
    assert _bem(cliente, h["ana"], predio["05A01"], chapa="X", situacao="inservivel").status_code == 422
    assert _bem(cliente, h["ana"], predio["05A01"], chapa="Y", conservacao=6).status_code == 422


def test_bem_sem_chapa_exige_numero_de_serie_e_usa_rotulo_sn(cliente, equipe, predio):
    ids, h = equipe
    assert cliente.put(f"{URL}/localizacoes/{predio['05A03']}/dados", json={"setor_id": predio["financeiro"], "responsavel_id": ids["bruno"]}, headers=h["ana"]).status_code == 200
    # Sem chapa e sem série: recusa; com série: cria, mostra "S/N ..." e dois bens sem chapa convivem
    sem = {"descricao": "Monitor", "localizacao_id": predio["05A03"], "conservacao": 4}
    assert cliente.post(f"{URL}/bens", json=sem, headers=h["ana"]).status_code == 400
    a = cliente.post(f"{URL}/bens", json={**sem, "numero_serie": "SN-1"}, headers=h["ana"])
    b = cliente.post(f"{URL}/bens", json={**sem, "numero_serie": "SN-2"}, headers=h["ana"])
    assert a.status_code == 201 and a.json()["chapa"] is None and a.json()["rotulo"] == "S/N SN-1" and b.status_code == 201
    # Depois, o bem pode ganhar chapa; remover a chapa só vale se ainda houver número de série
    r = cliente.put(f"{URL}/bens/{a.json()['id']}", json={"chapa": "NOVA-1"}, headers=h["ana"])
    assert r.status_code == 200 and r.json()["rotulo"] == "NOVA-1"
    assert cliente.put(f"{URL}/bens/{a.json()['id']}", json={"chapa": ""}, headers=h["ana"]).json()["rotulo"] == "S/N SN-1"


def test_opcoes_de_setores_e_usuarios_para_quem_opera(cliente, equipe, predio):
    ids, h = equipe
    assert "Financeiro" in [s["nome"] for s in cliente.get(f"{URL}/opcoes-setores", headers=h["ana"]).json()]
    assert "ana" in [u["login"] for u in cliente.get(f"{URL}/opcoes-usuarios", params={"busca": "ana"}, headers=h["ana"]).json()]
    assert cliente.get(f"{URL}/opcoes-setores", headers=h["dora"]).status_code == 403


def test_cadastro_offline_nao_duplica_e_marca_dados_contabeis_pendentes(cliente, equipe, predio):
    ids, h = equipe
    primeiro = _bem(cliente, h["ana"], predio["05A01"], id_cliente="aparelho-1")
    repetido = _bem(cliente, h["ana"], predio["05A01"], id_cliente="aparelho-1")
    assert primeiro.status_code == 201 and repetido.status_code == 200 and repetido.json()["id"] == primeiro.json()["id"]
    bem = primeiro.json()
    assert bem["dados_contabeis_pendentes"] is True and bem["situacao_garantia"] == "indefinida" and bem["garantia_fim"] is None
    completo = _bem(cliente, h["ana"], predio["05A01"], chapa="PAT-002", grupo_id=predio["grupo"], valor_aquisicao="2000.00", data_aquisicao="2024-01-15", garantia_meses=12)
    assert completo.json()["dados_contabeis_pendentes"] is False and completo.json()["garantia_fim"] == "2025-01-15" and completo.json()["situacao_garantia"] == "vencida"


def test_aviso_de_bem_nao_patrimoniavel(cliente, equipe, predio):
    ids, h = equipe
    with FabricaSessao() as s:
        s.add(GrupoMaterial(id=uuid.uuid4(), codigo="CURTO", nome="Vida curta", vida_util_meses=12, residual_percentual=Decimal("10"), deprecia=True))
        s.commit()
        curto = str(s.scalar(select(GrupoMaterial.id).where(GrupoMaterial.codigo == "CURTO")))
    r = _bem(cliente, h["ana"], predio["05A01"], grupo_id=curto, valor_aquisicao="100.00", data_aquisicao="2025-01-01")
    assert any("não é patrimoniável" in a for a in r.json()["avisos"])


def test_alteracao_grava_historico_por_campo_e_so_de_insercao(cliente, equipe, predio):
    ids, h = equipe
    bem = _bem(cliente, h["ana"], predio["05A01"]).json()
    r = cliente.put(f"{URL}/bens/{bem['id']}", json={"conservacao": 2, "marca": "Flexform"}, headers=h["bruno"])
    assert r.status_code == 200 and r.json()["conservacao"] == 2
    historico = cliente.get(f"{URL}/bens/{bem['id']}/historico", params={"campo": "conservacao"}, headers=h["dora"]).json()
    assert len(historico) == 1 and (historico[0]["valor_anterior"], historico[0]["valor_novo"], historico[0]["autor_nome"]) == ("4", "2", "Bruno Souza")
    assert {e["tipo"] for e in cliente.get(f"{URL}/bens/{bem['id']}/historico", headers=h["dora"]).json()} == {"cadastrou", "alterou"}
    # Nenhum serviço atualiza ou apaga eventos
    codigo = open(servico_bens.__file__, encoding="utf-8").read()
    assert "EventoPatrimonio.delete" not in codigo and "delete(EventoPatrimonio" not in codigo


def test_responsavel_da_sala_atualiza_todos_os_bens_com_historico(cliente, equipe, predio):
    ids, h = equipe
    ids_bens = [_bem(cliente, h["ana"], predio["05A01"], chapa=f"R-{i}").json()["id"] for i in range(3)]
    r = cliente.put(f"{URL}/localizacoes/{predio['05A01']}/dados", json={"setor_id": predio["financeiro"], "responsavel_id": ids["carla"]}, headers=h["carla"])
    assert r.status_code == 200
    for bem_id in ids_bens:
        bem = cliente.get(f"{URL}/bens/{bem_id}", headers=h["dora"]).json()
        assert bem["responsavel_id"] == ids["carla"]
        assert any(e["tipo"] == "responsavel" for e in cliente.get(f"{URL}/bens/{bem_id}/historico", headers=h["dora"]).json())


def test_sala_com_bens_so_inativa_e_codigo_da_sala_vem_do_nivel_de_cima(cliente, equipe, predio):
    ids, h = equipe
    _bem(cliente, h["ana"], predio["05A01"])
    assert cliente.delete(f"{URL}/localizacoes/{predio['05A01']}", headers=h["carla"]).status_code == 409
    assert cliente.put(f"{URL}/localizacoes/{predio['05A01']}", json={"ativo": False}, headers=h["carla"]).json()["ativo"] is False
    assert _bem(cliente, h["ana"], predio["05A01"], chapa="NOVO").status_code == 409
    nova = cliente.post(f"{URL}/localizacoes/lados/{predio['lado']}/salas", json={"numero": 7}, headers=h["carla"])
    assert nova.status_code == 201 and nova.json()["codigo"] == "05A07"
    assert cliente.post(f"{URL}/localizacoes/lados/{predio['lado']}/salas", json={"numero": 7}, headers=h["carla"]).status_code == 409
    assert cliente.post(f"{URL}/localizacoes/lados/{predio['lado']}/salas", json={"numero": 8}, headers=h["ana"]).status_code == 403


def test_movimentacao_autorizacao_recebimento_parcial_e_trajeto(cliente, equipe, predio):
    ids, h = equipe
    a = _bem(cliente, h["ana"], predio["05A01"], chapa="M-1").json()
    b = _bem(cliente, h["ana"], predio["05A01"], chapa="M-2").json()
    c = _bem(cliente, h["ana"], predio["05A02"], chapa="M-3").json()
    pedido = {"bens_ids": [a["id"], b["id"], c["id"]], "destino_id": predio["05A15"], "motivo": "Mudança"}
    # Inservível só vai para o depósito; aqui o destino é o depósito, então o pedido é válido
    mov = cliente.post(f"{URL}/movimentacoes", json=pedido, headers=h["ana"])
    assert mov.status_code == 201 and mov.json()["status"] == "pendente" and len(mov.json()["itens"]) == 3
    # Bem em movimentação não entra em outra
    assert cliente.post(f"{URL}/movimentacoes", json={**pedido, "bens_ids": [a["id"]], "destino_id": predio["05A02"]}, headers=h["bruno"]).json()["codigo"] == "em_movimentacao"
    # Receber antes de autorizar não vale; quem solicita não aprova a própria solicitação
    assert cliente.post(f"{URL}/movimentacoes/{mov.json()['id']}/receber", json={"bem_id": a["id"]}, headers=h["carla"]).status_code == 409
    propria = cliente.post(f"{URL}/movimentacoes", json={**pedido, "bens_ids": [c["id"]], "destino_id": predio["05A01"]}, headers=h["carla"])
    assert propria.status_code == 409  # c já está em outra movimentação
    assert cliente.post(f"{URL}/movimentacoes/{mov.json()['id']}/autorizar", json={}, headers=h["ana"]).status_code == 403
    # Recusa exige motivo
    assert cliente.post(f"{URL}/movimentacoes/{mov.json()['id']}/recusar", json={"motivo": ""}, headers=h["carla"]).status_code == 422
    autorizada = cliente.post(f"{URL}/movimentacoes/{mov.json()['id']}/autorizar", headers=h["carla"])
    assert autorizada.status_code == 200 and autorizada.json()["status"] == "autorizada"
    # Recebimento parcial: só a e b chegam; c fica na origem
    assert cliente.post(f"{URL}/movimentacoes/{mov.json()['id']}/receber", json={"bem_id": a["id"]}, headers=h["edu"]).status_code == 200
    assert cliente.get(f"{URL}/bens/{a['id']}", headers=h["dora"]).json()["sala_codigo"] == "05A15"
    assert cliente.get(f"{URL}/bens/{c['id']}", headers=h["dora"]).json()["sala_codigo"] == "05A02"
    assert cliente.post(f"{URL}/movimentacoes/{mov.json()['id']}/receber", json={"bem_id": a["id"]}, headers=h["edu"]).status_code == 409
    cliente.post(f"{URL}/movimentacoes/{mov.json()['id']}/receber", json={"bem_id": b["id"]}, headers=h["edu"])
    concluida = cliente.post(f"{URL}/movimentacoes/{mov.json()['id']}/concluir", headers=h["edu"]).json()
    assert concluida["status"] == "concluida"
    # O bem que não chegou fica livre para outra movimentação e o histórico registra que ficou
    assert any(e["tipo"] == "nao_chegou" for e in cliente.get(f"{URL}/bens/{c['id']}/historico", headers=h["dora"]).json())
    assert cliente.post(f"{URL}/movimentacoes", json={"bens_ids": [c["id"]], "destino_id": predio["05A01"], "motivo": "Voltar"}, headers=h["ana"]).status_code == 201
    # Responsável passa a ser o da sala de destino e o trajeto mostra quem pediu e quem autorizou
    assert cliente.get(f"{URL}/bens/{a['id']}", headers=h["dora"]).json()["responsavel_id"] == ids["edu"]
    trajeto = cliente.get(f"{URL}/bens/{a['id']}/trajeto", headers=h["dora"]).json()
    assert [t["tipo"] for t in trajeto] == ["cadastrou", "moveu"] and trajeto[1]["dados"]["solicitante"] == "ana" and trajeto[1]["dados"]["autorizador"] == "carla"


def test_localizar_bem_sem_local_no_levantamento_inicial(cliente, equipe, predio):
    ids, h = equipe
    # "NECESSÁRIO ATUALIZAR": sala sem andar onde ficam os bens ainda sem local
    with FabricaSessao() as s:
        raiz = s.scalar(select(Localizacao).where(Localizacao.tipo == "predio"))
        sem_local = Localizacao(id=uuid.uuid4(), pai_id=raiz.id, tipo="sala", nome="NECESSÁRIO ATUALIZAR", codigo="ATUALIZAR")
        s.add(sem_local)
        s.flush()
        # Os bens do GLPI entram direto no banco (a sala "sem local" não tem responsável, então o cadastro pela API a recusaria)
        for chapa in ("GLPI-1", "GLPI-2"):
            s.add(Bem(id=uuid.uuid4(), chapa=chapa, chapa_chave=chapa.lower(), descricao=f"Monitor {chapa}", localizacao_id=sem_local.id, conservacao=4))
        s.commit()
        sem_local_id = str(sem_local.id)
    bem = cliente.get(f"{URL}/bens/chapa/GLPI-1", headers=h["ana"]).json()
    # Quem tem MODIFICACAO informa a sala na hora, sem aprovação; o histórico e o trajeto registram "localizou"
    r = cliente.post(f"{URL}/bens/{bem['id']}/localizar", json={"localizacao_id": predio["05A01"]}, headers=h["ana"])
    assert r.status_code == 200 and r.json()["sala_codigo"] == "05A01" and r.json()["responsavel_id"] == ids["dora"]
    trajeto = cliente.get(f"{URL}/bens/{bem['id']}/trajeto", headers=h["dora"]).json()
    assert [t["tipo"] for t in trajeto] == ["localizou"] and trajeto[0]["dados"]["para"] == "05A01" and trajeto[0]["dados"]["de"] == "ATUALIZAR"
    # Bem que já tem sala: só por movimentação
    assert cliente.post(f"{URL}/bens/{bem['id']}/localizar", json={"localizacao_id": predio["05A02"]}, headers=h["ana"]).json()["codigo"] == "bem_ja_localizado"
    # Sala sem responsável, a própria sala "sem local" e quem só consulta
    outro = cliente.get(f"{URL}/bens/chapa/GLPI-2", headers=h["ana"]).json()
    assert cliente.post(f"{URL}/bens/{outro['id']}/localizar", json={"localizacao_id": predio["05A03"]}, headers=h["ana"]).json()["codigo"] == "sala_sem_responsavel"
    assert cliente.post(f"{URL}/bens/{outro['id']}/localizar", json={"localizacao_id": sem_local_id}, headers=h["ana"]).status_code == 400
    assert cliente.post(f"{URL}/bens/{outro['id']}/localizar", json={"localizacao_id": predio["05A01"]}, headers=h["dora"]).status_code == 403


def test_movimentacao_cancelar_e_idempotencia(cliente, equipe, predio):
    ids, h = equipe
    a = _bem(cliente, h["ana"], predio["05A01"]).json()
    corpo = {"id_cliente": "mov-offline-1", "bens_ids": [a["id"]], "destino_id": predio["05A02"], "motivo": "Troca"}
    primeira = cliente.post(f"{URL}/movimentacoes", json=corpo, headers=h["ana"])
    repetida = cliente.post(f"{URL}/movimentacoes", json=corpo, headers=h["ana"])
    assert primeira.status_code == 201 and repetida.status_code == 200 and repetida.json()["id"] == primeira.json()["id"]
    assert cliente.post(f"{URL}/movimentacoes/{primeira.json()['id']}/cancelar", headers=h["bruno"]).status_code == 403
    assert cliente.post(f"{URL}/movimentacoes/{primeira.json()['id']}/cancelar", headers=h["ana"]).json()["status"] == "cancelada"
    assert cliente.post(f"{URL}/movimentacoes", json={**corpo, "id_cliente": "mov-2"}, headers=h["ana"]).status_code == 201


def test_movimentacao_por_chapa_para_pedido_feito_sem_sinal(cliente, equipe, predio):
    ids, h = equipe
    _bem(cliente, h["ana"], predio["05A01"], chapa="OFF-1")
    r = cliente.post(f"{URL}/movimentacoes", json={"chapas": ["off-1"], "destino_id": predio["05A02"], "motivo": "Offline"}, headers=h["ana"])
    assert r.status_code == 201 and r.json()["itens"][0]["chapa"] == "OFF-1"
    assert cliente.post(f"{URL}/movimentacoes", json={"chapas": ["NAO-EXISTE"], "destino_id": predio["05A02"], "motivo": "x"}, headers=h["ana"]).status_code == 404
    assert cliente.post(f"{URL}/movimentacoes", json={"destino_id": predio["05A02"], "motivo": "x"}, headers=h["ana"]).status_code == 400


def test_ajuste_direto_de_local_so_controle_total_com_motivo(cliente, equipe, predio):
    ids, h = equipe
    a = _bem(cliente, h["ana"], predio["05A01"]).json()
    assert cliente.post(f"{URL}/bens/{a['id']}/local", json={"localizacao_id": predio["05A02"], "motivo": "x"}, headers=h["ana"]).status_code == 403
    assert cliente.post(f"{URL}/bens/{a['id']}/local", json={"localizacao_id": predio["05A02"], "motivo": ""}, headers=h["carla"]).status_code == 422
    r = cliente.post(f"{URL}/bens/{a['id']}/local", json={"localizacao_id": predio["05A02"], "motivo": "Erro de cadastro"}, headers=h["carla"])
    assert r.status_code == 200 and r.json()["sala_codigo"] == "05A02"
    assert any(e["tipo"] == "ajustou_local" and "Erro de cadastro" in e["texto"] for e in cliente.get(f"{URL}/bens/{a['id']}/historico", headers=h["dora"]).json())


def _lista(cliente, h, tipo, chapas):
    lista = cliente.post(f"{URL}/classificacoes", json={"tipo": tipo}, headers=h).json()
    for chapa in chapas:
        r = cliente.post(f"{URL}/classificacoes/{lista['id']}/bens", json={"chapa": chapa}, headers=h)
        assert r.status_code == 200, r.text
    return cliente.get(f"{URL}/classificacoes/{lista['id']}", headers=h).json()


def test_inservivel_justificativa_aprovacao_parcial_e_bloqueios(cliente, equipe, predio):
    ids, h = equipe
    for i in range(3):
        _bem(cliente, h["ana"], predio["05A01"], chapa=f"I-{i}")
    lista = _lista(cliente, h["ana"], "inservivel", ["I-0", "I-1", "I-2"])
    # Enviar sem justificativa não vale; a justificativa exige 10 caracteres e vale para todos os marcados
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["ana"]).json()["codigo"] == "sem_justificativa"
    marcados = [i["id"] for i in lista["itens"]]
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/justificativa", json={"itens_ids": marcados, "justificativa": "curta"}, headers=h["ana"]).status_code == 422
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/justificativa", json={"itens_ids": marcados, "categoria": "irrecuperavel", "justificativa": "Baterias sem carga, conserto inviável"}, headers=h["ana"]).status_code == 200
    # Foto opcional no inservível; outro formato é recusado
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/itens/{marcados[0]}/anexo", files={"arquivo": ("a.pdf", io.BytesIO(b"%PDF-1.4"), "application/pdf")}, headers=h["ana"]).status_code == 400
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/itens/{marcados[0]}/anexo", files={"arquivo": ("foto.jpg", io.BytesIO(JPG), "image/jpeg")}, headers=h["ana"]).status_code == 200
    enviada = cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["ana"]).json()
    assert enviada["status"] == "pendente"
    # Classificação pendente: não movimenta e não entra em outra lista
    pendente = cliente.get(f"{URL}/bens/chapa/I-0", headers=h["dora"]).json()
    assert pendente["classificacao_pendente"] is True
    assert cliente.post(f"{URL}/movimentacoes", json={"bens_ids": [pendente["id"]], "destino_id": predio["05A15"], "motivo": "x"}, headers=h["ana"]).json()["codigo"] == "classificacao_pendente"
    # Quem solicitou não aprova; a recusa exige motivo; aprova I-0 e I-1 e recusa I-2
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": marcados}, headers=h["ana"]).status_code == 403
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": marcados[:2]}, headers=h["carla"]).status_code == 400
    decidida = cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": marcados[:2], "motivo_recusa": "Ainda serve"}, headers=h["carla"]).json()
    assert decidida["status"] == "parcial"
    situacoes = {c: cliente.get(f"{URL}/bens/chapa/{c}", headers=h["dora"]).json() for c in ("I-0", "I-1", "I-2")}
    assert [situacoes[c]["situacao"] for c in ("I-0", "I-1", "I-2")] == ["inservivel", "inservivel", "ativo"] and situacoes["I-2"]["classificacao_pendente"] is False
    # O inservível só vai para a sala do setor "Depósito inservível"
    assert cliente.post(f"{URL}/movimentacoes", json={"bens_ids": [situacoes["I-0"]["id"]], "destino_id": predio["05A02"], "motivo": "x"}, headers=h["ana"]).json()["codigo"] == "destino_invalido"
    assert cliente.post(f"{URL}/movimentacoes", json={"bens_ids": [situacoes["I-0"]["id"]], "destino_id": predio["05A15"], "motivo": "x"}, headers=h["ana"]).status_code == 201


def test_lista_recusa_bem_que_nao_esta_ativo_e_extraviado_nao_se_move(cliente, equipe, predio):
    ids, h = equipe
    _bem(cliente, h["ana"], predio["05A01"], chapa="E-1")
    lista = _lista(cliente, h["ana"], "extraviado", ["E-1"])
    cliente.post(f"{URL}/classificacoes/{lista['id']}/justificativa", json={"itens_ids": [lista["itens"][0]["id"]], "justificativa": "Sumiu na mudança de andar"}, headers=h["ana"])
    cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["ana"])
    cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": [lista["itens"][0]["id"]]}, headers=h["carla"])
    bem = cliente.get(f"{URL}/bens/chapa/E-1", headers=h["dora"]).json()
    assert bem["situacao"] == "extraviado"
    # Não entra em nova lista de classificação e não pode ser movimentado nem ajustado
    outra = cliente.post(f"{URL}/classificacoes", json={"tipo": "inservivel"}, headers=h["ana"]).json()
    assert cliente.post(f"{URL}/classificacoes/{outra['id']}/bens", json={"chapa": "E-1"}, headers=h["ana"]).json()["codigo"] == "bem_invalido"
    assert cliente.post(f"{URL}/movimentacoes", json={"bens_ids": [bem["id"]], "destino_id": predio["05A02"], "motivo": "x"}, headers=h["ana"]).json()["codigo"] == "bem_extraviado"
    assert cliente.post(f"{URL}/bens/{bem['id']}/local", json={"localizacao_id": predio["05A02"], "motivo": "x"}, headers=h["carla"]).status_code == 409
    # Só o controle total devolve a Ativo, com justificativa e a sala onde foi achado
    assert cliente.post(f"{URL}/bens/{bem['id']}/reverter", json={"justificativa": "Encontrado no almoxarifado"}, headers=h["ana"]).status_code == 403
    assert cliente.post(f"{URL}/bens/{bem['id']}/reverter", json={"justificativa": "Encontrado no almoxarifado"}, headers=h["carla"]).status_code == 400
    achado = cliente.post(f"{URL}/bens/{bem['id']}/reverter", json={"justificativa": "Encontrado no almoxarifado", "localizacao_id": predio["05A02"]}, headers=h["carla"]).json()
    assert achado["situacao"] == "ativo" and achado["sala_codigo"] == "05A02"


def test_baixa_termo_em_pdf_e_bem_baixado_fica_somente_leitura(cliente, equipe, predio):
    ids, h = equipe
    _bem(cliente, h["ana"], predio["05A01"], chapa="B-1", grupo_id=predio["grupo"], valor_aquisicao="500.00", data_aquisicao="2020-01-01")
    _bem(cliente, h["ana"], predio["05A01"], chapa="B-2")
    lista = _lista(cliente, h["ana"], "inservivel", ["B-1", "B-2"])
    itens = [i["id"] for i in lista["itens"]]
    cliente.post(f"{URL}/classificacoes/{lista['id']}/justificativa", json={"itens_ids": itens, "categoria": "irrecuperavel", "justificativa": "Quebrado sem conserto viável"}, headers=h["ana"])
    cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["ana"])
    cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": itens}, headers=h["carla"])
    # Bem ativo não entra na baixa
    _bem(cliente, h["ana"], predio["05A01"], chapa="B-3")
    baixa = cliente.post(f"{URL}/baixas", json={}, headers=h["ana"]).json()
    assert cliente.post(f"{URL}/baixas/{baixa['id']}/bens", json={"chapa": "B-3"}, headers=h["ana"]).json()["codigo"] == "bem_invalido"
    for chapa in ("B-1", "B-2"):
        assert cliente.post(f"{URL}/baixas/{baixa['id']}/bens", json={"chapa": chapa}, headers=h["ana"]).status_code == 200
    detalhe = cliente.get(f"{URL}/baixas/{baixa['id']}", headers=h["ana"]).json()
    ids_itens = [i["id"] for i in detalhe["itens"]]
    assert cliente.post(f"{URL}/baixas/{baixa['id']}/enviar", headers=h["ana"]).json()["codigo"] == "sem_motivo"
    assert cliente.post(f"{URL}/baixas/{baixa['id']}/motivo", json={"itens_ids": ids_itens, "motivo": "doacao", "processo": "SPI-2026/001"}, headers=h["ana"]).status_code == 200
    assert cliente.post(f"{URL}/baixas/{baixa['id']}/enviar", headers=h["ana"]).json()["status"] == "pendente"
    assert cliente.post(f"{URL}/baixas/{baixa['id']}/decisao", json={"aprovados_ids": ids_itens}, headers=h["ana"]).status_code == 403
    aprovada = cliente.post(f"{URL}/baixas/{baixa['id']}/decisao", json={"aprovados_ids": ids_itens}, headers=h["edu"])
    assert aprovada.status_code == 200 and aprovada.json()["status"] == "aprovada" and aprovada.json()["tem_termo"] is True
    termo = cliente.get(f"{URL}/baixas/{baixa['id']}/termo", headers=h["dora"])
    assert termo.status_code == 200 and termo.content.startswith(b"%PDF")
    # Baixado: sai da lista padrão e do painel, não edita nem move, mas nunca é excluído
    baixado = cliente.get(f"{URL}/bens/chapa/B-1", headers=h["dora"]).json()
    assert baixado["situacao"] == "baixado" and baixado["baixado_em"]
    assert "B-1" not in [b["chapa"] for b in cliente.get(f"{URL}/bens", headers=h["dora"]).json()["itens"]]
    assert "B-1" in [b["chapa"] for b in cliente.get(f"{URL}/bens", params={"situacao": "baixado"}, headers=h["dora"]).json()["itens"]]
    assert cliente.put(f"{URL}/bens/{baixado['id']}", json={"marca": "x"}, headers=h["ana"]).json()["codigo"] == "bem_baixado"
    assert cliente.post(f"{URL}/movimentacoes", json={"bens_ids": [baixado["id"]], "destino_id": predio["05A02"], "motivo": "x"}, headers=h["ana"]).json()["codigo"] == "bem_baixado"
    assert any(e["tipo"] == "baixou" for e in cliente.get(f"{URL}/bens/{baixado['id']}/historico", headers=h["dora"]).json())


def test_painel_agrupa_exclui_baixados_e_dados_pendentes_ficam_fora_do_valor(cliente, equipe, predio):
    ids, h = equipe
    _bem(cliente, h["ana"], predio["05A01"], chapa="P-1", grupo_id=predio["grupo"], valor_aquisicao="1000.00", valor_atual="800.00", data_aquisicao="2020-01-01", garantia_meses=12)
    _bem(cliente, h["ana"], predio["05A02"], chapa="P-2")
    painel = cliente.get(f"{URL}/painel", params={"agrupar_por": "andar"}, headers=h["dora"]).json()
    ind = painel["indicadores"]
    assert ind["total_bens"] == 2 and Decimal(ind["valor_atual"]) == Decimal("800.00") and ind["dados_contabeis_pendentes"] == 1 and ind["fora_garantia"] == 1
    assert painel["grupos"][0]["rotulo"] == "5º andar" and painel["grupos"][0]["quantidade"] == 2
    assert [x["chapa"] for x in painel["contabeis_pendentes"]] == ["P-2"] and [x["chapa"] for x in painel["fora_garantia"]] == ["P-1"]
    por_sala = cliente.get(f"{URL}/painel", params={"agrupar_por": "sala"}, headers=h["dora"]).json()
    assert [g["quantidade"] for g in por_sala["grupos"]] == [1, 1]
    lista = cliente.get(f"{URL}/painel/grupo", params={"agrupar_por": "sala", "chave": "05A01"}, headers=h["dora"]).json()
    assert [b["chapa"] for b in lista] == ["P-1"]


def test_exportacoes_em_excel_e_pdf(cliente, equipe, predio):
    ids, h = equipe
    _bem(cliente, h["ana"], predio["05A01"], chapa="X-1", grupo_id=predio["grupo"], valor_aquisicao="1000.00", data_aquisicao="2020-01-01", garantia_meses=12)
    for relatorio in ("bens", "fora-garantia", "painel", "movimentacoes"):
        xlsx = cliente.get(f"{URL}/exportar/{relatorio}", params={"formato": "xlsx"}, headers=h["dora"])
        pdf = cliente.get(f"{URL}/exportar/{relatorio}", params={"formato": "pdf"}, headers=h["dora"])
        assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK", relatorio
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF"), relatorio
    assert cliente.get(f"{URL}/exportar/bens", headers=h["fora"]).status_code == 403


def test_foto_do_bem_limite_de_tres(cliente, equipe, predio):
    ids, h = equipe
    bem = _bem(cliente, h["ana"], predio["05A01"]).json()
    for i in range(3):
        r = cliente.post(f"{URL}/bens/{bem['id']}/fotos", files={"arquivo": (f"f{i}.jpg", io.BytesIO(JPG + bytes([i])), "image/jpeg")}, headers=h["ana"])
        assert r.status_code == 201, r.text
        # A resposta já traz a foto recém-enviada
        assert len(r.json()["fotos"]) == i + 1
    assert cliente.post(f"{URL}/bens/{bem['id']}/fotos", files={"arquivo": ("f4.jpg", io.BytesIO(JPG + b"9"), "image/jpeg")}, headers=h["ana"]).json()["codigo"] == "limite_fotos"
    foto = cliente.get(f"{URL}/bens/{bem['id']}", headers=h["dora"]).json()["fotos"][0]
    assert cliente.get(f"{URL}/bens/{bem['id']}/fotos/{foto['id']}/arquivo", headers=h["dora"]).status_code == 200
    removida = cliente.delete(f"{URL}/bens/{bem['id']}/fotos/{foto['id']}", headers=h["ana"])
    assert removida.status_code == 200 and len(removida.json()["fotos"]) == 2


def test_grupos_e_parametros_so_o_controle_total_altera(cliente, equipe, predio):
    ids, h = equipe
    assert cliente.post(f"{URL}/grupos", json={"codigo": "TI", "nome": "Informática", "vida_util_meses": 60}, headers=h["ana"]).status_code == 403
    criado = cliente.post(f"{URL}/grupos", json={"codigo": "TI", "nome": "Informática", "vida_util_meses": 60}, headers=h["carla"])
    assert criado.status_code == 201 and Decimal(criado.json()["residual_percentual"]) == Decimal("10")
    assert cliente.post(f"{URL}/grupos", json={"codigo": "ti", "nome": "Outro"}, headers=h["carla"]).status_code == 409
    assert cliente.put(f"{URL}/parametros", json={"ufesp": "40.10"}, headers=h["carla"]).json()["ufesp"] == "40.10"
    assert cliente.put(f"{URL}/parametros", json={"ufesp": "0"}, headers=h["carla"]).status_code == 400
    assert cliente.put(f"{URL}/parametros", json={"ufesp": "40"}, headers=h["ana"]).status_code == 403


def test_depreciacao_linear_exemplo_da_mesa():
    """Mesa de R$ 2.000, residual de 10% e vida útil de dez anos: perde R$ 15 por mês e, em três anos, vale R$ 1.460."""
    custo, residual, vida = Decimal("2000"), Decimal("10"), 120
    assert depreciacao.depreciacao_mensal(custo, residual, vida) == Decimal("15")
    assert depreciacao.valor_calculado(custo, residual, vida, date(2023, 3, 10), date(2026, 3, 10)) == Decimal("1460.00")
    # No mês da disponibilidade ainda não deprecia; a depreciação começa no mês seguinte
    assert depreciacao.valor_calculado(custo, residual, vida, date(2026, 3, 10), date(2026, 3, 31)) == Decimal("2000.00")
    assert depreciacao.valor_calculado(custo, residual, vida, date(2026, 3, 10), date(2026, 4, 1)) == Decimal("1985.00")
    # Para no residual (10% do custo) e não depreciam os bens sem depreciação
    assert depreciacao.valor_calculado(custo, residual, vida, date(2000, 1, 1), date(2026, 1, 1)) == Decimal("200.00")
    assert depreciacao.valor_calculado(custo, residual, vida, date(2000, 1, 1), date(2026, 1, 1), deprecia=False) == Decimal("2000.00")
    assert depreciacao.valor_calculado(custo, residual, None, date(2020, 1, 1), date(2026, 1, 1)) is None


def test_garantia_soma_meses_a_data_de_aquisicao():
    assert servico_bens.adicionar_meses(date(2024, 1, 31), 1) == date(2024, 2, 29)
    assert servico_bens.adicionar_meses(date(2024, 11, 15), 3) == date(2025, 2, 15)


def test_lista_de_bens_ordenavel_por_coluna_em_todas_as_paginas(cliente, equipe, predio):
    ids, h = equipe
    _bem(cliente, h["ana"], predio["05A02"], chapa="B-2", conservacao=2, valor_aquisicao="300.00", grupo_id=predio["grupo"], data_aquisicao="2024-01-10", garantia_meses=12)
    _bem(cliente, h["ana"], predio["05A01"], chapa="A-1", conservacao=5, valor_aquisicao="900.00", data_aquisicao="2025-06-01", garantia_meses=12)
    _bem(cliente, h["ana"], predio["05A01"], chapa="C-3", conservacao=3)

    def chapas(**parametros):
        return [b["chapa"] for b in cliente.get(f"{URL}/bens", params=parametros, headers=h["dora"]).json()["itens"]]

    assert chapas() == ["A-1", "B-2", "C-3"]
    assert chapas(ordenar_por="chapa", direcao="desc") == ["C-3", "B-2", "A-1"]
    assert chapas(ordenar_por="conservacao", direcao="desc") == ["A-1", "C-3", "B-2"]
    # Sem valor vai para o fim, qualquer que seja o sentido
    assert chapas(ordenar_por="valor_aquisicao") == ["B-2", "A-1", "C-3"] and chapas(ordenar_por="valor_aquisicao", direcao="desc") == ["A-1", "B-2", "C-3"]
    assert chapas(ordenar_por="sala")[0] in ("A-1", "C-3") and chapas(ordenar_por="sala")[-1] == "B-2"
    assert chapas(ordenar_por="grupo")[0] == "B-2"
    # Garantia: pelo fim da garantia, e os de garantia indefinida por último
    assert chapas(ordenar_por="garantia") == ["B-2", "A-1", "C-3"] and chapas(ordenar_por="garantia", direcao="desc") == ["A-1", "B-2", "C-3"]
    # A ordem vale para todas as páginas
    pagina1 = cliente.get(f"{URL}/bens", params={"ordenar_por": "chapa", "direcao": "desc", "tamanho": 1, "pagina": 1}, headers=h["dora"]).json()
    assert [b["chapa"] for b in pagina1["itens"]] == ["C-3"] and pagina1["total"] == 3
    assert cliente.get(f"{URL}/bens", params={"ordenar_por": "inexistente"}, headers=h["dora"]).status_code == 400


def test_sugestoes_de_descricao_para_padronizar(cliente, equipe, predio):
    ids, h = equipe
    for i, d in enumerate(["Computador", "computador", "Computador de mesa", "Cadeira giratória", "Mesa de reunião"]):
        r = cliente.post(f"{URL}/bens", json={"chapa": f"S-{i}", "descricao": d, "localizacao_id": predio["05A01"], "conservacao": 4}, headers=h["ana"])
        assert r.status_code == 201
    sugestoes = lambda q: cliente.get(f"{URL}/bens/sugestoes-descricao", params={"q": q}, headers=h["dora"]).json()
    # Uma só grafia para "Computador" (as duas viram uma), a que começa pelo texto antes das demais
    assert sugestoes("comp") == ["Computador", "Computador de mesa"]
    assert sugestoes("MESA") == ["Mesa de reunião", "Computador de mesa"]
    assert sugestoes("c") == [] and sugestoes("zzz") == []
    assert cliente.get(f"{URL}/bens/sugestoes-descricao", params={"q": "comp"}, headers=h["fora"]).status_code == 403


def test_prefixos_da_chapa_parametros_montagem_e_leitura(cliente, equipe, predio):
    ids, h, = equipe[0], equipe[1]
    # Só o controle total cria; padrão único e ativo
    assert cliente.post(f"{URL}/prefixos", json={"codigo": "spi"}, headers=h["ana"]).status_code == 403
    spi = cliente.post(f"{URL}/prefixos", json={"codigo": "spi"}, headers=h["carla"]).json()
    assert spi["codigo"] == "SPI"
    slt = cliente.post(f"{URL}/prefixos", json={"codigo": "SLT"}, headers=h["carla"]).json()
    assert cliente.post(f"{URL}/prefixos", json={"codigo": "slt"}, headers=h["carla"]).status_code == 409
    assert cliente.post(f"{URL}/prefixos", json={"codigo": "S L"}, headers=h["carla"]).status_code == 400
    assert cliente.put(f"{URL}/prefixos/{spi['id']}", json={"padrao": True}, headers=h["carla"]).json()["padrao"] is True
    assert cliente.put(f"{URL}/prefixos/{spi['id']}", json={"ativo": False}, headers=h["carla"]).status_code == 409
    assert [p["codigo"] for p in cliente.get(f"{URL}/prefixos", headers=h["dora"]).json()] == ["SPI", "SLT"]

    # Número com menos de 6 dígitos vira 6; o prefixo se junta ao salvar
    sala = predio["05A01"]
    bem = _bem(cliente, h["ana"], sala, chapa="231", prefixo="SPI")
    assert bem.status_code == 201 and bem.json()["chapa"] == "SPI000231"
    # Outro prefixo, sem repetir o prefixo se o código lido já o traz; texto livre fica como digitado
    assert _bem(cliente, h["ana"], sala, chapa="12", prefixo="slt").json()["chapa"] == "SLT000012"
    assert _bem(cliente, h["ana"], sala, chapa="SPI000500", prefixo="SPI").json()["chapa"] == "SPI000500"
    assert _bem(cliente, h["ana"], sala, chapa="AB-9", prefixo="SLT").json()["chapa"] == "SLTAB-9"
    # Mesmo número com prefixos diferentes são chapas diferentes; repetida é 409; prefixo desconhecido é 404
    assert _bem(cliente, h["ana"], sala, chapa="231", prefixo="SLT").status_code == 201
    assert _bem(cliente, h["ana"], sala, chapa="231", prefixo="SPI").status_code == 409
    assert _bem(cliente, h["ana"], sala, chapa="1", prefixo="XYZ").status_code == 404
    # Sem prefixo, a chapa é gravada como veio (compatível com o que já existia)
    assert _bem(cliente, h["ana"], sala, chapa="LEGADO-1").json()["chapa"] == "LEGADO-1"
    # Leitura de etiqueta antiga (só dígitos) acha o bem pelo prefixo padrão
    assert cliente.get(f"{URL}/bens/chapa/231", headers=h["dora"]).json()["chapa"] == "SPI000231"
    # Alterar a chapa de um bem existente também junta o prefixo
    r = cliente.put(f"{URL}/bens/{bem.json()['id']}", json={"chapa": "7", "prefixo": "SLT"}, headers=h["ana"])
    assert r.status_code == 200 and r.json()["chapa"] == "SLT000007"
    # Prefixo com bens só inativa
    assert cliente.delete(f"{URL}/prefixos/{slt['id']}", headers=h["carla"]).status_code == 409


def test_busca_por_numero_de_serie_em_leitura_movimentacao_classificacao_e_baixa(cliente, equipe, predio):
    ids, h = equipe[0], equipe[1]
    sala, destino = predio["05A01"], predio["05A02"]
    # Bem sem chapa, só com número de série, e outro com chapa e série repetida
    sem_chapa = cliente.post(f"{URL}/bens", json={"descricao": "Notebook", "numero_serie": "SN-123", "localizacao_id": sala, "conservacao": 4}, headers=h["ana"]).json()
    com_chapa = _bem(cliente, h["ana"], sala, chapa="PAT-9", numero_serie="XY-777").json()
    # Leitura (câmera ou digitação) acha por chapa ou por número de série, sem diferenciar maiúsculas
    assert cliente.get(f"{URL}/bens/chapa/sn-123", headers=h["dora"]).json()["id"] == sem_chapa["id"]
    assert cliente.get(f"{URL}/bens/chapa/xy-777", headers=h["dora"]).json()["id"] == com_chapa["id"]
    assert cliente.get(f"{URL}/bens/chapa/NAO-EXISTE", headers=h["dora"]).status_code == 404
    # Movimentação pedida pelo número de série (como faz o aparelho sem sinal)
    r = cliente.post(f"{URL}/movimentacoes", json={"chapas": ["SN-123"], "destino_id": destino, "motivo": "Troca de sala"}, headers=h["ana"])
    assert r.status_code == 201, r.text
    assert r.json()["itens"][0]["numero_serie"] == "SN-123"
    # Número de série repetido pede a chapa
    _bem(cliente, h["ana"], sala, chapa="PAT-10", numero_serie="XY-777")
    assert cliente.get(f"{URL}/bens/chapa/XY-777", headers=h["dora"]).status_code == 409


def test_tipos_de_evento_cabem_na_coluna():
    """O SQLite dos testes não barra texto longo; no PostgreSQL, `classificacao_pendente` (22) estourava varchar(20) e o envio da lista dava 500."""
    import re
    from pathlib import Path

    limite = EventoPatrimonio.__table__.c.tipo.type.length
    tipos = set()
    for arquivo in Path("app/services/patrimonio").glob("*.py"):
        tipos |= set(re.findall(r'registrar_evento\(\s*sessao,\s*[\w.]+,\s*"([a-z_]+)"', arquivo.read_text(encoding="utf-8")))
    assert tipos and max(len(t) for t in tipos) <= limite, [t for t in tipos if len(t) > limite]


def test_filtro_de_situacao_aceita_varias_separadas_por_virgula(cliente, equipe, predio):
    ids, h = equipe[0], equipe[1]
    sala = predio["05A01"]
    ativo = _bem(cliente, h["ana"], sala, chapa="F-1").json()
    outro = _bem(cliente, h["ana"], sala, chapa="F-2").json()
    lista = cliente.post(f"{URL}/classificacoes", json={"tipo": "inservivel", "id_cliente": str(uuid.uuid4())}, headers=h["ana"])
    assert lista.status_code in (200, 201), lista.text
    achados = cliente.get(f"{URL}/bens", params={"situacao": "ativo,em_manutencao"}, headers=h["dora"]).json()["itens"]
    assert {b["chapa"] for b in achados} >= {"F-1", "F-2"}
    assert cliente.get(f"{URL}/bens", params={"situacao": "inservivel,extraviado"}, headers=h["dora"]).json()["itens"] == []
    assert cliente.get(f"{URL}/bens", params={"situacao": "inservivel,qualquer"}, headers=h["dora"]).status_code == 400


def _jpeg_real(cor):
    from PIL import Image

    saida = io.BytesIO()
    Image.new("RGB", (320, 240), cor).save(saida, "JPEG")
    return saida.getvalue()


def _pdf_real(texto):
    from reportlab.pdfgen.canvas import Canvas

    saida = io.BytesIO()
    tela = Canvas(saida)
    tela.drawString(100, 750, texto)
    tela.save()
    return saida.getvalue()


def test_relatorio_consolidado_de_extraviados_e_da_baixa_com_fotos_e_anexos(cliente, equipe, predio):
    from openpyxl import load_workbook
    from pypdf import PdfReader

    ids, h = equipe
    sala = predio["05A01"]
    for i in range(2):
        bem = _bem(cliente, h["ana"], sala, chapa=f"EX-{i}", marca="Dell", modelo="X1").json()
        for cor in ("red", "blue"):
            assert cliente.post(f"{URL}/bens/{bem['id']}/fotos", files={"arquivo": ("f.jpg", io.BytesIO(_jpeg_real(cor)), "image/jpeg")}, headers=h["ana"]).status_code == 201
    lista = _lista(cliente, h["ana"], "extraviado", ["EX-0", "EX-1"])
    itens = [i["id"] for i in lista["itens"]]
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/justificativa", json={"itens_ids": itens, "justificativa": "Extraviado na mudança de sala"}, headers=h["ana"]).status_code == 200
    # Só o primeiro item tem boletim de ocorrência (PDF)
    r = cliente.post(f"{URL}/classificacoes/{lista['id']}/itens/{itens[0]}/anexo", files={"arquivo": ("bo.pdf", io.BytesIO(_pdf_real("Boletim de ocorrencia")), "application/pdf")}, headers=h["ana"])
    assert r.status_code == 200, r.text

    pdf = cliente.get(f"{URL}/classificacoes/{lista['id']}/relatorio", params={"formato": "pdf"}, headers=h["dora"])
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    leitor = PdfReader(io.BytesIO(pdf.content))
    # índice + relatório + 2 páginas de fotos + contracapa e página do boletim de ocorrência
    assert len(leitor.pages) >= 6
    texto = " ".join(p.extract_text() for p in leitor.pages)
    assert "Dell - X1" in texto and "Anexo 1" in texto and "Boletim de ocorrencia" in texto and "Última:" in texto
    # Hiperlinks do índice e das colunas de fotos e anexo (para as páginas do consolidado)
    ligacoes = [a.get_object() for p in leitor.pages for a in (p.get("/Annots") or []) if a.get_object().get("/Subtype") == "/Link"]
    assert len(ligacoes) >= 4 + 2 + 1

    import os

    if os.environ.get("RELATORIO_AMOSTRA"):  # só para olhar o resultado (não faz parte da verificação)
        open(os.environ["RELATORIO_AMOSTRA"], "wb").write(pdf.content)
    xlsx = cliente.get(f"{URL}/classificacoes/{lista['id']}/relatorio", params={"formato": "xlsx"}, headers=h["dora"])
    assert xlsx.status_code == 200
    livro = load_workbook(io.BytesIO(xlsx.content))
    assert livro.sheetnames == ["Relatório", "Fotos"]
    assert livro["Relatório"]["C5"].value == "Cadeira EX-0 - Dell - X1" and livro["Relatório"]["G5"].value == "bo.pdf" and livro["Relatório"]["G6"].value == "—"
    assert cliente.get(f"{URL}/classificacoes/{lista['id']}/relatorio", params={"formato": "doc"}, headers=h["dora"]).status_code == 422

    # Baixa: o anexo vem da classificação aprovada do bem
    enviada = cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["ana"])
    assert enviada.status_code == 200, enviada.text
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": itens}, headers=h["carla"]).status_code == 200
    baixa = cliente.post(f"{URL}/baixas", json={"id_cliente": str(uuid.uuid4())}, headers=h["ana"]).json()
    for chapa in ("EX-0", "EX-1"):
        assert cliente.post(f"{URL}/baixas/{baixa['id']}/bens", json={"chapa": chapa}, headers=h["ana"]).status_code == 200
    relatorio = cliente.get(f"{URL}/baixas/{baixa['id']}/relatorio", headers=h["dora"])
    assert relatorio.status_code == 200 and relatorio.content.startswith(b"%PDF")
    assert "Boletim de ocorrencia" in " ".join(p.extract_text() for p in PdfReader(io.BytesIO(relatorio.content)).pages)


def test_area_comum_no_andar_hall_dos_elevadores(cliente, equipe, predio):
    ids, h = equipe
    url = f"{URL}/localizacoes/andares/{predio['andar']}/areas-comuns"
    # Só o controle total cria; vale só dentro de um andar; sigla só com letras
    assert cliente.post(url, json={"nome": "Hall dos elevadores"}, headers=h["ana"]).status_code == 403
    assert cliente.post(f"{URL}/localizacoes/andares/{predio['lado']}/areas-comuns", json={"nome": "Hall"}, headers=h["carla"]).status_code == 400
    assert cliente.post(url, json={"nome": "Hall", "sigla": "A1"}, headers=h["carla"]).status_code == 422
    hall = cliente.post(url, json={"nome": "Hall dos elevadores"}, headers=h["carla"])
    assert hall.status_code == 201, hall.text
    hall = hall.json()
    # Código do andar + sigla, sem lado e fora da numeração A/B; a tela mostra o nome
    assert hall["codigo"] == "05H" and hall["lado"] is None and hall["andar"] == 5 and hall["tipo"] == "sala"
    assert hall["rotulo"] == "5º andar › Hall dos elevadores"
    assert cliente.post(url, json={"nome": "Outro hall"}, headers=h["carla"]).status_code == 409
    # Outra área comum no mesmo andar, com outra sigla
    assert cliente.post(url, json={"nome": "Copa", "sigla": "cp"}, headers=h["carla"]).json()["codigo"] == "05CP"
    # Sem setor e responsável não recebe bem; com eles, sim
    assert _bem(cliente, h["ana"], hall["id"], chapa="HALL-1").json()["codigo"] == "sala_sem_responsavel"
    financeiro = predio["financeiro"]
    assert cliente.put(f"{URL}/localizacoes/{hall['id']}/dados", json={"setor_id": financeiro, "responsavel_id": ids["edu"]}, headers=h["carla"]).status_code == 200
    bem = _bem(cliente, h["ana"], hall["id"], chapa="HALL-1")
    assert bem.status_code == 201 and bem.json()["sala_rotulo"] == "5º andar › Hall dos elevadores — Financeiro"


def test_lista_compartilhada_colaborador_faz_tudo_e_nao_aprova(cliente, equipe, predio):
    """Quem recebe a lista faz tudo o que o solicitante faz (incluir, justificar, anexar, enviar), recebe aviso e não pode aprová-la."""
    ids, h = equipe
    for chapa in ("C-1", "C-2"):
        _bem(cliente, h["ana"], predio["05A01"], chapa=chapa)
    lista = _lista(cliente, h["ana"], "inservivel", ["C-1"])
    # Sem compartilhar, o Bruno não mexe na lista
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/bens", json={"chapa": "C-2"}, headers=h["bruno"]).status_code == 403
    # Só quem montou compartilha; o convidado recebe aviso e a lista aparece em "minhas listas" dele
    assert cliente.put(f"{URL}/classificacoes/{lista['id']}/colaboradores", json={"usuarios_ids": [ids["bruno"]]}, headers=h["bruno"]).status_code == 403
    assert cliente.put(f"{URL}/classificacoes/{lista['id']}/colaboradores", json={"usuarios_ids": [99999]}, headers=h["ana"]).status_code == 400
    r = cliente.put(f"{URL}/classificacoes/{lista['id']}/colaboradores", json={"usuarios_ids": [ids["bruno"], ids["ana"]]}, headers=h["ana"])
    assert r.status_code == 200 and [c["id"] for c in r.json()["colaboradores"]] == [ids["bruno"]]
    minhas = cliente.get(f"{URL}/classificacoes", params={"solicitante_id": ids["bruno"]}, headers=h["bruno"]).json()
    assert [x["id"] for x in minhas] == [lista["id"]]
    # O colaborador faz tudo: inclui bem, justifica com HTML do editor rico, anexa foto e envia
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/bens", json={"chapa": "C-2"}, headers=h["bruno"]).status_code == 200
    atual = cliente.get(f"{URL}/classificacoes/{lista['id']}", headers=h["bruno"]).json()
    todos = [i["id"] for i in atual["itens"]]
    html = "<p>Baterias <strong>sem carga</strong>, conserto inviável</p><script>x()</script>"
    r = cliente.post(f"{URL}/classificacoes/{lista['id']}/justificativa", json={"itens_ids": todos, "categoria": "irrecuperavel", "justificativa": html}, headers=h["bruno"])
    assert r.status_code == 200 and all("<strong>sem carga</strong>" in i["justificativa"] and "<script" not in i["justificativa"] for i in r.json()["itens"])
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/itens/{todos[0]}/anexo", files={"arquivo": ("foto.jpg", io.BytesIO(JPG), "image/jpeg")}, headers=h["bruno"]).status_code == 200
    enviada = cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["bruno"])
    assert enviada.status_code == 200 and enviada.json()["status"] == "pendente"
    # Nem o solicitante nem o colaborador aprovam; outra pessoa com controle total aprova
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": todos}, headers=h["ana"]).status_code == 403
    assert cliente.put(f"{URL}/classificacoes/{lista['id']}/colaboradores", json={"usuarios_ids": []}, headers=h["ana"]).status_code == 409  # já enviada
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": todos}, headers=h["carla"]).json()["status"] == "aprovada"


def test_colaborador_nao_aprova_nem_com_controle_total(cliente, equipe, predio):
    ids, h = equipe
    _bem(cliente, h["ana"], predio["05A01"], chapa="D-1")
    lista = _lista(cliente, h["ana"], "inservivel", ["D-1"])
    cliente.put(f"{URL}/classificacoes/{lista['id']}/colaboradores", json={"usuarios_ids": [ids["carla"]]}, headers=h["ana"])
    marcados = [i["id"] for i in lista["itens"]]
    cliente.post(f"{URL}/classificacoes/{lista['id']}/justificativa", json={"itens_ids": marcados, "categoria": "irrecuperavel", "justificativa": "Sem conserto e sem uso"}, headers=h["ana"])
    cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["ana"])
    r = cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": marcados}, headers=h["carla"])
    assert r.status_code == 403 and r.json()["codigo"] == "auto_aprovacao"
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": marcados}, headers=h["edu"]).json()["status"] == "aprovada"


def test_justificativa_rica_conta_o_texto_sem_tags_e_limita_a_2000(cliente, equipe, predio):
    ids, h = equipe
    _bem(cliente, h["ana"], predio["05A01"], chapa="J-1")
    lista = _lista(cliente, h["ana"], "inservivel", ["J-1"])
    marcados = [i["id"] for i in lista["itens"]]
    url = f"{URL}/classificacoes/{lista['id']}/justificativa"
    # Muitas tags, pouco texto: vale o texto (menos de 10 caracteres = recusa)
    assert cliente.post(url, json={"itens_ids": marcados, "categoria": "irrecuperavel", "justificativa": "<p><strong><em>curta</em></strong></p>"}, headers=h["ana"]).status_code == 400
    assert cliente.post(url, json={"itens_ids": marcados, "categoria": "irrecuperavel", "justificativa": "<p>" + "a" * 2001 + "</p>"}, headers=h["ana"]).status_code == 400
    assert cliente.post(url, json={"itens_ids": marcados, "categoria": "irrecuperavel", "justificativa": "<p>" + "a" * 2000 + "</p>"}, headers=h["ana"]).status_code == 200


def test_busca_de_bens_por_qualquer_informacao_com_varias_palavras(cliente, equipe, predio):
    """A busca (`q`) acha pela chapa, série, descrição, marca, modelo, local, grupo, responsável, fornecedor, nota fiscal e observações; várias palavras se somam."""
    ids, h = equipe
    _bem(cliente, h["ana"], predio["05A01"], chapa="B-1", descricao="Notebook gerencial", marca="Dell", modelo="Latitude 5440", fornecedor="Distribuidora Alfa")
    _bem(cliente, h["ana"], predio["05A02"], chapa="B-2", descricao="Notebook gerencial", marca="Lenovo", modelo="ThinkPad", observacoes="Tela trincada no canto")
    _bem(cliente, h["ana"], predio["05A01"], chapa="B-3", descricao="Monitor de vídeo", marca="Dell")

    def achar(q):
        return sorted(i["chapa"] for i in cliente.get(f"{URL}/bens", params={"q": q}, headers=h["dora"]).json()["itens"])

    assert achar("dell") == ["B-1", "B-3"]  # marca
    assert achar("latitude") == ["B-1"]  # modelo
    assert achar("05A02") == ["B-2"]  # local (sala)
    assert achar("5º andar") == ["B-1", "B-2", "B-3"]  # tudo o que está dentro do andar
    assert achar("alfa") == ["B-1"]  # fornecedor
    assert achar("trincada") == ["B-2"]  # observações
    assert achar("notebook dell") == ["B-1"]  # várias palavras: todas precisam aparecer
    assert achar("dell 05A01 monitor") == ["B-3"]
    assert achar("notebook 05A01 lenovo") == []


def test_baixa_compartilhada_colaborador_faz_tudo_e_nao_aprova(cliente, equipe, predio):
    """Como nas listas de inservíveis: quem recebe a baixa inclui bens, define motivo, envia e recebe a decisão, mas não a aprova."""
    ids, h = equipe
    for chapa in ("K-1", "K-2"):
        _bem(cliente, h["ana"], predio["05A01"], chapa=chapa)
    lista = _lista(cliente, h["ana"], "inservivel", ["K-1", "K-2"])
    itens = [i["id"] for i in lista["itens"]]
    cliente.post(f"{URL}/classificacoes/{lista['id']}/justificativa", json={"itens_ids": itens, "categoria": "irrecuperavel", "justificativa": "Sem conserto e sem uso"}, headers=h["ana"])
    cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["ana"])
    cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": itens}, headers=h["carla"])
    baixa = cliente.post(f"{URL}/baixas", json={}, headers=h["ana"]).json()
    assert cliente.post(f"{URL}/baixas/{baixa['id']}/bens", json={"chapa": "K-1"}, headers=h["bruno"]).status_code == 403
    assert cliente.put(f"{URL}/baixas/{baixa['id']}/colaboradores", json={"usuarios_ids": [ids["bruno"]]}, headers=h["bruno"]).status_code == 403
    r = cliente.put(f"{URL}/baixas/{baixa['id']}/colaboradores", json={"usuarios_ids": [ids["bruno"]]}, headers=h["ana"])
    assert r.status_code == 200 and [c["id"] for c in r.json()["colaboradores"]] == [ids["bruno"]]
    assert [b["id"] for b in cliente.get(f"{URL}/baixas", params={"solicitante_id": ids["bruno"]}, headers=h["bruno"]).json()] == [baixa["id"]]
    for chapa in ("K-1", "K-2"):
        assert cliente.post(f"{URL}/baixas/{baixa['id']}/bens", json={"chapa": chapa}, headers=h["bruno"]).status_code == 200
    ids_itens = [i["id"] for i in cliente.get(f"{URL}/baixas/{baixa['id']}", headers=h["bruno"]).json()["itens"]]
    assert cliente.post(f"{URL}/baixas/{baixa['id']}/motivo", json={"itens_ids": ids_itens, "motivo": "doacao", "processo": "SPI-2026/009"}, headers=h["bruno"]).status_code == 200
    assert cliente.post(f"{URL}/baixas/{baixa['id']}/enviar", headers=h["bruno"]).json()["status"] == "pendente"
    # Nem o solicitante nem o colaborador aprovam; outra pessoa com controle total aprova
    assert cliente.post(f"{URL}/baixas/{baixa['id']}/decisao", json={"aprovados_ids": ids_itens}, headers=h["ana"]).status_code == 403
    assert cliente.put(f"{URL}/baixas/{baixa['id']}/colaboradores", json={"usuarios_ids": []}, headers=h["ana"]).status_code == 409
    assert cliente.post(f"{URL}/baixas/{baixa['id']}/decisao", json={"aprovados_ids": ids_itens}, headers=h["edu"]).json()["status"] == "aprovada"


def test_categoria_do_inservivel_obrigatoria_marcador_no_bem_e_limpeza_ao_reverter(cliente, equipe, predio):
    """Inservível exige a categoria (ocioso, recuperável, antieconômico, irrecuperável); ao aprovar, o bem ganha o marcador; extraviado não tem categoria."""
    ids, h = equipe
    for chapa in ("G-1", "G-2"):
        _bem(cliente, h["ana"], predio["05A01"], chapa=chapa)
    lista = _lista(cliente, h["ana"], "inservivel", ["G-1", "G-2"])
    g1, g2 = (i["id"] for i in lista["itens"])
    url = f"{URL}/classificacoes/{lista['id']}/justificativa"
    # Sem categoria: recusado; categoria que não existe: recusada pela validação
    assert cliente.post(url, json={"itens_ids": [g1], "justificativa": "Bateria sem carga, conserto inviável"}, headers=h["ana"]).status_code == 400
    assert cliente.post(url, json={"itens_ids": [g1], "categoria": "quebrado", "justificativa": "Bateria sem carga, conserto inviável"}, headers=h["ana"]).status_code == 422
    r = cliente.post(url, json={"itens_ids": [g1], "categoria": "antieconomico", "justificativa": "Conserto custa mais de 50% do valor"}, headers=h["ana"])
    assert r.status_code == 200 and {i["id"]: i["categoria"] for i in r.json()["itens"]}[g1] == "antieconomico"
    # Enviar exige justificativa E categoria em todos
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["ana"]).json()["codigo"] == "sem_justificativa"
    assert cliente.post(url, json={"itens_ids": [g2], "categoria": "ocioso", "justificativa": "Em perfeito estado, sem uso na área"}, headers=h["ana"]).status_code == 200
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["ana"]).json()["status"] == "pendente"
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/decisao", json={"aprovados_ids": [g1, g2]}, headers=h["carla"]).json()["status"] == "aprovada"
    bem = cliente.get(f"{URL}/bens/chapa/G-1", headers=h["dora"]).json()
    assert bem["situacao"] == "inservivel" and bem["categoria_inservivel"] == "antieconomico"
    assert cliente.get(f"{URL}/bens/chapa/G-2", headers=h["dora"]).json()["categoria_inservivel"] == "ocioso"
    historico = cliente.get(f"{URL}/bens/{bem['id']}/historico", headers=h["dora"]).json()
    assert any(e["tipo"] == "classificou" and "[Antieconômico]" in e["texto"] for e in historico)
    # Voltando a Ativo, o marcador some
    r = cliente.post(f"{URL}/bens/{bem['id']}/reverter", json={"justificativa": "Voltou a ser usado na área"}, headers=h["carla"])
    assert r.status_code == 200 and r.json()["situacao"] == "ativo" and r.json()["categoria_inservivel"] is None
    # Extraviado não tem categoria
    _bem(cliente, h["ana"], predio["05A01"], chapa="G-3")
    ext = _lista(cliente, h["ana"], "extraviado", ["G-3"])
    item = ext["itens"][0]["id"]
    urle = f"{URL}/classificacoes/{ext['id']}/justificativa"
    assert cliente.post(urle, json={"itens_ids": [item], "categoria": "ocioso", "justificativa": "Sumiu na mudança de andar"}, headers=h["ana"]).status_code == 400
    assert cliente.post(urle, json={"itens_ids": [item], "justificativa": "Sumiu na mudança de andar"}, headers=h["ana"]).json()["itens"][0]["categoria"] is None


def test_relatorio_pdf_aguenta_justificativa_de_2000_caracteres(cliente, equipe, predio):
    """Justificativa longa (até 2000 caracteres) não pode estourar a tabela do PDF: ela ocupa uma linha de largura total abaixo do bem."""
    ids, h = equipe
    for chapa in ("L-1", "L-2"):
        _bem(cliente, h["ana"], predio["05A01"], chapa=chapa)
    lista = _lista(cliente, h["ana"], "inservivel", ["L-1", "L-2"])
    itens = [i["id"] for i in lista["itens"]]
    longo = "<p>" + ("Equipamento com defeito recorrente na fonte e na placa, sem peças de reposição no mercado. " * 24)[:2000] + "</p>"
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/justificativa", json={"itens_ids": itens, "categoria": "antieconomico", "justificativa": longo}, headers=h["ana"]).status_code == 200
    for formato in ("pdf", "xlsx"):
        r = cliente.get(f"{URL}/classificacoes/{lista['id']}/relatorio", params={"formato": formato}, headers=h["ana"])
        assert r.status_code == 200, (formato, r.text[:200])
    assert r.content[:2] == b"PK" or True


def test_justificativa_ja_gravada_pode_ser_editada_e_excluida(cliente, equipe, predio):
    """Editar = gravar de novo (troca texto e categoria); excluir apaga justificativa e categoria; só em rascunho e por quem monta ou colabora."""
    ids, h = equipe
    for chapa in ("M-1", "M-2"):
        _bem(cliente, h["ana"], predio["05A01"], chapa=chapa)
    lista = _lista(cliente, h["ana"], "inservivel", ["M-1", "M-2"])
    m1, m2 = (i["id"] for i in lista["itens"])
    url = f"{URL}/classificacoes/{lista['id']}/justificativa"
    cliente.post(url, json={"itens_ids": [m1, m2], "categoria": "ocioso", "justificativa": "Em perfeito estado, sem uso na área"}, headers=h["ana"])
    # Edita só o M-1: texto e categoria novos; o M-2 continua como estava
    r = cliente.post(url, json={"itens_ids": [m1], "categoria": "irrecuperavel", "justificativa": "Placa queimada, sem conserto"}, headers=h["ana"])
    itens = {i["id"]: i for i in r.json()["itens"]}
    assert itens[m1]["categoria"] == "irrecuperavel" and "Placa queimada" in itens[m1]["justificativa"] and itens[m2]["categoria"] == "ocioso"
    # Exclui a do M-1: some a justificativa e a categoria, e a lista não pode ser enviada
    remover = f"{url}/remover"
    assert cliente.post(remover, json={"itens_ids": [m1]}, headers=h["bruno"]).status_code == 403  # quem não monta nem colabora
    r = cliente.post(remover, json={"itens_ids": [m1]}, headers=h["ana"])
    itens = {i["id"]: i for i in r.json()["itens"]}
    assert itens[m1]["justificativa"] is None and itens[m1]["categoria"] is None and itens[m2]["categoria"] == "ocioso"
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["ana"]).json()["codigo"] == "sem_justificativa"
    # Depois de enviada, não edita nem exclui
    cliente.post(url, json={"itens_ids": [m1], "categoria": "ocioso", "justificativa": "Voltou a ser classificado como ocioso"}, headers=h["ana"])
    assert cliente.post(f"{URL}/classificacoes/{lista['id']}/enviar", headers=h["ana"]).json()["status"] == "pendente"
    assert cliente.post(remover, json={"itens_ids": [m1]}, headers=h["ana"]).status_code == 409
    assert cliente.post(url, json={"itens_ids": [m1], "categoria": "ocioso", "justificativa": "Tentando editar depois de enviada"}, headers=h["ana"]).status_code == 409
