# Criado por José Eduardo Santana Martins
# Este arquivo serve para testar férias e licença-prêmio (Módulo RH, funcionalidade 2).
"""Férias e licença-prêmio: regras parametrizáveis, aprovação (autorizador, substituto, CGP), painel e exportação."""

from datetime import date

import pytest

from app.core.banco import FabricaSessao
from app.models.rh import Afastamento
from app.models.setor import Setor
from app.models.usuario import Usuario
from app.services.rh import servico_afastamentos
from app.tarefas import mensageria
from tests.apoio_rh import criar_cgp, funcionais, simular_smtp
from tests.conftest import cabecalho, criar_usuario
from tests.test_smtp import SmtpSimulado, dados_servidor

URL = "/api/rh/afastamentos"
HOJE = date(2026, 10, 1)  # quinta-feira


@pytest.fixture(autouse=True)
def _ambiente(monkeypatch):
    simular_smtp(monkeypatch)
    monkeypatch.setattr("app.services.rh.servico_afastamentos.hoje", lambda: HOJE)
    monkeypatch.setattr("app.services.contratos.servico_contratos.hoje", lambda: HOJE)


@pytest.fixture
def equipe(cliente, admin):
    """Chefe autoriza Ana e Bia (Diretoria A, com a Coordenadoria A1 filha); Rita é da CGP; Sílvia substitui o chefe."""
    ids = {login: criar_usuario(login, nome_completo=nome, email=f"{login}@sp.gov.br", departamento=depto) for login, nome, depto in (
        ("chefe", "Chefe Silva", "Diretoria A"), ("ana", "Ana Souza", "Diretoria A"), ("bia", "Bia Lima", "Coordenadoria A1"),
        ("rh", "Rita RH", "Coordenadoria de Gestão de Pessoas"), ("sub", "Sílvia Substituta", "Diretoria A"))}
    criar_cgp()
    with FabricaSessao() as sessao:
        diretoria = Setor(nome="Diretoria A")
        sessao.add(diretoria)
        sessao.flush()
        sessao.add(Setor(nome="Coordenadoria A1", setor_pai_id=diretoria.id))
        sessao.commit()
        diretoria_id = diretoria.id
    for pessoa in ("ana", "bia"):
        funcionais(ids[pessoa], autorizador_id=ids["chefe"], exercicio=2026, inicio_aquisitivo_dia=1, inicio_aquisitivo_mes=1, saldo_lp_dias=10)
    funcionais(ids["chefe"], autorizador_id=ids["rh"], substituto_id=ids["sub"], exercicio=2026, inicio_aquisitivo_dia=1, inicio_aquisitivo_mes=1,
               saldo_lp_dias=0)
    assert cliente.post("/api/smtp/servidores", json=dados_servidor(), headers=admin).status_code == 201
    return ids, {k: cabecalho(cliente, k) for k in ids}, diretoria_id


def _pedir(cliente, h, inicio, fim, tipo="ferias"):
    return cliente.post(URL, json={"tipo": tipo, "inicio": inicio, "fim": fim}, headers=h)


def test_regras_de_agendamento(cliente, equipe):
    ids, h, _ = equipe
    erro = lambda r: (r.status_code, r.json()["detalhe"])  # noqa: E731
    assert "pelo menos 5 dias" in erro(_pedir(cliente, h["ana"], "2026-11-03", "2026-11-05"))[1]
    assert "segunda-feira" in erro(_pedir(cliente, h["ana"], "2026-11-02", "2026-11-10"))[1]
    assert "antecedência" in erro(_pedir(cliente, h["ana"], "2026-10-13", "2026-10-20"))[1]
    assert "exercício" in erro(_pedir(cliente, h["ana"], "2026-12-29", "2027-01-05", "licenca_premio"))[1]
    assert "Saldo de licença-prêmio insuficiente: 10 dia(s)" in erro(_pedir(cliente, h["ana"], "2026-11-03", "2026-11-14", "licenca_premio"))[1]
    assert _pedir(cliente, h["ana"], "2026-11-03", "2026-11-12").status_code == 201
    assert "sobrepõe" in erro(_pedir(cliente, h["ana"], "2026-11-10", "2026-11-16", "licenca_premio"))[1]
    # Saldos independentes: licença-prêmio emendada logo depois das férias (permitido por padrão)
    assert _pedir(cliente, h["ana"], "2026-11-13", "2026-11-17", "licenca_premio").status_code == 201
    meus = cliente.get(f"{URL}/meus", headers=h["ana"]).json()
    assert meus["saldos"]["ferias"] == {"saldo": 30, "usado": 10, "disponivel": 20} and meus["saldos"]["licenca_premio"]["disponivel"] == 5
    # A CGP muda as regras: proíbe emenda e libera segunda-feira
    parametros = {**cliente.get("/api/rh/parametros", headers=h["ana"]).json(), "permite_emenda": False, "inicio_vedado_ferias": []}
    parametros = {k: v for k, v in parametros.items() if k not in ("atualizado_por_nome", "atualizado_em")}
    assert cliente.put("/api/rh/parametros", json=parametros, headers=h["ana"]).status_code == 403
    assert cliente.put("/api/rh/parametros", json=parametros, headers=h["rh"]).status_code == 200
    assert "emendar" in erro(_pedir(cliente, h["ana"], "2026-11-18", "2026-11-22"))[1]
    assert _pedir(cliente, h["ana"], "2026-11-23", "2026-11-27").status_code == 201  # segunda-feira liberada


def test_aprovacao_alteracao_cancelamento_e_emails(cliente, equipe, monkeypatch):
    ids, h, _ = equipe
    SmtpSimulado.enviadas.clear()
    pedido = _pedir(cliente, h["ana"], "2026-11-03", "2026-11-12").json()
    with FabricaSessao() as sessao:
        mensageria.enviar_emails_pendentes(sessao)
    assert sorted(para[0] for _, _, para, _ in SmtpSimulado.enviadas) == ["chefe@sp.gov.br", "rh@sp.gov.br"]
    assert SmtpSimulado.enviadas[0][0]["Subject"] == "Ana Souza agendou férias e aguarda aprovação"
    # Bia não decide; a própria Ana não decide; o chefe decide
    assert cliente.post(f"{URL}/{pedido['id']}/aprovar", headers=h["bia"]).status_code == 403
    assert cliente.post(f"{URL}/{pedido['id']}/recusar", json={"justificativa": ""}, headers=h["chefe"]).status_code == 400
    assert [p["id"] for p in cliente.get(f"{URL}/aprovacoes", headers=h["chefe"]).json()] == [pedido["id"]]
    aprovado = cliente.post(f"{URL}/{pedido['id']}/aprovar", headers=h["chefe"]).json()
    assert aprovado["status"] == "aprovado" and aprovado["decidido_por_nome"] == "Chefe Silva"
    # Alterar um aprovado cria um novo pedido; ao aprovar, o anterior é cancelado
    novo = cliente.put(f"{URL}/{pedido['id']}", json={"tipo": "ferias", "inicio": "2026-11-10", "fim": "2026-11-19"}, headers=h["ana"]).json()
    assert novo["status"] == "pendente" and novo["substitui_id"] == pedido["id"]
    cliente.post(f"{URL}/{novo['id']}/aprovar", headers=h["rh"])
    estados = {a["id"]: a["status"] for a in cliente.get(f"{URL}/meus", headers=h["ana"]).json()["afastamentos"]}
    assert estados == {pedido["id"]: "cancelado", novo["id"]: "aprovado"}
    # Fora do prazo (menos de 5 dias do início) o usuário não cancela; a CGP cancela
    monkeypatch.setattr("app.services.rh.servico_afastamentos.hoje", lambda: date(2026, 11, 7))
    r = cliente.post(f"{URL}/{novo['id']}/cancelar", json={}, headers=h["ana"])
    assert r.status_code == 400 and "prazo" in r.json()["detalhe"]
    assert cliente.post(f"{URL}/{novo['id']}/cancelar", json={"justificativa": "Pedido da chefia"}, headers=h["rh"]).json()["status"] == "cancelado"
    # O usuário recebeu e-mail a cada mudança de status
    SmtpSimulado.enviadas.clear()
    with FabricaSessao() as sessao:
        mensageria.enviar_emails_pendentes(sessao)
    assuntos = [m["Subject"] for m, _, para, _ in SmtpSimulado.enviadas if para == ["ana@sp.gov.br"]]
    assert assuntos.count("Férias aprovada(s)") == 2 and "Afastamento cancelado: férias" in assuntos


def test_substituto_decide_com_o_autorizador_afastado(cliente, equipe, monkeypatch):
    ids, h, _ = equipe
    pedido = _pedir(cliente, h["ana"], "2026-11-24", "2026-12-03").json()
    with FabricaSessao() as sessao:
        sessao.add(Afastamento(usuario_id=ids["chefe"], tipo="ferias", inicio=date(2026, 11, 3), fim=date(2026, 11, 20), dias=18,
                               exercicio=2026, status="aprovado"))
        sessao.commit()
    monkeypatch.setattr("app.services.rh.servico_afastamentos.hoje", lambda: date(2026, 11, 10))
    assert cliente.post(f"{URL}/{pedido['id']}/aprovar", headers=h["chefe"]).status_code == 403
    assert cliente.get("/api/rh/papeis", headers=h["sub"]).json()["autorizador"] is True
    assert cliente.post(f"{URL}/{pedido['id']}/aprovar", headers=h["sub"]).json()["status"] == "aprovado"
    # Depois do fim, a tarefa diária marca como gozado
    with FabricaSessao() as sessao:
        assert servico_afastamentos.marcar_gozados(sessao, date(2026, 12, 4)) == 2


def test_painel_escopo_setor_alerta_e_exportacao(cliente, equipe):
    ids, h, diretoria = equipe
    for pessoa in ("ana", "bia"):
        pedido = _pedir(cliente, h[pessoa], "2026-11-03", "2026-11-12").json()
        cliente.post(f"{URL}/{pedido['id']}/aprovar", headers=h["chefe"])
    _pedir(cliente, h["chefe"], "2026-11-10", "2026-11-16")
    params = {"visao": "mensal", "ano": 2026, "mes": 11}
    # CGP vê todos; o chefe vê os autorizados e a si mesmo; a Ana só a si
    assert len(cliente.get(f"{URL}/painel", params=params, headers=h["rh"]).json()["periodos"]) == 3
    chefe = cliente.get(f"{URL}/painel", params=params, headers=h["chefe"]).json()
    assert {p["nome"] for p in chefe["periodos"]} == {"Ana Souza", "Bia Lima", "Chefe Silva"}
    assert {p["nome"] for p in cliente.get(f"{URL}/painel", params=params, headers=h["ana"]).json()["periodos"]} == {"Ana Souza"}
    # Setor com filhos: Diretoria A inclui a Coordenadoria A1 (Bia)
    filtro = cliente.get(f"{URL}/painel", params={**params, "setor_id": diretoria, "tipo": "ferias"}, headers=h["rh"]).json()
    assert {p["nome"] for p in filtro["periodos"]} == {"Ana Souza", "Bia Lima", "Chefe Silva"}
    # Alerta: limite padrão de 3 pessoas do mesmo setor (Diretoria A: Ana e chefe; Bia é da A1) → baixa o limite para 2
    parametros = {k: v for k, v in cliente.get("/api/rh/parametros", headers=h["rh"]).json().items() if k not in ("atualizado_por_nome", "atualizado_em")}
    cliente.put("/api/rh/parametros", json={**parametros, "limite_alerta_setor": 2}, headers=h["rh"])
    alertas = cliente.get(f"{URL}/painel", params=params, headers=h["chefe"]).json()["alertas"]
    assert alertas == [{"setor": "Diretoria A", "inicio": "2026-11-10", "fim": "2026-11-12", "pessoas": 2}]
    # Anual e exportações
    assert len(cliente.get(f"{URL}/painel", params={"visao": "anual", "ano": 2026}, headers=h["rh"]).json()["periodos"]) == 3
    pdf = cliente.get(f"{URL}/painel/exportar", params={**params, "formato": "pdf"}, headers=h["rh"])
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    xlsx = cliente.get(f"{URL}/painel/exportar", params={**params, "formato": "xlsx"}, headers=h["rh"])
    assert xlsx.content[:2] == b"PK" and "ferias-licencas-2026-11.xlsx" in xlsx.headers["content-disposition"]
    with FabricaSessao() as sessao:
        assert sessao.query(Usuario).count() >= 5


def test_cgp_lanca_afastamento_em_nome_do_servidor(cliente, equipe):
    ids, h, _ = equipe
    url = f"{URL}/lancamento"
    # Férias já gozadas antes do sistema (retroativas, começando numa segunda-feira e sem antecedência)
    corpo = {"usuario_id": ids["ana"], "tipo": "ferias", "inicio": "2026-08-03", "fim": "2026-08-12", "situacao": "gozado",
             "justificativa": "Férias registradas na planilha da CGP"}
    assert cliente.post(url, json=corpo, headers=h["ana"]).status_code == 403
    r = cliente.post(url, json=corpo, headers=h["rh"])
    assert r.status_code == 201, r.text
    lancado = r.json()
    assert lancado["status"] == "gozado" and lancado["dias"] == 10 and lancado["decidido_por_nome"] == "Rita RH"
    assert any("Lançado pela CGP" in (e["justificativa"] or "") for e in lancado["eventos"])
    # Debitou o saldo do período (30 − 10) e aparece para o servidor
    meus = cliente.get(f"{URL}/meus", headers=h["ana"]).json()
    assert meus["periodo_vigente"]["disponivel"] == 20
    # "Gozado" só para períodos encerrados; sobreposição continua proibida
    futuro = {**corpo, "inicio": "2026-10-05", "fim": "2026-10-09"}
    assert "encerrados" in cliente.post(url, json=futuro, headers=h["rh"]).json()["detalhe"]
    assert "sobrepõe" in cliente.post(url, json={**corpo, "inicio": "2026-08-10", "fim": "2026-08-14"}, headers=h["rh"]).json()["detalhe"]
    # Aprovado sem antecedência (começa em 2 dias); saldo insuficiente só passa com "ignorar saldo"
    r = cliente.post(url, json={**futuro, "inicio": "2026-10-03", "fim": "2026-10-24", "situacao": "aprovado"}, headers=h["rh"])
    assert r.status_code == 400 and "insuficiente" in r.json()["detalhe"]
    r = cliente.post(url, json={**futuro, "inicio": "2026-10-03", "fim": "2026-10-24", "situacao": "aprovado", "ignorar_saldo": True}, headers=h["rh"])
    assert r.status_code == 201 and r.json()["status"] == "aprovado"
    assert cliente.post(url, json={**corpo, "justificativa": ""}, headers=h["rh"]).status_code == 422


def test_relatorio_de_saldos(cliente, equipe):
    from io import BytesIO

    from openpyxl import load_workbook

    ids, h, _ = equipe
    assert _pedir(cliente, h["ana"], "2026-11-03", "2026-11-12").status_code == 201
    url = "/api/rh/relatorios/saldos"
    assert cliente.get(url, headers=h["ana"]).status_code == 403
    r = cliente.get(url, params={"formato": "xlsx"}, headers=h["rh"])
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('saldos-ferias-lp-2026-10-01.xlsx"')
    folha = load_workbook(BytesIO(r.content)).active
    linhas = {row[1]: row for row in folha.iter_rows(values_only=True) if row and row[1] in ids}
    ana = linhas["ana"]
    assert ana[3] == "Chefe Silva" and ana[4] == "2026 · 01/01/2026 a 31/12/2026" and (ana[5], ana[6], ana[7]) == (30, 10, 20)
    assert (ana[10], ana[11], ana[13]) == (2026, 10, 10)
    # Sem início do período aquisitivo: "não informado"
    assert linhas["rh"][4] == "não informado"
    pdf = cliente.get(url, params={"formato": "pdf"}, headers=h["rh"])
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def _definir_superior(login_usuario: str, login_superior: str) -> None:
    with FabricaSessao() as sessao:
        from sqlalchemy import select

        superior = sessao.scalar(select(Usuario).where(Usuario.login == login_superior))
        sessao.scalar(select(Usuario).where(Usuario.login == login_usuario)).gestor_id = superior.id
        sessao.commit()


def _assuntos_enviados() -> list[tuple[str, list[str]]]:
    SmtpSimulado.enviadas.clear()
    with FabricaSessao() as sessao:
        mensageria.enviar_emails_pendentes(sessao)
    return [(m["Subject"], para) for m, _, para, _ in SmtpSimulado.enviadas]


def test_pedido_em_duas_etapas_ciente_do_superior_e_depois_aprovacao(cliente, equipe):
    ids, h, _ = equipe
    _definir_superior("ana", "sub")  # Sílvia é superior da Ana, mas quem aprova é o chefe
    SmtpSimulado.enviadas.clear()
    pedido = _pedir(cliente, h["ana"], "2026-11-03", "2026-11-12").json()
    assert pedido["aguarda_ciencia"] is True
    # Etapa 1: só o superior imediato é avisado (nem o aprovador nem a CGP)
    assert _assuntos_enviados() == [("Ana Souza agendou férias: ciente e de acordo", ["sub@sp.gov.br"])]
    # O aprovador ainda não vê o pedido nem aprova; a Ana não dá o próprio ciente; outro colega não dá ciente
    assert cliente.get(f"{URL}/aprovacoes", headers=h["chefe"]).json() == []
    r = cliente.post(f"{URL}/{pedido['id']}/aprovar", headers=h["chefe"])
    assert r.status_code == 409 and r.json()["codigo"] == "aguarda_ciencia"
    assert cliente.post(f"{URL}/{pedido['id']}/ciencia", headers=h["ana"]).status_code == 403
    assert cliente.post(f"{URL}/{pedido['id']}/ciencia", headers=h["bia"]).status_code == 403
    pendentes = cliente.get(f"{URL}/aprovacoes", headers=h["sub"]).json()
    assert [(p["id"], p["pode_dar_ciencia"], p["pode_decidir"]) for p in pendentes] == [(pedido["id"], True, False)]
    # Ciente e de acordo: não aprova; avisa a Ana e libera o pedido ao aprovador e à CGP
    dado = cliente.post(f"{URL}/{pedido['id']}/ciencia", headers=h["sub"]).json()
    assert dado["status"] == "pendente" and dado["aguarda_ciencia"] is False and dado["ciencia_por_nome"] == "Sílvia Substituta"
    assert cliente.post(f"{URL}/{pedido['id']}/ciencia", headers=h["sub"]).status_code == 400  # já deu
    enviados = _assuntos_enviados()
    assert ("Ciente e de acordo: férias", ["ana@sp.gov.br"]) in enviados
    assert sorted(para[0] for assunto, para in enviados if "aguarda aprovação" in assunto) == ["chefe@sp.gov.br", "rh@sp.gov.br"]
    # Etapa 2
    assert [p["id"] for p in cliente.get(f"{URL}/aprovacoes", headers=h["chefe"]).json()] == [pedido["id"]]
    assert cliente.post(f"{URL}/{pedido['id']}/aprovar", headers=h["chefe"]).json()["status"] == "aprovado"
    historico = [e["justificativa"] for e in cliente.get(f"{URL}/meus", headers=h["ana"]).json()["afastamentos"][0]["eventos"]]
    assert "Ciente e de acordo do superior imediato" in historico


def test_superior_recusa_na_etapa_1_e_etapa_dispensada(cliente, equipe):
    ids, h, _ = equipe
    _definir_superior("ana", "sub")
    pedido = _pedir(cliente, h["ana"], "2026-11-03", "2026-11-12").json()
    assert cliente.post(f"{URL}/{pedido['id']}/recusar", json={"justificativa": ""}, headers=h["sub"]).status_code == 400
    recusado = cliente.post(f"{URL}/{pedido['id']}/recusar", json={"justificativa": "Equipe desfalcada"}, headers=h["sub"]).json()
    assert recusado["status"] == "recusado" and recusado["justificativa"] == "Equipe desfalcada"
    # Superior que é o próprio aprovador (chefe) ou ausência de superior: vai direto à aprovação, como antes
    _definir_superior("ana", "chefe")
    assert _pedir(cliente, h["ana"], "2026-11-17", "2026-11-26").json()["aguarda_ciencia"] is False
    assert _pedir(cliente, h["bia"], "2026-11-17", "2026-11-26").json()["aguarda_ciencia"] is False


def _abrir_agendamento(data: date | None) -> None:
    with FabricaSessao() as sessao:
        servico_afastamentos.parametros(sessao).abertura_agendamento_ferias = data
        sessao.commit()


def test_agendamento_do_proximo_exercicio_abre_na_data_e_abate_da_janela(cliente, equipe, monkeypatch):
    ids, h, _ = equipe
    # Sem a regra, vale o comportamento de sempre
    assert cliente.get(f"{URL}/meus", headers=h["ana"]).json()["agendamento_antecipado"] is None
    _abrir_agendamento(date(2026, 10, 15))
    meus = cliente.get(f"{URL}/meus", headers=h["ana"]).json()
    assert meus["periodo_vigente"]["exercicio"] == 2026 and meus["proximo_periodo"]["exercicio"] == 2027
    assert meus["agendamento_antecipado"] == {"exercicio": 2027, "inicio": "2027-01-01", "fim": "2027-12-31", "abertura": "2026-10-15",
                                              "aberto": False, "limite_dias": 30, "agendados": 0}
    # Antes da data: o próximo exercício não pode ser agendado; o exercício vigente segue as regras normais
    r = _pedir(cliente, h["ana"], "2027-02-02", "2027-02-11")
    assert r.status_code == 400 and "exercício 2027 abre em 15/10/2026" in r.json()["detalhe"]
    assert _pedir(cliente, h["ana"], "2026-12-01", "2026-12-10").status_code == 201
    # A CGP lança/altera mesmo antes da abertura (não passa pela trava)
    # Depois da data: até 30 dias para o exercício 2027, mesmo antes de o saldo ser implantado, abatendo da janela de 2027
    monkeypatch.setattr("app.services.rh.servico_afastamentos.hoje", lambda: date(2026, 10, 20))
    assert cliente.get(f"{URL}/meus", headers=h["bia"]).json()["agendamento_antecipado"]["aberto"] is True
    assert _pedir(cliente, h["bia"], "2027-02-02", "2027-02-21").status_code == 201  # 20 dias
    r = _pedir(cliente, h["bia"], "2027-06-01", "2027-06-15")  # +15 passaria de 30
    assert r.status_code == 400 and "Saldo de férias do período aquisitivo 01/01/2027 a 31/12/2027 insuficiente" in r.json()["detalhe"]
    assert _pedir(cliente, h["bia"], "2027-06-01", "2027-06-10").status_code == 201  # +10 = 30
    assert cliente.get(f"{URL}/meus", headers=h["bia"]).json()["agendamento_antecipado"]["agendados"] == 30
    # Quem ainda não tem período aquisitivo informado agenda o ano seguinte pelo limite por ano civil
    criar_usuario("novo", nome_completo="Novo Servidor", email="novo@sp.gov.br")
    novo = cabecalho(cliente, "novo")
    assert cliente.get(f"{URL}/meus", headers=novo).json()["agendamento_antecipado"]["exercicio"] == 2027
    assert _pedir(cliente, novo, "2027-03-02", "2027-03-11").status_code == 201
    # Em 2027, com a janela nova em curso, a regra deixa de valer e os dias agendados contam como usados da janela
    monkeypatch.setattr("app.services.rh.servico_afastamentos.hoje", lambda: date(2027, 1, 5))
    meus = cliente.get(f"{URL}/meus", headers=h["bia"]).json()
    assert meus["agendamento_antecipado"] is None and meus["periodo_vigente"]["exercicio"] == 2027
    assert (meus["periodo_vigente"]["usado"], meus["periodo_vigente"]["disponivel"]) == (30, 0)


def test_janela_virando_em_15_03_debita_a_janela_de_cada_pedido(cliente, equipe, monkeypatch):
    """Quem tem a janela virando em 15/03: fevereiro debita o exercício vigente e abril o próximo, cada um com 30 dias."""
    ids, h, _ = equipe
    funcionais(ids["bia"], autorizador_id=ids["chefe"], exercicio=2026, inicio_aquisitivo_dia=15, inicio_aquisitivo_mes=3, saldo_lp_dias=10)
    _abrir_agendamento(date(2026, 10, 15))
    monkeypatch.setattr("app.services.rh.servico_afastamentos.hoje", lambda: date(2026, 10, 20))
    ag = cliente.get(f"{URL}/meus", headers=h["bia"]).json()["agendamento_antecipado"]
    assert (ag["exercicio"], ag["inicio"], ag["fim"]) == (2027, "2027-03-15", "2028-03-14")
    assert _pedir(cliente, h["bia"], "2027-02-02", "2027-03-01").status_code == 201  # 28 dias na janela vigente (2026/27)
    assert _pedir(cliente, h["bia"], "2027-04-06", "2027-05-05").status_code == 201  # 30 dias na janela seguinte
    meus = cliente.get(f"{URL}/meus", headers=h["bia"]).json()
    assert meus["periodo_vigente"]["usado"] == 28 and meus["proximo_periodo"]["usado"] == 30


def test_cancelamento_avisa_cgp_superior_e_aprovador_na_etapa_de_aprovacao(cliente, equipe):
    ids, h, _ = equipe
    _definir_superior("ana", "sub")  # Sílvia é a superior; o chefe é o aprovador; Rita é da CGP
    # Etapa 1 (aguardando o ciente do superior): cancelar avisa a CGP e o superior, mas não o aprovador
    pedido = _pedir(cliente, h["ana"], "2026-11-03", "2026-11-12").json()
    _assuntos_enviados()
    assert cliente.post(f"{URL}/{pedido['id']}/cancelar", json={"justificativa": "Mudança de planos"}, headers=h["ana"]).status_code == 200
    enviados = [(a, p[0]) for a, p in _assuntos_enviados() if "Cancelamento" in a]
    assert sorted(p for _, p in enviados) == ["rh@sp.gov.br", "sub@sp.gov.br"]
    assert enviados[0][0] == "Cancelamento: Ana Souza cancelou férias"
    # Etapa 2 (aguardando o aprovador): o aprovador também é avisado
    pedido = _pedir(cliente, h["ana"], "2026-11-17", "2026-11-26").json()
    assert cliente.post(f"{URL}/{pedido['id']}/ciencia", headers=h["sub"]).status_code == 200
    _assuntos_enviados()
    assert cliente.post(f"{URL}/{pedido['id']}/cancelar", json={}, headers=h["ana"]).status_code == 200
    destinos = sorted(p[0] for a, p in _assuntos_enviados() if "Cancelamento" in a)
    assert destinos == ["chefe@sp.gov.br", "rh@sp.gov.br", "sub@sp.gov.br"]
    # Pedido aprovado: avisa a CGP e o superior (o aprovador já decidiu); a própria Ana só recebe a confirmação
    pedido = _pedir(cliente, h["ana"], "2026-12-01", "2026-12-10").json()
    cliente.post(f"{URL}/{pedido['id']}/ciencia", headers=h["sub"])
    assert cliente.post(f"{URL}/{pedido['id']}/aprovar", headers=h["chefe"]).status_code == 200
    _assuntos_enviados()
    assert cliente.post(f"{URL}/{pedido['id']}/cancelar", json={}, headers=h["ana"]).status_code == 200
    todos = _assuntos_enviados()
    assert sorted(p[0] for a, p in todos if "Cancelamento" in a) == ["rh@sp.gov.br", "sub@sp.gov.br"]
    assert [p[0] for a, p in todos if a.startswith("Afastamento cancelado")] == ["ana@sp.gov.br"]
