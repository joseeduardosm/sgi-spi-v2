# Criado por José Eduardo Santana Martins
# Este arquivo serve para testar a carga em lote dos dados funcionais do RH por planilha XLSX.
"""Carga de dados funcionais: modelo com os servidores, prévia com erros por linha, importação e célula vazia preservada."""

from io import BytesIO

from openpyxl import load_workbook

from tests.apoio_rh import funcionais
from tests.test_rh_afastamentos import _ambiente, equipe  # noqa: F401  (fixtures reaproveitadas; hoje = 01/10/2026)

BASE = "/api/rh/cadastro/funcionais/importacao"


def _planilha(linhas: list[dict], modelo: bytes) -> bytes:
    """Preenche o modelo: cada dict tem o login e as colunas (pelo título) a escrever."""
    livro = load_workbook(BytesIO(modelo))
    aba = livro["Dados funcionais"]
    titulos = [c.value for c in aba[1]]
    por_login = {aba.cell(row=r, column=1).value: r for r in range(2, aba.max_row + 1)}
    for dados in linhas:
        linha = por_login.get(dados["Login"]) or aba.max_row + 1
        for titulo, valor in dados.items():
            aba.cell(row=linha, column=titulos.index(titulo) + 1, value=valor)
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _enviar(cliente, h, rota, conteudo):
    return cliente.post(rota, files={"arquivo": ("dados.xlsx", conteudo, "application/octet-stream")}, headers=h)


def test_modelo_previa_e_importacao(cliente, equipe):  # noqa: F811
    ids, h, _ = equipe
    assert cliente.get(f"{BASE}/modelo", headers=h["ana"]).status_code == 403
    modelo = cliente.get(f"{BASE}/modelo", headers=h["rh"])
    assert modelo.status_code == 200
    aba = load_workbook(BytesIO(modelo.content))["Dados funcionais"]
    logins = [aba.cell(row=r, column=1).value for r in range(2, aba.max_row + 1)]
    assert "bia" in logins and "root" not in logins
    # Valores atuais já vêm preenchidos (Ana: autorizador chefe, início 01/01)
    linha_ana = next(r for r in range(2, aba.max_row + 1) if aba.cell(row=r, column=1).value == "ana")
    assert aba.cell(row=linha_ana, column=4).value == "chefe" and aba.cell(row=linha_ana, column=7).value == "01/01"

    completa = {"Jornada (horas/semana)": 40, "Horário de trabalho": "9:00 às 18:00", "Intervalo de almoço e descanso": "12h às 13h",
                "RG/CIN nº": "12.345.678-9", "RS/PV nº": "1.234.567/8", "Regime de plantão": "Não", "Horário de estudante": "Sim"}
    conteudo = _planilha([
        {"Login": "bia", **completa, "Dias disponíveis no período vigente": 12},
        {"Login": "sub", "Autorizador (login)": "chefe", "Início do período aquisitivo (dd/mm)": "15/03"},
        {"Login": "rh", "Autorizador (login)": "fantasma", "Horário de trabalho": "nove às seis", "Topo da hierarquia": "talvez"},
        {"Login": "desconhecido"},
    ], modelo.content)
    previa = _enviar(cliente, h["rh"], f"{BASE}/previa", conteudo)
    assert previa.status_code == 200, previa.text
    p = previa.json()
    por_login = {l["login"]: l for l in p["linhas"]}
    assert not p["gravado"] and p["com_erro"] == 2
    assert "horario_estudante" in por_login["bia"]["mudancas"] and "dias disponíveis do período vigente" in por_login["bia"]["mudancas"]
    erros_rh = " | ".join(por_login["rh"]["erros"])
    assert "fantasma" in erros_rh and "Horário de trabalho" in erros_rh and "Sim ou Não" in erros_rh
    assert "não encontrado" in por_login["desconhecido"]["erros"][0]
    assert por_login["ana"]["mudancas"] == []  # linha sem alteração (célula vazia mantém o valor)
    # Com erro, nada é gravado
    r = _enviar(cliente, h["rh"], BASE, conteudo)
    assert r.status_code == 400 and "nada foi gravado" in r.json()["detalhe"]

    corrigida = _planilha([{"Login": "bia", **completa, "Dias disponíveis no período vigente": 12},
                           {"Login": "sub", "Autorizador (login)": "chefe", "Início do período aquisitivo (dd/mm)": "15/03"}], modelo.content)
    r = _enviar(cliente, h["rh"], BASE, corrigida)
    assert r.status_code == 200, r.text
    assert r.json()["gravado"] and r.json()["com_mudanca"] == 2
    bia = cliente.get(f"/api/rh/cadastro/usuarios/{ids['bia']}", headers=h["rh"]).json()["funcionais"]
    assert (bia["jornada_semanal_horas"], bia["horario_trabalho_inicio"], bia["intervalo_fim"], bia["horario_estudante"]) == (40, "09:00", "13:00", True)
    assert bia["autorizador_nome"] == "Chefe Silva"  # célula vazia manteve o autorizador
    vigente = next(p for p in bia["periodos"] if p["vigente"])
    assert (vigente["disponivel"], vigente["origem"]) == (12, "ajuste_cgp")
    sub = cliente.get(f"/api/rh/cadastro/usuarios/{ids['sub']}", headers=h["rh"]).json()["funcionais"]
    assert sub["inicio_periodo_aquisitivo"] == "15/03" and sub["autorizador_nome"] == "Chefe Silva"
    # Reenviar o modelo atual sem mexer em nada não altera ninguém (nem os dias disponíveis)
    atual = cliente.get(f"{BASE}/modelo", headers=h["rh"]).content
    assert _enviar(cliente, h["rh"], f"{BASE}/previa", atual).json()["com_mudanca"] == 0


def test_arquivo_invalido(cliente, equipe):  # noqa: F811
    _, h, _ = equipe
    r = cliente.post(f"{BASE}/previa", files={"arquivo": ("dados.xlsx", b"nao e planilha", "application/octet-stream")}, headers=h["rh"])
    assert r.status_code == 400 and ".xlsx" in r.json()["detalhe"]
    r = cliente.post(f"{BASE}/previa", files={"arquivo": ("dados.csv", b"a;b", "text/csv")}, headers=h["rh"])
    assert r.status_code == 400
