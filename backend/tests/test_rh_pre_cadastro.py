# Criado por José Eduardo Santana Martins
# Este arquivo serve para testar a ficha de pré-cadastro do RH: validações, login único (SGI e AD), vínculo automático com a conta e busca.
"""Pré-cadastro do RH."""

import pytest

from app.services.cliente_ldap import IdentidadeLdap
from app.services.rh.servico_pre_cadastro import cpf_valido, pis_valido
from tests.apoio_rh import criar_cgp
from tests.conftest import FabricaSessao, cabecalho, criar_usuario

RH = "/api/rh/pre-cadastros"
CPF_OK = "529.982.247-25"
PIS_OK = "12345678919"


def test_validadores_de_documentos():
    assert cpf_valido("52998224725") and not cpf_valido("52998224724") and not cpf_valido("11111111111") and not cpf_valido("123")
    assert pis_valido(PIS_OK) and not pis_valido("12345678910") and not pis_valido("00000000000")


@pytest.fixture
def rh(cliente, admin, monkeypatch):
    """Usuário do RH (controle total) e o AD simulado: só o login `ocupado.no.ad` existe lá."""
    criar_cgp(criar_usuario("rh-ficha", nome_completo="Rita do RH"))
    monkeypatch.setattr("app.services.servico_ldap.situacao_do_login", lambda sessao, login: ("encontrado", "Já existe no AD.") if login == "ocupado.no.ad" else ("livre", "Login livre no AD."))
    return cabecalho(cliente, "rh-ficha")


def _ficha(**extra):
    return {"nome_completo": "Maria da Silva Souza", "cpf": CPF_OK, "pis_pasep": PIS_OK, "cep": "01310-100", "uf": "sp", "rs_pv": "123456-7", "celular": "11 99999-0000", **extra}


def test_criar_normaliza_valida_e_guarda_formacoes(cliente, rh):
    r = cliente.post(RH, json=_ficha(login_proposto="Maria.Souza", status="aguardando_conta", formacoes=[{"nivel": "Graduação", "curso": "Direito", "instituicao": "USP", "ano_conclusao": 2015}],
                                      experiencias=[{"empresa": "Órgão X", "cargo": "Analista", "inicio": "2016-01-01"}]), headers=rh)
    assert r.status_code == 201, r.text
    f = r.json()
    assert f["cpf"] == "52998224725" and f["cep"] == "01310100" and f["uf"] == "SP" and f["login_proposto"] == "maria.souza" and f["status"] == "aguardando_conta"
    assert len(f["formacoes"]) == 1 and len(f["experiencias"]) == 1
    # CPF inválido, PIS inválido e login no formato errado
    assert cliente.post(RH, json=_ficha(cpf="111.111.111-11"), headers=rh).status_code == 400
    assert cliente.post(RH, json=_ficha(pis_pasep="12345678910"), headers=rh).status_code == 400
    assert cliente.post(RH, json=_ficha(cpf="168.995.350-09", login_proposto="1x"), headers=rh).status_code == 400
    # Aguardando conta exige login
    assert cliente.post(RH, json=_ficha(cpf="111.444.777-35", status="aguardando_conta"), headers=rh).status_code == 400
    # CPF repetido em outra ficha é 409
    assert cliente.post(RH, json=_ficha(login_proposto="outro.login"), headers=rh).status_code == 409


def test_login_unico_no_sgi_na_ficha_e_no_ad(cliente, rh):
    criar_usuario("ja.existe", nome_completo="Alguém")
    assert cliente.post(RH, json=_ficha(login_proposto="x.primeiro"), headers=rh).status_code == 201
    consultas = {
        "ja.existe": ("usuario", False), "x.primeiro": ("ficha", False), "ocupado.no.ad": ("ad", False), "livre.mesmo": ("livre", True), "Ab": ("formato", False),
    }
    for login, (motivo, livre) in consultas.items():
        r = cliente.get(f"{RH}/login-disponivel", params={"login": login}, headers=rh).json()
        assert (r["motivo"], r["disponivel"]) == (motivo, livre), login
    assert cliente.get(f"{RH}/login-disponivel", params={"login": "livre.mesmo"}, headers=rh).json()["verificado_no_ad"] is True
    # Gravar com login ocupado no AD é 409
    assert cliente.post(RH, json=_ficha(nome_completo="Outra Pessoa", cpf="168.995.350-09", login_proposto="ocupado.no.ad"), headers=rh).status_code == 409
    # A própria ficha pode manter o seu login
    ficha = cliente.get(RH, params={"busca": "maria"}, headers=rh).json()[0]
    r = cliente.put(f"{RH}/{ficha['id']}", json=_ficha(login_proposto="x.primeiro"), headers=rh)
    assert r.status_code == 200, r.text


def test_vinculo_automatico_quando_a_conta_aparece(cliente, rh):
    ficha = cliente.post(RH, json=_ficha(login_proposto="nova.conta", status="aguardando_conta", data_nascimento="1990-05-20"), headers=rh).json()
    from app.models.diretorio_ldap import DiretorioLdap
    from app.services import servico_ldap

    with FabricaSessao() as sessao:
        diretorio = DiretorioLdap(nome="AD teste", servidor="ad.local", porta=389, usar_ssl=False, base_dn="DC=a,DC=b", bind_dn="x", senha_bind_cifrada="x", ativo=True)
        sessao.add(diretorio)
        sessao.commit()
        servico_ldap.aplicar_fotografia(sessao, diretorio, [IdentidadeLdap(login="nova.conta", nome_completo="Maria da Silva Souza", email="maria@spi.sp.gov.br", dn="CN=m", id_externo="guid-1")])
        sessao.commit()
    r = cliente.get(f"{RH}/{ficha['id']}", headers=rh).json()
    assert r["status"] == "vinculado" and r["usuario_id"] and r["vinculado_em"]
    # Conta vinculada: ficha não muda mais nem se cancela; histórico registra o vínculo
    assert cliente.put(f"{RH}/{ficha['id']}", json=_ficha(), headers=rh).status_code == 409
    assert cliente.delete(f"{RH}/{ficha['id']}", headers=rh).status_code == 409
    hist = cliente.get(f"/api/rh/usuarios/{r['usuario_id']}/historico", headers=rh).json()
    assert any("vinculada à ficha" in e["titulo"] for e in hist["eventos"])
    # A busca de usuários encontra pela ficha (CPF) para o RH
    achados = cliente.get("/api/usuarios", params={"busca": "52998224725"}, headers=rh).json()
    assert [u["login"] for u in achados["itens"]] == ["nova.conta"]


def test_cancelar_libera_o_login_e_so_o_rh_acessa(cliente, rh):
    ficha = cliente.post(RH, json=_ficha(login_proposto="vai.cancelar"), headers=rh).json()
    assert cliente.delete(f"{RH}/{ficha['id']}", headers=rh).status_code == 204
    assert cliente.get(f"{RH}/login-disponivel", params={"login": "vai.cancelar"}, headers=rh).json()["disponivel"] is True
    assert cliente.get(f"{RH}/9999", headers=rh).status_code == 404
    criar_usuario("sem-rh", nome_completo="Sem RH")
    h = cabecalho(cliente, "sem-rh")
    assert cliente.get(RH, headers=h).status_code == 403 and cliente.post(RH, json=_ficha(), headers=h).status_code == 403


def test_importacao_pelo_formulario_xlsx_de_uma_pessoa(cliente, rh):
    """O modelo é o formulário de uma pessoa (uma aba, protegida, só as respostas liberadas); a prévia aponta erros sem gravar e a importação cria a ficha em rascunho."""
    import io

    from openpyxl import load_workbook

    r = cliente.get(f"{RH}/importacao/modelo", headers=rh)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    modelo = r.content
    livro = load_workbook(io.BytesIO(modelo))
    assert livro.sheetnames == ["Ficha"]
    aba = livro["Ficha"]
    assert aba.protection.sheet is True
    destino = lambda nome: next(iter(livro.defined_names[nome].destinations))[1].replace("$", "")
    # Rótulos bloqueados e respostas liberadas; listas e alertas nas respostas
    assert aba[destino("f_cpf")].protection.locked is False and aba["B14"].protection.locked is True
    validacoes = {str(v.sqref): v for v in aba.data_validations.dataValidation}
    assert validacoes[destino("f_sexo")].type == "list" and "Feminino" in validacoes[destino("f_sexo")].formula1
    assert validacoes[destino("f_cpf")].type == "custom" and "CPF" in validacoes[destino("f_cpf")].error and validacoes[destino("f_cpf")].showErrorMessage
    assert validacoes[destino("f_email_pessoal")].error.startswith("Digite um e-mail")
    # Instruções no topo da mesma aba
    assert "COMO PREENCHER" in str(aba["B3"].value)

    def preencher(**valores):
        wb = load_workbook(io.BytesIO(modelo))
        for chave, valor in valores.items():
            wb["Ficha"][next(iter(wb.defined_names[chave].destinations))[1].replace("$", "")] = valor
        saida = io.BytesIO()
        wb.save(saida)
        return saida.getvalue()

    def enviar(rota, conteudo, nome="ficha.xlsx"):
        return cliente.post(f"{RH}/importacao{rota}", files={"arquivo": (nome, io.BytesIO(conteudo), "application/octet-stream")}, headers=rh)

    # Em branco e com CPF inválido: erros na prévia, nada gravado
    em_branco = enviar("/previa", modelo).json()["linhas"][0]
    assert any("Nome completo em branco" in e for e in em_branco["erros"]) and any("CPF em branco" in e for e in em_branco["erros"])
    ruim = enviar("/previa", preencher(f_nome_completo="Ana Souza", f_cpf="111.111.111-11")).json()
    assert ruim["com_erro"] == 1 and any("CPF inválido" in e for e in ruim["linhas"][0]["erros"])
    assert enviar("", preencher(f_nome_completo="Ana Souza", f_cpf="111.111.111-11")).status_code == 400
    assert cliente.get(RH, headers=rh).json() == []
    # Preenchida certa: prévia sem erro e importação cria o rascunho com formações e experiências
    ok = preencher(f_nome_completo="Ana Souza", f_cpf="529.982.247-25", f_data_nascimento=__import__("datetime").datetime(1990, 3, 15), f_sexo="Feminino", f_celular="(11) 99999-8888",
                   f_email_pessoal="ana@exemplo.com", f_cep="01310-100", f_uf="SP", f_vinculo_tipo="Servidor efetivo cedido", f_rs_pv="123456-7",
                   fo1_nivel="Graduação", fo1_curso="Direito", fo1_instituicao="USP", fo1_ano_conclusao=2015, fo1_em_andamento="Não",
                   ex1_empresa="Órgão X", ex1_cargo="Analista", ex1_inicio=__import__("datetime").datetime(2016, 1, 1), ex1_fim=__import__("datetime").datetime(2019, 12, 31), ex1_descricao="Atividades")
    previa = enviar("/previa", ok).json()
    assert previa["com_erro"] == 0 and previa["gravado"] is False and previa["linhas"][0]["nome"] == "Ana Souza"
    r = enviar("", ok)
    assert r.status_code == 201, r.text
    ficha_id = r.json()["linhas"][0]["ficha_id"]
    detalhe = cliente.get(f"{RH}/{ficha_id}", headers=rh).json()
    assert detalhe["status"] == "rascunho" and detalhe["cpf"] == "52998224725" and detalhe["data_nascimento"] == "1990-03-15" and detalhe["cep"] == "01310100"
    assert detalhe["formacoes"][0]["curso"] == "Direito" and detalhe["experiencias"][0]["fim"] == "2019-12-31" and detalhe["cargo_id"] is None and detalhe["login_proposto"] is None
    # CPF já cadastrado, fim anterior ao início e arquivos que não são o formulário
    assert any("já está na ficha" in e for e in enviar("/previa", preencher(f_nome_completo="Outra Ana", f_cpf="529.982.247-25")).json()["linhas"][0]["erros"])
    invertido = enviar("/previa", preencher(f_nome_completo="Beto", f_cpf="168.995.350-09", ex1_empresa="X", ex1_inicio=__import__("datetime").datetime(2020, 1, 1), ex1_fim=__import__("datetime").datetime(2019, 1, 1))).json()
    assert any("fim não pode ser anterior" in e for e in invertido["linhas"][0]["erros"])
    assert cliente.post(f"{RH}/importacao/previa", files={"arquivo": ("x.txt", io.BytesIO(b"x"), "text/plain")}, headers=rh).status_code == 400
    from openpyxl import Workbook

    outra = io.BytesIO()
    Workbook().save(outra)
    assert enviar("/previa", outra.getvalue()).status_code == 400
