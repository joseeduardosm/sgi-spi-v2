# Criado por José Eduardo Santana Martins
# Este arquivo serve para testar as subtarefas e as dependências entre tarefas.
"""Subtarefas (`POST /{numero}/subtarefas`) e dependências (`PUT /{numero}/dependencias`)."""

from datetime import datetime, timedelta, timezone

from tests.test_tarefas import URL, _nova, _prazo, equipe  # noqa: F401  (fixture da equipe)


def _mover(cliente, h, numero, acao, **extra):
    return cliente.post(f"{URL}/{numero}/mover", json={"acao": acao, **extra}, headers=h)


def test_subtarefa_herda_da_mae_e_a_mae_se_move_sem_depender_dela(cliente, equipe):
    ids, h, equipe_id = equipe
    marcador = cliente.post(f"{URL}/equipes/{equipe_id}/marcadores", json={"nome": "Urgente"}, headers=h["lia"]).json()
    mae = _nova(cliente, h["lia"], equipe_id, responsaveis_ids=[ids["ana"], ids["beto"]], marcadores_ids=[marcador["id"]])
    r = cliente.post(f"{URL}/{mae['numero']}/subtarefas", json={"titulo": "Levantar dados"}, headers=h["lia"])
    assert r.status_code == 201, r.text
    sub = r.json()
    assert sub["tarefa_pai"]["numero"] == mae["numero"] and sub["equipe"]["id"] == equipe_id and [m["id"] for m in sub["marcadores"]] == [marcador["id"]]
    assert sub["responsavel"]["id"] == ids["ana"] and len(sub["responsaveis"]) == 2 and sub["prazo"] == mae["prazo"]
    assert sub["pode_criar_subtarefa"] is False  # subtarefa não tem subtarefa
    assert cliente.post(f"{URL}/{sub['numero']}/subtarefas", json={"titulo": "Neto"}, headers=h["lia"]).status_code == 400
    # Prazo da subtarefa não passa do da mãe
    depois = (datetime.fromisoformat(mae["prazo"]) + timedelta(days=1)).isoformat()
    assert cliente.post(f"{URL}/{mae['numero']}/subtarefas", json={"titulo": "Tarde", "prazo": depois}, headers=h["lia"]).status_code == 400
    # Quem não edita a mãe não cria
    assert cliente.post(f"{URL}/{mae['numero']}/subtarefas", json={"titulo": "X"}, headers=h["caio"]).status_code in (403, 404)

    detalhe = cliente.get(f"{URL}/{mae['numero']}", headers=h["lia"]).json()
    assert [s["numero"] for s in detalhe["subtarefas"]] == [sub["numero"]] and detalhe["subtarefas_total"] == 1 and detalhe["subtarefas_concluidas"] == 0
    # A mãe se move livremente com subtarefa aberta; a subtarefa só ganha um registro, sem sair do lugar
    antes = cliente.get(f"{URL}/{sub['numero']}", headers=h["lia"]).json()
    assert _mover(cliente, h["ana"], mae["numero"], "iniciar").status_code == 200
    assert _mover(cliente, h["ana"], mae["numero"], "entregar").status_code == 200
    assert _mover(cliente, h["lia"], mae["numero"], "validar").status_code == 200
    depois = cliente.get(f"{URL}/{sub['numero']}", headers=h["lia"]).json()
    assert (depois["status"], depois["estagio_id"], depois["versao"]) == (antes["status"], antes["estagio_id"], antes["versao"])
    linha = cliente.get(f"{URL}/{sub['numero']}/linha-do-tempo", headers=h["lia"]).json()["itens"]
    registros = [e["titulo"] for e in linha if f"Tarefa mãe #{mae['numero']} movida" in e["titulo"]]
    assert len(registros) == 3 and "Em andamento → Em validação" in " ".join(registros)
    # A subtarefa segue o próprio pipeline, independente da mãe
    _mover(cliente, h["ana"], sub["numero"], "iniciar")
    assert _mover(cliente, h["ana"], sub["numero"], "entregar").status_code == 200
    assert _mover(cliente, h["lia"], sub["numero"], "validar").status_code == 200


def test_subtarefas_ficam_fora_das_listas_e_tem_quadro_proprio(cliente, equipe):
    ids, h, equipe_id = equipe
    mae = _nova(cliente, h["lia"], equipe_id, responsavel_id=ids["ana"])
    sub = cliente.post(f"{URL}/{mae['numero']}/subtarefas", json={"titulo": "Parte 1"}, headers=h["lia"]).json()
    # Equipe, minhas e pessoa só têm a mãe
    for params, quem in (({"escopo": "equipe", "equipe_id": equipe_id}, "lia"), ({"escopo": "minhas"}, "ana"), ({"escopo": "pessoa", "login": "ana"}, "lia")):
        numeros = [i["numero"] for i in cliente.get(URL, params=params, headers=h[quem]).json()["itens"]]
        assert mae["numero"] in numeros and sub["numero"] not in numeros, params
    # Quadro próprio das subtarefas
    r = cliente.get(URL, params={"escopo": "subtarefas", "tarefa": mae["numero"]}, headers=h["lia"])
    assert r.status_code == 200, r.text
    corpo = r.json()
    assert [i["numero"] for i in corpo["itens"]] == [sub["numero"]]
    assert corpo["contexto"]["tipo"] == "subtarefas" and corpo["contexto"]["titulo"].endswith(f"Tarefa {mae['titulo']} - Subtarefas")
    assert corpo["contexto"]["equipe_id"] == equipe_id and corpo["contexto"]["tarefa_numero"] == mae["numero"]
    assert cliente.get(URL, params={"escopo": "subtarefas"}, headers=h["lia"]).status_code == 400
    # Quem não vê a mãe não vê o quadro; busca global também não traz a subtarefa
    assert cliente.get(URL, params={"escopo": "subtarefas", "tarefa": mae["numero"]}, headers=h["caio"]).status_code == 404
    achados = cliente.get("/api/busca", params={"q": "Parte 1"}, headers=h["ana"]).json()
    assert all(x.get("id") != str(sub["numero"]) for x in (achados if isinstance(achados, list) else achados.get("resultados", [])))


def test_mudar_estagio_da_mae_registra_nas_subtarefas(cliente, equipe):
    ids, h, equipe_id = equipe
    mae = _nova(cliente, h["lia"], equipe_id, responsavel_id=ids["ana"])
    sub = cliente.post(f"{URL}/{mae['numero']}/subtarefas", json={"titulo": "Parte 1"}, headers=h["lia"]).json()
    base = f"{URL}/equipes/{equipe_id}/estagios"
    a, b, v, c = [e["id"] for e in cliente.get(base, headers=h["lia"]).json()]
    novo = [{"id": a, "nome": "A fazer", "categoria": "a_fazer", "cor_indice": 0}, {"nome": "Triagem", "categoria": "a_fazer", "cor_indice": 1},
            {"id": b, "nome": "Em andamento", "categoria": "em_andamento", "cor_indice": 4}, {"id": v, "nome": "Em validação", "categoria": "em_validacao", "cor_indice": 5},
            {"id": c, "nome": "Concluída", "categoria": "concluida", "cor_indice": 10}]
    r = cliente.put(base, json={"estagios": novo}, headers=h["lia"])
    assert r.status_code == 200, r.text
    triagem = next(e["id"] for e in r.json() if e["nome"] == "Triagem")
    assert cliente.post(f"{URL}/{mae['numero']}/estagio", json={"estagio_id": triagem}, headers=h["lia"]).status_code == 200
    linha = cliente.get(f"{URL}/{sub['numero']}/linha-do-tempo", headers=h["lia"]).json()["itens"]
    assert any("estágio" in e["titulo"] and "movida" in e["titulo"] for e in linha)
    assert cliente.get(f"{URL}/{sub['numero']}", headers=h["lia"]).json()["estagio_id"] == a


def test_dependencias_bloqueiam_o_inicio_e_impedem_ciclos(cliente, equipe):
    ids, h, equipe_id = equipe
    a = _nova(cliente, h["lia"], equipe_id, responsavel_id=ids["ana"])
    b = _nova(cliente, h["lia"], equipe_id, responsavel_id=ids["ana"])
    c = _nova(cliente, h["lia"], equipe_id, responsavel_id=ids["ana"])
    base = f"{URL}/{b['numero']}/dependencias"
    r = cliente.put(base, json={"numeros": [a["numero"]]}, headers=h["lia"])
    assert r.status_code == 200 and [x["numero"] for x in r.json()["bloqueada_por"]] == [a["numero"]] and r.json()["bloqueada"] is True
    assert [x["numero"] for x in cliente.get(f"{URL}/{a['numero']}", headers=h["lia"]).json()["bloqueia"]] == [b["numero"]]
    # B só inicia depois de A concluída
    r = _mover(cliente, h["ana"], b["numero"], "iniciar")
    assert r.status_code == 400 and f"#{a['numero']}" in r.json()["detalhe"]
    _mover(cliente, h["ana"], a["numero"], "iniciar")
    assert _mover(cliente, h["lia"], a["numero"], "concluir").status_code == 200
    assert cliente.get(f"{URL}/{b['numero']}", headers=h["lia"]).json()["bloqueada"] is False
    assert _mover(cliente, h["ana"], b["numero"], "iniciar").status_code == 200
    # Regras: a si mesma, ciclo (A depende de B que depende de A), número inexistente
    assert cliente.put(base, json={"numeros": [b["numero"]]}, headers=h["lia"]).status_code == 400
    assert cliente.put(f"{URL}/{c['numero']}/dependencias", json={"numeros": [b["numero"]]}, headers=h["lia"]).status_code == 200
    r = cliente.put(f"{URL}/{b['numero']}/dependencias", json={"numeros": [c["numero"]]}, headers=h["lia"])
    assert r.status_code == 400 and "ciclo" in r.json()["detalhe"]
    assert cliente.put(base, json={"numeros": [9999]}, headers=h["lia"]).status_code == 400
    # Limpar a lista libera a tarefa
    assert cliente.put(f"{URL}/{c['numero']}/dependencias", json={"numeros": []}, headers=h["lia"]).json()["bloqueada_por"] == []


def test_excluir_a_mae_leva_as_subtarefas(cliente, equipe):
    ids, h, equipe_id = equipe
    mae = _nova(cliente, h["lia"], equipe_id, responsavel_id=ids["ana"])
    sub = cliente.post(f"{URL}/{mae['numero']}/subtarefas", json={"titulo": "Parte 1"}, headers=h["lia"]).json()
    assert cliente.delete(f"{URL}/{mae['numero']}", headers=h["lia"]).status_code == 204
    assert cliente.get(f"{URL}/{sub['numero']}", headers=h["lia"]).status_code == 404


def test_memorial_da_tarefa_pdf_e_xlsx_com_todos_os_acontecimentos(cliente, equipe):
    import io

    from openpyxl import load_workbook

    ids, h, equipe_id = equipe
    mae = _nova(cliente, h["lia"], equipe_id, responsavel_id=ids["ana"])
    cliente.post(f"{URL}/{mae['numero']}/comentarios", data={"texto": "Primeiro comentário"}, headers=h["lia"])
    assert _mover(cliente, h["ana"], mae["numero"], "iniciar").status_code == 200
    detalhe = cliente.get(f"{URL}/{mae['numero']}", headers=h["lia"]).json()
    assert detalhe["dias_em_aberto"] == 0
    pdf = cliente.get(f"{URL}/{mae['numero']}/memorial", params={"formato": "pdf"}, headers=h["lia"])
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and "memorial-tarefa" in pdf.headers["content-disposition"]
    xlsx = cliente.get(f"{URL}/{mae['numero']}/memorial", params={"formato": "xlsx"}, headers=h["lia"])
    assert xlsx.status_code == 200
    planilha = load_workbook(io.BytesIO(xlsx.content))
    assert planilha.sheetnames == ["Resumo", "Acontecimentos"]
    texto = " ".join(str(c.value) for linha in planilha["Acontecimentos"].iter_rows() for c in linha if c.value)
    assert "Primeiro comentário" in texto and "Tarefa iniciada" in texto and "Situação: A fazer → Em andamento" in texto
    resumo = " ".join(str(c.value) for linha in planilha["Resumo"].iter_rows() for c in linha if c.value)
    assert "dia(s) em aberto" in resumo
    # Quem não vê a tarefa não emite o memorial
    assert cliente.get(f"{URL}/{mae['numero']}/memorial", headers=h["caio"]).status_code == 404


def test_subtarefa_em_validacao_aparece_na_mae_e_nos_indicadores_sem_mudar_a_mae(cliente, equipe):
    ids, h, equipe_id = equipe
    mae = _nova(cliente, h["lia"], equipe_id, responsavel_id=ids["ana"])
    sub = cliente.post(f"{URL}/{mae['numero']}/subtarefas", json={"titulo": "Conferir"}, headers=h["lia"]).json()
    assert _mover(cliente, h["ana"], sub["numero"], "iniciar").status_code == 200
    assert _mover(cliente, h["ana"], sub["numero"], "entregar").json()["status"] == "em_validacao"

    lista = cliente.get(URL, params={"escopo": "equipe", "equipe_id": equipe_id}, headers=h["lia"]).json()
    # Os contadores da lista seguem sem subtarefas; o campo novo traz as que aguardam validação
    assert lista["indicadores"]["em_validacao"] == 0 and lista["indicadores"]["subtarefas_em_validacao"] == 1
    item = next(i for i in lista["itens"] if i["numero"] == mae["numero"])
    assert item["status"] == "a_fazer" and item["subtarefas_em_validacao"] == 1
    equipes = cliente.get(f"{URL}/equipes", headers=h["lia"]).json()
    assert next(e for e in equipes if e["id"] == equipe_id)["indicadores"]["subtarefas_em_validacao"] == 1
    # Validada, o selo some
    assert _mover(cliente, h["lia"], sub["numero"], "validar").status_code == 200
    lista = cliente.get(URL, params={"escopo": "equipe", "equipe_id": equipe_id}, headers=h["lia"]).json()
    assert lista["indicadores"]["subtarefas_em_validacao"] == 0


def test_resumo_traz_os_numeros_das_tarefas_que_bloqueiam(cliente, equipe):
    ids, h, equipe_id = equipe
    a = _nova(cliente, h["lia"], equipe_id, responsavel_id=ids["ana"], titulo="A")
    b = _nova(cliente, h["lia"], equipe_id, responsavel_id=ids["ana"], titulo="B")
    c = _nova(cliente, h["lia"], equipe_id, responsavel_id=ids["ana"], titulo="C")
    r = cliente.put(f"{URL}/{c['numero']}/dependencias", json={"numeros": [a["numero"], b["numero"]]}, headers=h["lia"])
    assert r.status_code == 200, r.text
    lista = cliente.get(URL, params={"escopo": "equipe", "equipe_id": equipe_id}, headers=h["lia"]).json()["itens"]
    por_numero = {i["numero"]: i for i in lista}
    assert por_numero[c["numero"]]["bloqueada_por_numeros"] == sorted([a["numero"], b["numero"]]) and por_numero[a["numero"]]["bloqueada_por_numeros"] == []
    # Concluir uma bloqueadora tira o cadeado, mas a dependência continua desenhada
    assert _mover(cliente, h["lia"], a["numero"], "concluir").status_code == 200
    assert _mover(cliente, h["lia"], b["numero"], "concluir").status_code == 200
    depois = {i["numero"]: i for i in cliente.get(URL, params={"escopo": "equipe", "equipe_id": equipe_id}, headers=h["lia"]).json()["itens"]}
    assert depois[c["numero"]]["bloqueada"] is False and depois[c["numero"]]["bloqueada_por_numeros"] == sorted([a["numero"], b["numero"]])


def test_subtarefa_pode_depender_da_anterior_aberta(cliente, equipe):
    """`depende_da_anterior`: a nova subtarefa nasce bloqueada pela última subtarefa ainda não concluída (e não inicia antes dela)."""
    ids, h, equipe_id = equipe
    mae = _nova(cliente, h["lia"], equipe_id, responsaveis_ids=[ids["ana"]])
    url = f"{URL}/{mae['numero']}/subtarefas"
    a = cliente.post(url, json={"titulo": "A", "depende_da_anterior": True}, headers=h["lia"]).json()
    assert a["bloqueada_por"] == []  # a primeira não tem anterior
    b = cliente.post(url, json={"titulo": "B", "depende_da_anterior": True}, headers=h["lia"]).json()
    assert [x["numero"] for x in b["bloqueada_por"]] == [a["numero"]]
    c = cliente.post(url, json={"titulo": "C"}, headers=h["lia"]).json()
    assert c["bloqueada_por"] == []  # sem a opção, nada muda
    d = cliente.post(url, json={"titulo": "D", "depende_da_anterior": True}, headers=h["lia"]).json()
    assert [x["numero"] for x in d["bloqueada_por"]] == [c["numero"]]
    # B não inicia enquanto A não estiver concluída
    assert _mover(cliente, h["ana"], b["numero"], "iniciar").status_code == 400


def test_subtarefa_nasce_com_varias_dependencias_e_valida_antes_de_criar(cliente, equipe):
    ids, h, equipe_id = equipe
    mae = _nova(cliente, h["lia"], equipe_id, responsaveis_ids=[ids["ana"]])
    outra = _nova(cliente, h["lia"], equipe_id, responsaveis_ids=[ids["ana"]])
    url = f"{URL}/{mae['numero']}/subtarefas"
    a = cliente.post(url, json={"titulo": "A"}, headers=h["lia"]).json()
    b = cliente.post(url, json={"titulo": "B", "bloqueada_por_numeros": [a["numero"], outra["numero"]]}, headers=h["lia"])
    assert b.status_code == 201, b.text
    assert sorted(x["numero"] for x in b.json()["bloqueada_por"]) == sorted([a["numero"], outra["numero"]])
    antes = len(cliente.get(f"{URL}/{mae['numero']}", headers=h["lia"]).json()["subtarefas"])
    # Não pode depender da própria mãe nem de tarefa inexistente, e nada é criado nesses casos
    assert cliente.post(url, json={"titulo": "C", "bloqueada_por_numeros": [mae["numero"]]}, headers=h["lia"]).status_code == 400
    assert cliente.post(url, json={"titulo": "D", "bloqueada_por_numeros": [999999]}, headers=h["lia"]).status_code == 400
    assert len(cliente.get(f"{URL}/{mae['numero']}", headers=h["lia"]).json()["subtarefas"]) == antes


def test_subtarefa_herda_responsaveis_da_mae_e_acompanha_as_mudancas(cliente, equipe):
    ids, h, equipe_id = equipe
    mae = _nova(cliente, h["lia"], equipe_id, responsaveis_ids=[ids["ana"]])
    # Na criação, os da mãe vêm sempre; o informado entra além deles
    sub = cliente.post(f"{URL}/{mae['numero']}/subtarefas", json={"titulo": "Levantar dados", "responsaveis_ids": [ids["lia"]]}, headers=h["lia"]).json()
    assert [p["id"] for p in sub["responsaveis"]] == [ids["ana"], ids["lia"]] and sub["responsavel"]["id"] == ids["ana"]

    def ler(n):
        return cliente.get(f"{URL}/{n}", headers=h["lia"]).json()

    def editar_mae(responsaveis):
        corpo = {"titulo": mae["titulo"], "prioridade": "normal", "responsaveis_ids": responsaveis, "versao": ler(mae["numero"])["versao"]}
        return cliente.put(f"{URL}/{mae['numero']}", json=corpo, headers=h["lia"])

    # Quem entra na mãe entra na subtarefa (que mantém os seus extras)
    assert editar_mae([ids["ana"], ids["beto"]]).status_code == 200
    assert {p["id"] for p in ler(sub["numero"])["responsaveis"]} == {ids["ana"], ids["beto"], ids["lia"]}
    # A subtarefa não perde os da mãe, mas pode ganhar outros
    corpo = {"titulo": sub["titulo"], "prioridade": "normal", "versao": ler(sub["numero"])["versao"]}
    assert cliente.put(f"{URL}/{sub['numero']}", json={**corpo, "responsaveis_ids": [ids["lia"]]}, headers=h["lia"]).status_code == 400
    assert cliente.put(f"{URL}/{sub['numero']}", json={**corpo, "responsaveis_ids": [ids["ana"], ids["beto"]]}, headers=h["lia"]).status_code == 200
    # Quem sai da mãe sai da subtarefa; o principal passa para quem ficou
    assert editar_mae([ids["beto"]]).status_code == 200
    depois = ler(sub["numero"])
    assert [p["id"] for p in depois["responsaveis"]] == [ids["beto"]] and depois["responsavel"]["id"] == ids["beto"]
