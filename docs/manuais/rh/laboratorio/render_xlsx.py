# Desenha uma planilha XLSX como HTML (para tirar print do formulário sem precisar do Excel): células, mesclagens, cores, bordas, larguras e listas de opções.
import html
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


def _cor(cor):
    rgb = getattr(cor, "rgb", None)
    return "#" + rgb[-6:] if isinstance(rgb, str) and len(rgb) >= 6 and rgb[-6:] != "000000" else None


def renderizar(caminho_ou_bytes, linhas=(1, 60), colunas=(1, 8)) -> str:
    livro = load_workbook(caminho_ou_bytes)
    aba = livro.active
    listas = set()
    for v in aba.data_validations.dataValidation:
        if v.type == "list":
            listas.update(str(v.sqref).split())
    mesclas = {}
    ocupadas = set()
    for m in aba.merged_cells.ranges:
        c1, l1, c2, l2 = range_boundaries(str(m))
        mesclas[(l1, c1)] = (l2 - l1 + 1, c2 - c1 + 1)
        for l in range(l1, l2 + 1):
            for c in range(c1, c2 + 1):
                if (l, c) != (l1, c1):
                    ocupadas.add((l, c))
    corpo = []
    for l in range(linhas[0], linhas[1] + 1):
        altura = (aba.row_dimensions[l].height or 15) * 1.33
        celulas = []
        for c in range(colunas[0], colunas[1] + 1):
            if (l, c) in ocupadas:
                continue
            cel = aba.cell(row=l, column=c)
            estilo = []
            fundo = _cor(cel.fill.fgColor) if cel.fill and cel.fill.fill_type == "solid" else None
            if fundo:
                estilo.append(f"background:{fundo}")
            if cel.font:
                estilo.append(f"font-size:{(cel.font.size or 10) * 1.3:.0f}px;font-weight:{'700' if cel.font.bold else '400'};color:{_cor(cel.font.color) or '#000'}")
            if cel.border and cel.border.left and cel.border.left.style:
                estilo.append("border:1px solid #bfc5cc")
            if cel.alignment and cel.alignment.wrap_text:
                estilo.append("white-space:pre-wrap")
            rs, cs = mesclas.get((l, c), (1, 1))
            texto = html.escape(str(cel.value)) if cel.value is not None else ""
            seta = '<span class="seta">▼</span>' if cel.coordinate in listas else ""
            celulas.append(f'<td id="c_{cel.coordinate}" rowspan="{rs}" colspan="{cs}" style="{";".join(estilo)}">{texto}{seta}</td>')
        corpo.append(f'<tr style="height:{altura:.0f}px">{"".join(celulas)}</tr>')
    larguras = "".join(f'<col style="width:{(aba.column_dimensions[get_column_letter(c)].width or 8.43) * 7.4:.0f}px">' for c in range(colunas[0], colunas[1] + 1))
    return f"""<html><head><meta charset="utf-8"><style>
    body {{ margin: 0; background: #e9ecef; font-family: Calibri, Arial, sans-serif; }}
    .folha {{ width: fit-content; margin: 24px auto; background: #fff; box-shadow: 0 1px 6px rgba(0,0,0,.25); }}
    table {{ border-collapse: collapse; table-layout: fixed; }} td {{ padding: 2px 6px; vertical-align: middle; overflow: hidden; }}
    .seta {{ float: right; color: #6b7280; font-size: 11px; margin-top: 3px; border-left: 1px solid #bfc5cc; padding-left: 5px; }}
    .alerta {{ position: absolute; width: 420px; background: #f0f0f0; border: 1px solid #6e6e6e; box-shadow: 4px 4px 12px rgba(0,0,0,.35); font: 14px Segoe UI, Arial, sans-serif; z-index: 5; }}
    .alerta .barra {{ background: #fff; padding: 6px 10px; border-bottom: 1px solid #ccc; }} .alerta .msg {{ padding: 14px; display: flex; gap: 12px; background: #fff; }}
    .alerta .x {{ color: #fff; background: #d9261c; border-radius: 50%; width: 30px; height: 30px; flex: none; text-align: center; line-height: 30px; font-weight: 700; }}
    .alerta .botoes {{ padding: 10px; text-align: right; }} .alerta button {{ margin-left: 6px; padding: 3px 18px; border: 1px solid #adadad; background: #e1e1e1; }}
    .dica {{ position: absolute; width: 330px; background: #ffffe1; border: 1px solid #777; padding: 5px 8px; font: 13px Segoe UI, Arial, sans-serif; box-shadow: 2px 2px 5px rgba(0,0,0,.3); z-index: 5; }}
    </style></head><body><div class="folha"><table><colgroup>{larguras}</colgroup>{"".join(corpo)}</table></div></body></html>"""
