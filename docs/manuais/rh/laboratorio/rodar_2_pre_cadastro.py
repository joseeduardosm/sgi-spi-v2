# Prints do manual: pré-cadastro (lista, ficha, login proposto) e ficha por planilha (formulário XLSX, avisos e importação).
import io, sys; sys.path.insert(0, ".")
from cap import *
import render_xlsx
from openpyxl import load_workbook


def ir(pg, rota, espera=800):
    pg.goto(BASE + rota); pg.wait_for_load_state("networkidle"); pg.wait_for_timeout(espera)


def preencher(modelo: bytes, **valores) -> bytes:
    livro = load_workbook(io.BytesIO(modelo))
    for chave, valor in valores.items():
        livro["Ficha"][next(iter(livro.defined_names[chave].destinations))[1].replace("$", "")] = valor
    saida = io.BytesIO(); livro.save(saida); return saida.getvalue()


def coord(modelo: bytes, nome: str) -> str:
    return next(iter(load_workbook(io.BytesIO(modelo)).defined_names[nome].destinations))[1].replace("$", "")


with sync_playwright() as p:
    nav, ctx = abrir(p)
    pg = entrar(ctx, "rita")

    # Lista de fichas
    ir(pg, "/rh/pre-cadastros")
    foto(pg, "p01-lista-fichas", [(pg.get_by_role("button", name="Ficha por planilha"), 1, "cima"), (pg.get_by_role("link", name="Nova ficha"), 2, "cima"), (pg.locator(".chips-filtro"), 3, "esq"),
                                  (pg.locator(".tabela-gestao tbody tr").first.locator("td").nth(4), 4, "esq"), (pg.locator(".tabela-gestao tbody tr").first.get_by_role("link"), 5, "cima")],
         legendas=["Preenche uma ficha por planilha: baixa o formulário em Excel e importa o que a pessoa devolveu", "Abre a ficha em branco para o RH preencher na tela", "Filtra por situação: rascunho, aguardando a conta, vinculada, cancelada ou aguardando validação (fichas que a própria pessoa enviou)", "Situação da ficha", "Abre a ficha"])

    # Ficha na tela
    ir(pg, "/rh/pre-cadastros/nova")
    foto(pg, "p02-ficha-identificacao", [(pg.get_by_role("button", name="Preencher por planilha (formulário XLSX)"), 1, "esq"), ("#f-nome_completo", 2, "dir"), ("#f-data_nascimento", 3, "esq")],
         legendas=["Alternativa: gerar o formulário em Excel e deixar a pessoa preencher", "Só o nome completo e o CPF são obrigatórios; o resto pode ser completado depois", "Datas no formato dd/mm/aaaa"], rolar="#sec-identificacao")
    foto(pg, "p03-ficha-documentos", [("#f-cpf", 1, "esq"), ("#f-pis_pasep", 2, "esq"), ("#f-rg_uf", 3, "cima")],
         legendas=["CPF: os dígitos verificadores são conferidos ao salvar, e o mesmo CPF não pode estar em duas fichas", "PIS/PASEP: também conferido", "UF, listas de opções e demais documentos (RG, CIN, título, CTPS, CNH)"], rolar="#sec-documentos")
    pg.select_option("#f-setor", label="Subsecretaria de Gestão Corporativa"); pg.wait_for_timeout(300)
    pg.select_option("#f-cargo", index=1); pg.wait_for_timeout(300)
    foto(pg, "p04-ficha-vinculo", [("#f-vinculo_tipo", 1, "esq"), ("#f-setor", 2, "esq"), ("#f-cargo", 3, "cima")],
         legendas=["Tipo de vínculo: servidor cedido, comissionado sem vínculo, empregado público ou outro", "Setor pretendido: só aparecem setores com vaga livre", "Cargo pretendido: as vagas livres do setor escolhido. Trocar o setor limpa o cargo"], rolar="#sec-vinculo")
    pg.get_by_role("button", name="+ Formação").click(); pg.get_by_role("button", name="+ Experiência").click(); pg.wait_for_timeout(300)
    pg.fill("#fo-curso-0", "Gestão Pública"); pg.fill("#fo-inst-0", "Universidade Exemplo"); pg.fill("#ex-emp-0", "Prefeitura Exemplo"); pg.fill("#ex-cargo-0", "Analista")
    foto(pg, "p05-formacao", [("#fo-nivel-0", 1, "esq"), (pg.get_by_role("button", name="Remover formação"), 2, "cima"), (pg.get_by_role("button", name="+ Formação"), 3, "dir")],
         legendas=["Nível da formação (lista) e, ao lado, curso, instituição e ano de conclusão", "Remove a linha", "Acrescenta outra formação (quantas forem necessárias)"], rolar="#sec-formacoes")
    foto(pg, "p05b-experiencia", [("#ex-emp-0", 1, "esq"), ("#ex-desc-0", 2, "dir"), (pg.get_by_role("button", name="+ Experiência"), 3, "dir")],
         legendas=["Órgão ou empresa, cargo, início e fim", "Descrição das atividades", "Acrescenta outra experiência"], rolar="#sec-experiencias")
    pg.fill("#f-login", "ana.souza"); pg.wait_for_timeout(1800)
    foto(pg, "p06-login-ocupado", [("#f-login", 1, "esq"), (pg.locator("#f-login ~ small"), 2, "dir")],
         legendas=["Login que o RH propõe para a TI criar no AD (letras minúsculas, números, ponto, hífen e sublinhado)", "O sistema confere na hora no SGI, nas outras fichas e no AD. Aqui o login já existe, e a ficha não deixa salvar"], rolar="#sec-login")
    pg.fill("#f-login", "carlos.novo"); pg.wait_for_timeout(1800)
    foto(pg, "p07-login-livre", [(pg.locator("#f-login ~ small"), 1, "dir"), (pg.locator("label.opcao").last, 2, "cima")],
         legendas=["Login livre: pode seguir", "Ficha pronta: aguardando a criação da conta. Marcada, exige o login e a ficha fica esperando a conta aparecer. Desmarcada, é um rascunho"], rolar="#sec-login")

    # Ficha por planilha: o formulário em Excel
    modelo = open("modelo.xlsx", "rb").read()  # gerado por gerar_modelo.sh com o código do sistema
    pg.get_by_role("button", name="Preencher por planilha (formulário XLSX)").click(); pg.wait_for_timeout(600)
    foto(pg, "p08-janela-planilha", [(pg.locator(".passos-importacao li").nth(0), 1, "esq"), (pg.locator(".passos-importacao li").nth(1), 2, "esq"), (pg.locator(".passos-importacao li").nth(2), 3, "esq")],
         legendas=["Baixe o formulário (XLSX) e envie à pessoa", "Quando ela devolver o arquivo preenchido, escolha-o aqui; cada arquivo é a ficha de uma pessoa", "Confira a prévia e importe: a ficha entra como rascunho, e o RH completa o setor e o cargo pretendidos, o login proposto e o início previsto"])
    pg.keyboard.press("Escape")
    pg_x = ctx.new_page(); pg_x.set_viewport_size({"width": 1366, "height": 860})
    pg_x.set_content(render_xlsx.renderizar(io.BytesIO(modelo), (1, 36))); pg_x.wait_for_timeout(300)
    foto(pg_x, "p09-formulario-excel", [("#c_B1", 1, "esq"), ("#c_B3", 2, "esq"), ("#c_B7", 3, "esq"), (f"#c_{coord(modelo, 'f_nome_completo')}", 4, "dir"), (f"#c_{coord(modelo, 'f_sexo')}", 5, "dir")],
         legendas=["Cabeçalho da ficha", "Instruções de preenchimento, no topo da própria planilha", "Células cinzas: o formulário, bloqueadas (não dá para alterar)", "Células brancas: onde a pessoa preenche; com * são obrigatórias", "Células com a seta ▼ têm lista de opções: é só escolher"])
    # Aviso do Excel quando o preenchimento está fora do padrão (simulação do aviso, com o texto real da regra da planilha)
    livro = load_workbook(io.BytesIO(modelo))
    validacoes = {str(v.sqref): v for v in livro["Ficha"].data_validations.dataValidation}
    c_cpf = coord(modelo, "f_cpf"); v = validacoes[c_cpf]; v_mail = validacoes[coord(modelo, "f_email_pessoal")]
    pg_x.set_content(render_xlsx.renderizar(io.BytesIO(modelo), (1, 36)))
    pg_x.evaluate("""([cel, texto, titulo, msg, prompt_titulo, prompt]) => {
      const td = document.getElementById('c_' + cel); td.textContent = texto; td.style.outline = '2px solid #1a73e8';
      const r = td.getBoundingClientRect();
      const a = document.createElement('div'); a.className = 'alerta'; a.style.left = (r.left + 40) + 'px'; a.style.top = (r.bottom + 12 + window.scrollY) + 'px';
      a.innerHTML = '<div class="barra">Microsoft Excel</div><div class="msg"><div class="x">✕</div><div><b>' + titulo + '</b><br>' + msg + '</div></div><div class="botoes"><button>Repetir</button><button>Cancelar</button></div>';
      document.body.appendChild(a);
    }""", [c_cpf, "123.456", v.errorTitle, v.error, v.promptTitle, v.prompt])
    foto(pg_x, "p10-aviso-excel", [(".alerta", 1, "dir"), (f"#c_{c_cpf}", 2, "esq")],
         legendas=["Se o que foi digitado fugir do padrão, o Excel bloqueia e explica como preencher certo (aqui, o CPF com a quantidade errada de números). O mesmo vale para PIS, CEP, e-mail, datas, ano e textos longos demais", "A célula não aceita o valor até ser corrigido (Repetir) ou desfeito (Cancelar). Ao selecionar uma célula, uma dica mostra o que escrever"])
    pg_x.close()

    # Fórmula da ficha: tabelas de formação e experiências
    pg_y = ctx.new_page(); pg_y.set_viewport_size({"width": 1100, "height": 860})
    linha_fo = int(coord(modelo, "fo1_curso")[1:]); linha_ex = int(coord(modelo, "ex1_empresa")[1:])
    pg_y.set_content(render_xlsx.renderizar(io.BytesIO(modelo), (linha_fo - 3, linha_ex + 5))); pg_y.wait_for_timeout(300)
    foto(pg_y, "p10b-formacao-excel", [(f"#c_{coord(modelo, 'fo1_curso')}", 1, "dir"), (f"#c_{coord(modelo, 'ex1_descricao')}", 2, "esq")],
         legendas=["Formação: até 5 linhas; deixe em branco as que sobrarem", "Experiências anteriores: até 5 linhas; a descrição das atividades ocupa duas colunas"])
    pg_y.close()

    # Importação: arquivo certo e arquivo com erro
    pg.get_by_role("button", name="Preencher por planilha (formulário XLSX)").click(); pg.wait_for_timeout(500)
    ok = preencher(modelo, f_nome_completo="Rafael Almeida Santos", f_cpf="168.995.350-09", f_sexo="Masculino", f_celular="(11) 98888-7777", f_email_pessoal="rafael@exemplo.com", f_uf="SP", f_cep="01310-100", fo1_nivel="Graduação", fo1_curso="Administração")
    open("/tmp/ficha-ok.xlsx", "wb").write(ok)
    pg.locator(".modal-importacao-rh input[type=file]").set_input_files("/tmp/ficha-ok.xlsx"); pg.wait_for_timeout(1500)
    foto(pg, "p11-previa-ok", [(pg.locator(".resumo-importacao"), 1, "esq"), (pg.get_by_role("button", name="Importar a ficha"), 2, "cima")],
         legendas=["A prévia confere tudo sem gravar: aqui a ficha está sem erros", "Importar cria a ficha como rascunho (com as formações e experiências do arquivo)"])
    ruim = preencher(modelo, f_nome_completo="Rafael Almeida Santos", f_cpf="111.111.111-11", f_pis_pasep="12345678910")
    open("/tmp/ficha-ruim.xlsx", "wb").write(ruim)
    pg.locator(".modal-importacao-rh input[type=file]").set_input_files("/tmp/ficha-ruim.xlsx"); pg.wait_for_timeout(1500)
    foto(pg, "p12-previa-erro", [(pg.locator(".resumo-importacao"), 1, "esq"), (pg.locator(".texto-erro").first, 2, "dir")],
         legendas=["Com erro, a importação fica bloqueada e nada é gravado", "Cada erro diz o que corrigir na planilha (CPF inválido, PIS inválido, e-mail, datas, CPF já cadastrado...)"])
    nav.close()
